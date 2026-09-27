"""Non-destructive sample reading and factual profiles (no model inference)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import threading
from collections import Counter
from pathlib import Path
from typing import Any, Literal

from src.workbench.intake_models import SampleSource, SourceRow

_CSV_READ_LOCK = threading.Lock()


def _resolve_excel_sheet(book: Any, sheet: str | int | None) -> tuple[str, bool]:
    """把 sheet 选择解析为确定的工作表名，返回（名称, 是否为显式指定）。

    None 读第一个 sheet（默认行为不变）；int 是 1 起始的序号（1=第一个）；
    str 先按名称精确匹配，无匹配且为纯数字时按序号（CLI 传参皆为字符串）。
    解析不到即报错并如实列出全部 sheet，不静默回退。
    """
    names = list(book.sheet_names)
    if sheet is None:
        return names[0], False
    hint = f"该文件含 {len(names)} 个 sheet：{'、'.join(names)}。序号从 1 开始（1=第一个）。"
    if isinstance(sheet, int):
        if not 1 <= sheet <= len(names):
            raise ValueError(f"sheet 序号 {sheet} 超出范围；{hint}")
        return names[sheet - 1], True
    text = str(sheet).strip()
    if not text:
        return names[0], False
    if text in names:
        return text, True
    if text.isdigit():
        ordinal = int(text)
        if 1 <= ordinal <= len(names):
            return names[ordinal - 1], True
        raise ValueError(f"sheet 序号 {ordinal} 超出范围；{hint}")
    raise ValueError(f"找不到 sheet「{text}」；{hint}")


def _read_csv(
    text: str, delimiter: str
) -> tuple[list[str], list[tuple[int, dict[str, Any]]], list[int]]:
    # csv's field limit is process-global. Serialize our readers and restore it even
    # on malformed input; the supplied file already bounds the required field size.
    with _CSV_READ_LOCK:
        previous_limit = csv.field_size_limit()
        try:
            csv.field_size_limit(max(previous_limit, len(text)))
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            records = []
            blank_lines: list[int] = []
            try:
                columns = _headers(next(reader, []))
                while True:
                    first_line = reader.line_num + 1
                    values = next(reader, None)
                    if values is None:
                        break
                    if not values:
                        blank_lines.append(first_line)
                        continue
                    if len(values) != len(columns):
                        raise ValueError(
                            f"第 {first_line} 行有 {len(values)} 列，表头为 {len(columns)} 列；"
                            "请检查分隔符或引号，未静默丢弃该行。"
                        )
                    records.append((first_line, dict(zip(columns, values))))
            except csv.Error as exc:
                raise ValueError(f"CSV 第 {reader.line_num} 行解析失败：{exc}") from exc
            return columns, records, blank_lines
        finally:
            csv.field_size_limit(previous_limit)


def canonical(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )


def content_digest(value: Any) -> str:
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _decode(data: bytes, encoding: str | None) -> tuple[str, str]:
    candidates = [encoding] if encoding else ["utf-8-sig", "gb18030"]
    # UTF-16 BOM(FF FE / FE FF,Excel「Unicode 文本」导出的常见形态)是明确的编码证据:
    # 0xFF/0xFE+0xFF 开头的字节永远不可能是合法 UTF-8 或 GBK,按 BOM 识别不会误判。
    if not encoding and data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        candidates.append("utf-16")
    for candidate in candidates:
        try:
            return data.decode(candidate), candidate
        except UnicodeDecodeError:
            continue
        except LookupError as exc:
            raise ValueError("无法识别指定编码，请选择有效编码。") from exc
    raise ValueError("无法解码文件，请明确指定编码；原始数据未修改。")


def _headers(values: list[Any]) -> list[str]:
    if not values or any(not isinstance(v, str) or not v.strip() for v in values):
        raise ValueError("表头必须是非空文本；请确认文件的第一行是字段名。")
    if len(values) != len(set(values)):
        raise ValueError("存在重复列名，无法明确引用字段；请先区分列名。")
    return values


def _unique_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"JSON 存在重复字段：{key}")
        result[key] = value
    return result


def parse_json_value(text: str) -> Any:
    """Parse structured cells without silently discarding duplicate keys or NaN."""
    value = json.loads(text, object_pairs_hook=_unique_object)
    canonical(value)
    return value


def _excel_sheet_note(sheet_names: list[str], read_name: str, explicit: bool) -> str:
    """多 Sheet 工作簿的读取范围说明:如实列出读了哪个 sheet、哪些未读取。"""

    unread = [name for name in sheet_names if name != read_name]
    listed = "、".join(unread[:5]) + ("等" if len(unread) > 5 else "")
    head = (
        f"该文件含 {len(sheet_names)} 个 sheet，按指定读取「{read_name}」"
        if explicit
        else f"该文件含 {len(sheet_names)} 个 sheet，仅读取第一个「{sheet_names[0]}」"
    )
    return f"{head}；其余 {len(unread)} 个（{listed}）未读取。"


def _excel_notes(
    data: bytes, sheet_name: str, columns: list[str], row_count: int
) -> tuple[str, str, str]:
    """xlsx 合并单元格、无缓存公式格与隐藏行/列的如实说明:点名事实,不自动修复。

    pandas 走 openpyxl 只读模式读值:合并区除左上角外均读为空串;公式格只有
    缓存计算结果才读得到值,由脚本/报表工具写出的 xlsx 常常没有缓存,同样读为
    空串——用户被「缺少监督答案」拦下时无从知道根因。openpyxl 完整加载能拿到
    合并范围与公式格清单(只读模式没有这些),据此如实点名;带缓存值的公式格
    正常读取,不列入。隐藏行/列(筛选后保存或手工隐藏)读取时照常进数据——
    Excel 里看不到的行也会进入分析与训练,同样点名;是否取消隐藏、删除不需要
    的行列由用户决定,没有自动排除。完全落在数据区之外或未读取 sheet 的情况
    不涉及本次读取,不列入。
    """
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    book = load_workbook(io.BytesIO(data), read_only=False)
    try:
        sheet = book[sheet_name]
        merged = sorted(sheet.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col))
        # 数据区公式格:(表头名, 坐标)。data_type=='f' 的格读取时若无比文件缓存值则读空。
        formulas = [
            (columns[cell.column - 1], cell.coordinate)
            for row in sheet.iter_rows(
                min_row=2, max_row=row_count + 1, min_col=1, max_col=len(columns)
            )
            for cell in row
            if cell.data_type == "f"
        ]
        # 隐藏行/列(数据区内):筛选后保存或手工隐藏的行列读取时照常进数据。
        hidden_rows = sorted(
            index
            for index, dim in sheet.row_dimensions.items()
            if dim.hidden and 2 <= index <= row_count + 1
        )
        hidden_columns = []
        for letter, dim in sorted(sheet.column_dimensions.items()):
            if not dim.hidden:
                continue
            position = column_index_from_string(letter)
            if 1 <= position <= len(columns):
                hidden_columns.append(columns[position - 1])
    finally:
        book.close()
    cached: set[str] = set()
    if formulas:
        values_book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
        try:
            values_sheet = values_book[sheet_name]
            cached = {
                coordinate
                for _, coordinate in formulas
                if values_sheet[coordinate].value is not None
            }
        finally:
            values_book.close()
    overlapping = [
        item for item in merged if item.min_row <= row_count + 1 and item.min_col <= len(columns)
    ]
    merged_note = ""
    if overlapping:
        described = [f"{columns[item.min_col - 1]} {item.coord}" for item in overlapping[:5]]
        listing = "、".join(described) + ("等" if len(overlapping) > 5 else "")
        merged_note = (
            f"该 sheet 含 {len(overlapping)} 处合并单元格（{listing}）："
            "合并区除左上角外均读为空值，涉及答案列时这些行会按缺少监督答案处理。"
            "请取消合并并逐行填写受影响的值；没有自动填充。"
        )
    uncached = [(name, coordinate) for name, coordinate in formulas if coordinate not in cached]
    formula_note = ""
    if uncached:
        described = [f"{name} {coordinate}" for name, coordinate in uncached[:5]]
        listing = "、".join(described) + ("等" if len(uncached) > 5 else "")
        formula_note = (
            f"该 sheet 含 {len(uncached)} 个没有缓存计算结果的公式单元格（{listing}）："
            "这些公式读为空值，涉及答案列时这些行会按缺少监督答案处理。"
            "请用 Excel 等软件打开并保存以生成计算结果；没有自动计算。"
        )
    hidden_note = ""
    if hidden_rows:
        described = [f"第 {index} 行" for index in hidden_rows[:5]]
        listing = "、".join(described) + ("等" if len(hidden_rows) > 5 else "")
        hidden_note += (
            f"该 sheet 含 {len(hidden_rows)} 个隐藏行（{listing}）："
            "隐藏行照常读入——Excel 中看不到的行也会进入分析与训练。"
            "请取消隐藏并删除不需要的行；没有自动排除。"
        )
    if hidden_columns:
        listed = "、".join(hidden_columns[:5]) + ("等" if len(hidden_columns) > 5 else "")
        hidden_note += (
            f"该 sheet 含 {len(hidden_columns)} 个隐藏列（{listed}）："
            "隐藏列照常读入——Excel 中看不到的列也会出现在可用字段中。"
            "请删除不需要的列；没有自动排除。"
        )
    return merged_note, formula_note, hidden_note


def _list_line_numbers(lines: list[int]) -> str:
    """行号清单:前 5 个逐一点名,超过 5 个以「等」收尾(与各 note 同口径)。"""
    described = [f"第 {line} 行" for line in lines[:5]]
    return "、".join(described) + ("等" if len(lines) > 5 else "")


def _skipped_blank_lines_note(lines: list[int]) -> str:
    """CSV/JSONL 被跳过空行的如实说明:点名行号,读取行为不变。"""
    return (
        f"该文件有 {len(lines)} 个空行（{_list_line_numbers(lines)}）已跳过——"
        "空行不进入分析与训练。请核对空行位置是否丢了数据；没有自动补行。"
    )


def _kept_blank_rows_note(lines: list[int]) -> str:
    """Excel 全空行(照常读入为全空记录)的如实说明:点名行号,不自动排除。"""
    return (
        f"该 sheet 有 {len(lines)} 个全空行（{_list_line_numbers(lines)}）照常读入——"
        "全空行按缺少监督答案与分组标识处理，会在样例确认或全量验证被拦下。"
        "请删除空行或补全数据；没有自动排除。"
    )


def _dup_header_note(lines: list[int]) -> str:
    """重复表头行(每个单元格都等于其列名的数据行)的如实说明:点名行号,两侧行为分述。"""
    return (
        f"该文件有 {len(lines)} 行与表头完全相同（{_list_line_numbers(lines)}）"
        "照常读入为普通数据行——通常是导出拼接产生的重复表头，输入与答案都会是"
        "列名；样例侧不拦（以「就绪」进入预览、可能参与对比核验），全量侧会被"
        "全量验证硬拦。请删除重复表头行；没有自动删行。"
    )


def read_source(
    name: str,
    data: bytes,
    *,
    scope: Literal["sample", "full"] = "sample",
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | int | None = None,
) -> SampleSource:
    """Read supplied bytes only; never follow paths mentioned inside the data.

    sheet 选择仅对 Excel 有效：按名称或 1 起始的序号指定工作表，None（默认）
    读第一个 sheet，读取行为与此前完全一致。xlsx 读取的 sheet 存在与数据区
    相交的合并单元格、数据区存在没有缓存计算结果的公式格、或数据区存在
    隐藏行/列时，来源分别携带 merged_note / formula_note / hidden_note
    如实点名（合并区除左上角外、无缓存公式格均读为空值；隐藏行/列照常
    读入）；不自动填充、不自动计算、不自动排除，修复由用户决定。

    空行处理跨格式如实点名（blank_note）：CSV/JSONL 的空行读取时被跳过
    （读取行为不变），Excel 的全空行照常读入为全空记录——来源携带
    blank_note 点名行号；不自动补行、不自动排除。Excel 尾部空行在解析时
    自然消失、无从检测，不列入。

    重复表头行跨格式如实点名（dup_header_note）：数据区存在与表头完全相同
    的行（每个单元格都等于其列名，常见于导出拼接）时，来源携带
    dup_header_note 点名行号——该行按普通数据行读入，样例侧不拦，全量侧由
    全量验证硬拦（repeated_header_rows）；没有自动删行。判定口径与全量侧
    完全一致（列数 ≥ 2 且每格等于列名）。
    """
    suffix = Path(name).suffix.lower().lstrip(".")
    digest = hashlib.sha256(data).hexdigest()
    records: list[tuple[int, dict[str, Any]]] = []
    blank_lines: list[int] = []
    actual_encoding, actual_delimiter = "", ""
    resolved_sheet, sheet_note, merged_note, formula_note, hidden_note = "", "", "", "", ""
    if suffix in {"csv", "jsonl"}:
        if sheet is not None:
            raise ValueError(f"sheet 选择仅对 Excel 文件有效；当前文件是 {suffix}。")
        text, actual_encoding = _decode(data, encoding)
        if suffix == "csv":
            if delimiter is not None and delimiter not in (",", ";", "\t", "|"):
                raise ValueError("分隔符支持逗号、分号、Tab 或竖线。")
            try:
                sniffed = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
            except csv.Error:
                sniffed = ","
            actual_delimiter = delimiter or sniffed
            columns, records, blank_lines = _read_csv(text, actual_delimiter)
        else:
            columns = []
            for line_no, line in enumerate(text.splitlines(), 1):
                if not line.strip():
                    blank_lines.append(line_no)
                    continue
                try:
                    record = parse_json_value(line)
                    canonical(record)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f"JSONL 第 {line_no} 行无效：{exc}") from exc
                if not isinstance(record, dict):
                    raise ValueError(f"JSONL 第 {line_no} 行必须是对象。")
                _headers(list(record))
                for key in record:
                    if key not in columns:
                        columns.append(key)
                records.append((line_no, record))
    elif suffix in {"xlsx", "xls"}:
        import pandas as pd

        try:
            # ExcelFile 与 read_excel(BytesIO, sheet_name=0) 走同一条解析路径:
            # 默认仍读第一个 sheet(行为不变),指定 sheet 时按名称选中同一解析器;
            # 借此拿到 sheet 清单,多 Sheet 时如实告知读取范围。
            book = pd.ExcelFile(io.BytesIO(data))
            resolved_sheet, explicit = _resolve_excel_sheet(book, sheet)
            frame = book.parse(
                sheet_name=resolved_sheet, header=None, dtype=object, keep_default_na=False
            )
        except ImportError as exc:
            raise ValueError(
                "读取 Excel 缺少对应引擎，请安装 openpyxl（xlsx）或 xlrd（xls）。"
            ) from exc
        if frame.empty:
            raise ValueError("文件没有表头或数据。")
        columns = _headers(frame.iloc[0].tolist())
        for i, values in enumerate(frame.iloc[1:].itertuples(index=False, name=None), 2):
            normalized = [v.isoformat() if hasattr(v, "isoformat") else v for v in values]
            records.append((i, dict(zip(columns, normalized))))
        sheet_names = list(book.sheet_names)
        if len(sheet_names) > 1:
            sheet_note = _excel_sheet_note(sheet_names, resolved_sheet, explicit)
        # xlsx 检测与数据区相交的合并单元格、数据区内无缓存值的公式格与
        # 数据区内的隐藏行/列 (xls 引擎不提供这些信息,不检测——如实边界)。
        merged_note, formula_note, hidden_note = (
            _excel_notes(data, resolved_sheet, columns, len(records))
            if suffix == "xlsx"
            else ("", "", "")
        )
    else:
        raise ValueError("当前数据入口支持 CSV、Excel、JSONL。")
    # 空行的如实说明(跨格式):CSV/JSONL 的空行读取时被跳过,Excel 的全空行照常
    # 读入为全空记录——两者都不在记录里留下痕迹,用户被拦时无从知道根因。
    # Excel 尾部空行在解析时自然消失、无从检测,不列入。
    blank_note = ""
    if suffix in {"csv", "jsonl"}:
        if blank_lines:
            blank_note = _skipped_blank_lines_note(blank_lines)
    else:
        kept_blank_rows = [
            line for line, record in records if all(is_missing(v) for v in record.values())
        ]
        if kept_blank_rows:
            blank_note = _kept_blank_rows_note(kept_blank_rows)
    # 重复表头行的如实说明(跨格式):每个单元格都等于其列名的行几乎必然是导出
    # 拼接产生的重复表头。判定口径与全量侧 repeated_header_rows 完全一致
    # (列数 ≥ 2 且每格等于列名),两侧披露与拦截永不互相矛盾。样例侧此前对该行
    # 完全无声(场景 30 实测:以「就绪」面目进入预览并可能参与对比核验)。
    dup_header_note = ""
    if len(columns) >= 2:
        dup_header_lines = [
            line
            for line, record in records
            if all(record.get(column) == column for column in columns)
        ]
        if dup_header_lines:
            dup_header_note = _dup_header_note(dup_header_lines)
    if not records:
        raise ValueError("文件没有数据行。")
    rows = [
        SourceRow(row_id=f"r{i:06d}", line=line, values=record)
        for i, (line, record) in enumerate(records, 1)
    ]
    return SampleSource(
        name=Path(name).name,
        digest=digest,
        scope=scope,
        format=suffix,
        encoding=actual_encoding,
        delimiter=actual_delimiter,
        sheet=resolved_sheet,
        sheet_note=sheet_note,
        merged_note=merged_note,
        formula_note=formula_note,
        hidden_note=hidden_note,
        blank_note=blank_note,
        dup_header_note=dup_header_note,
        columns=columns,
        rows=rows,
    )


def _value_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (dict, list)):
        return "object" if isinstance(value, dict) else "array"
    if isinstance(value, (int, float)):
        return "number"
    return "text"


def profile_source(source: SampleSource) -> dict[str, Any]:
    """All counts apply exclusively to the provided file, even when called a sample.

    多 Sheet Excel 的读取范围说明(sheet_note)随来源对象本身携带:read_source
    读取时按实际 sheet 选择生成并持久化,同进程重启后、同一文件按不同 sheet
    重复读取都不会串味;历史会话的 profile 快照原样保留既有标注。
    """
    fields: dict[str, Any] = {}
    evidence: list[str] = []

    def add(row_id: str) -> None:
        if row_id not in evidence:
            evidence.append(row_id)

    for row in source.rows[:3]:
        add(row.row_id)
    for column in source.columns:
        values = [row.values.get(column) for row in source.rows]
        present = [(i, v) for i, v in enumerate(values) if not is_missing(v)]
        counts = Counter(canonical(v) for _, v in present)
        missing_ids = [row.row_id for row, value in zip(source.rows, values) if is_missing(value)]
        if missing_ids:
            add(missing_ids[0])
        if present:
            shortest = min(present, key=lambda pair: len(str(pair[1])))
            longest = max(present, key=lambda pair: len(str(pair[1])))
            add(source.rows[shortest[0]].row_id)
            add(source.rows[longest[0]].row_id)
            rare = min(counts, key=counts.get)
            add(source.rows[next(i for i, v in present if canonical(v) == rare)].row_id)
        texts = [v for _, v in present if isinstance(v, str)]
        numeric_candidates = 0
        for value in texts:
            try:
                numeric_candidates += int(math.isfinite(float(value)))
            except ValueError:
                continue
        fields[column] = {
            "missing_count": len(missing_ids),
            "missing_row_ids": missing_ids[:5],
            "distinct_count": len(counts),
            "value_types": dict(Counter(_value_type(v) for _, v in present)),
            "min_length": min((len(str(v)) for _, v in present), default=0),
            "max_length": max((len(str(v)) for _, v in present), default=0),
            "numeric_text_count": numeric_candidates,
            "leading_zero_examples": [
                v for v in texts if len(v) > 1 and v.startswith("0") and v.isdigit()
            ][:3],
            "common_values": [
                {"value": json.loads(k), "count": n} for k, n in counts.most_common(5)
            ],
        }
    duplicates: dict[str, list[str]] = {}
    for row in source.rows:
        duplicates.setdefault(canonical(row.values), []).append(row.row_id)
    groups = [ids for ids in duplicates.values() if len(ids) > 1]
    profile: dict[str, Any] = {
        "source_digest": source.digest,
        "source_scope": source.scope,
        "record_count": len(source.rows),
        "columns": fields,
        "duplicate_record_count": sum(len(ids) - 1 for ids in groups),
        "duplicate_groups": groups[:10],
        "evidence_row_ids": evidence,
        "scope_note": (
            "以上仅描述用户提供的样例，不能推断全量覆盖、比例或训练就绪。"
            if source.scope == "sample"
            else "以上描述本次提供的文件；统计通过不代表监督含义正确或已具备独立评测条件。"
        ),
    }
    sheet_note = source.sheet_note
    if sheet_note:
        profile["sheet_note"] = sheet_note
    if source.merged_note:
        profile["merged_note"] = source.merged_note
    if source.formula_note:
        profile["formula_note"] = source.formula_note
    if source.hidden_note:
        profile["hidden_note"] = source.hidden_note
    if source.blank_note:
        profile["blank_note"] = source.blank_note
    if source.dup_header_note:
        profile["dup_header_note"] = source.dup_header_note
    return profile
