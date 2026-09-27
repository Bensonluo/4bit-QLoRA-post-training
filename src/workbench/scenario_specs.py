"""Built-in scenario matrix specs: the seed distribution for coverage measurement.

期望结局记录当前真相:unexpected_* 是产品缺陷;「通过但带已知局限」的场景
用 expect_note 记录待办(泄漏预警已上线;高基数标签拦截是已知待办)。
"""

from src.workbench.scenario_matrix import ScenarioSpec

_CLEAN_SAMPLE = ("编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n").encode()
_CLEAN_FULL = (
    "编号,客户描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode()

# GBK 编码样例:utf-8-sig 解码失败,入口自动回退 gb18030(GBK 超集)可正确解码。
_GBK_SAMPLE = "编号,描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n".encode("gbk")
_GBK_FULL = (
    "编号,描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode("gbk")

# 宽表:28 列业务噪声全被用户排除,只留编号(分组)+客户描述(输入)+类别(答案)。
_WIDE_JUNK_COLUMNS = tuple(f"附加信息{i:02d}" for i in range(1, 29))
_WIDE_HEADER = ",".join(("编号", "客户描述", "类别", *_WIDE_JUNK_COLUMNS))
_WIDE_SAMPLE = (
    _WIDE_HEADER + "\n"
    "001,杯子破损,质量," + ",".join(["噪声01"] * 28) + "\n"
    "002,物流未更新,物流," + ",".join(["噪声02"] * 28) + "\n"
).encode()
_WIDE_FULL = (
    _WIDE_HEADER + "\n"
    "001,杯子破损,质量," + ",".join(["备注甲"] * 28) + "\n"
    "002,物流未更新,物流,"
    + ",".join(["备注乙"] * 28)
    + "\n"
    + "".join(
        f"{i:03d},问题{i},{'质量' if i % 2 else '物流'}," + ",".join([f"字段值{i % 7}"] * 28) + "\n"
        for i in range(3, 11)
    )
).encode()

# 混合类型列:工单号同时出现纯数字样与文本样编号,当前路径不做类型区分。
_MIXED_IDS = (
    "10023",
    "AB-1002",
    "10024",
    "AB-1003",
    "10025",
    "10026",
    "AB-1004",
    "10027",
    "10028",
    "AB-1005",
)
_MIXED_SAMPLE = (
    "编号,工单号,客户描述,类别\n001,10023,杯子破损,质量\n002,AB-1002,物流未更新,物流\n".encode()
)
_MIXED_FULL = (
    "编号,工单号,客户描述,类别\n"
    + "".join(
        f"{i:03d},{_MIXED_IDS[i - 1]},问题{i},{'质量' if i % 2 else '物流'}\n" for i in range(1, 11)
    )
).encode()

# Excel「Unicode 文本」导出:UTF-16(带 BOM)+ Tab 分隔,中文 Windows Excel 用户的常见产物。
_UTF16_ROWS = "编号\t客户描述\t类别\n001\t杯子破损\t质量\n002\t物流未更新\t物流\n"
_UTF16_SAMPLE = _UTF16_ROWS.encode("utf-16")
_UTF16_FULL = (
    "编号\t客户描述\t类别\n"
    + "".join(f"{i:03d}\t问题{i}\t{'质量' if i % 2 else '物流'}\n" for i in range(1, 11))
).encode("utf-16")

# 空答案行混入:样例两条都有标签,全量第 005/008 行答案为空(导出缺字段/漏标注)。
_EMPTY_LABEL_FULL = (
    "编号,客户描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode()

# 列名前后空格:外部系统导出的表头带不可见的前后空格,用户按业务口径写「编号/类别」。
_SPACED_HEADER = " 编号 ,客户描述, 类别"
_SPACED_SAMPLE = (_SPACED_HEADER + "\n001,杯子破损,质量\n002,物流未更新,物流\n").encode()
_SPACED_FULL = (
    _SPACED_HEADER + "\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode()

# 超长单行:单个输入单元格里粘贴了数万字符的运行日志(含逗号,按 CSV 规范加引号),
# 单行数十 KB。日志单元格由 csv 模块按规范写出:带引号单元格里的逗号不是分隔符。
_LONG_CELL_PREFIX = (
    "2026-09-27 10:23:01 INFO 收到客户端请求,开始处理订单流程,读取配置项共 42 项;"
    "2026-09-27 10:23:02 WARN 缓存未命中,回源查询耗时 187ms;"
)

# JSONL 超长行:每行一个 JSON 对象,单条记录的输入字段带数万字符日志,
# 单行数十 KB(JSONL 路径逐行 json.loads,无按行截断)。
_JSONL_LONG_CELL = (
    "2026-09-27 10:23:01 INFO 收到客户端请求,处理订单流程,读取配置项共 42 项,回源查询耗时 187ms;"
)


def _jsonl_long_line_text() -> str:
    import json as _json

    lines = [_json.dumps({"编号": "001", "描述": "杯子破损", "类别": "质量"}, ensure_ascii=False)]
    lines.append(
        _json.dumps({"编号": "002", "描述": "物流未更新", "类别": "物流"}, ensure_ascii=False)
    )
    for i in range(3, 11):
        lines.append(
            _json.dumps(
                {
                    "编号": f"{i:03d}",
                    "描述": _JSONL_LONG_CELL * (400 + i * 40),
                    "类别": "质量" if i % 2 else "物流",
                },
                ensure_ascii=False,
            )
        )
    return "\n".join(lines) + "\n"


_JSONL_SAMPLE = (
    "\n".join(
        (
            '{"编号": "001", "描述": "杯子破损", "类别": "质量"}',
            '{"编号": "002", "描述": "物流未更新", "类别": "物流"}',
        )
    )
    + "\n"
).encode()
_JSONL_FULL = _jsonl_long_line_text().encode()

# 重复表头:导出工具把表头行追加在数据中部(常见拼接产物),表头行会变成一条普通数据。
_DUP_HEADER_FULL = (
    "编号,客户描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "编号,客户描述,类别\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode()

# 全角数字:编号与描述里出现全角数字(０１２３),零密钥路径不做全角/半角归一。
_FW_FULL = (
    "编号,客户描述,类别\n"
    "００１,订单１２３号商品杯子破损,质量\n００２,物流未更新,物流\n００３,屏幕碎裂第２次,质量\n"
    "００４,快递丢失,物流\n００５,开不了机,质量\n００６,地址填错,物流\n"
    "００７,异味,质量\n００８,延迟送达,物流\n００９,无法充电,质量\n０１０,包装破损,物流\n"
).encode()
_FW_SAMPLE = (
    "编号,客户描述,类别\n００１,订单１２３号商品杯子破损,质量\n００２,物流未更新,物流\n".encode()
)


# 只有 1 行数据的样例:用户拿一条记录试用产品;表头+1 行共 2 行,该行本身有标签,
# 拦截纯粹因为「一条样例不足以完成配对核验」,与缺标签/脏数据无关。
_ONE_ROW_SAMPLE = "编号,客户描述,类别\n001,杯子破损,质量\n".encode()

# 目标列全空:表头有「类别」但所有值为空(导出漏了标签值列,只保留了列名)。
_ALL_EMPTY_TARGET_SAMPLE = "编号,客户描述,类别\n001,杯子破损,\n002,物流未更新,\n".encode()
_ALL_EMPTY_TARGET_FULL = (
    "编号,客户描述,类别\n" + "".join(f"{i:03d},问题{i},\n" for i in range(1, 11))
).encode()

# 无 BOM 的 UTF-16:Excel「Unicode 文本」导出通常带 BOM,经其他工具转存/再导出可能丢 BOM;
# 交错的 NUL/换行字节让 utf-8-sig 与 gb18030 兜底都解不开(测试钉住这一点,防夹具漂移)。
_UTF16_NOBOM_SAMPLE = _UTF16_ROWS.encode("utf-16-le")
_UTF16_NOBOM_FULL = (
    "编号\t客户描述\t类别\n"
    + "".join(f"{i:03d}\t问题{i}\t{'质量' if i % 2 else '物流'}\n" for i in range(1, 11))
).encode("utf-16-le")

# 答案列单一取值:样例与全量所有行都是同一类别(某一时期只产生单一类别工单,
# 或标注者把所有行标成了同一类)——样例取满 10 行,让不均衡检出有机会说话。
_CONSTANT_TARGET_SAMPLE = (
    "编号,客户描述,类别\n" + "".join(f"{i:03d},问题{i},质量\n" for i in range(1, 11))
).encode()
_CONSTANT_TARGET_FULL = _CONSTANT_TARGET_SAMPLE

# 答案列全是空格:表头保留、每行答案字段是 3 个空格(复制粘贴/导出填充产生的
# 「看起来填了」的列)——空格是否等价于空,以实测为准。
_WHITESPACE_TARGET_SAMPLE = ("编号,客户描述,类别\n001,杯子破损,   \n002,物流未更新,   \n").encode()
_WHITESPACE_TARGET_FULL = (
    "编号,客户描述,类别\n" + "".join(f"{i:03d},问题{i},   \n" for i in range(1, 11))
).encode()

# 单格 3 万字符:一条记录的输入单元格粘贴了精确 30 000 字符的运行日志(含逗号,
# 按 CSV 规范加引号,引号内逗号不是分隔符);其余行是普通短文本。
# 行号嵌入日志开头,保证长格内容唯一(内容全同的长格会按重复输入被检出)。
_LOG_LINE_UNIT = "2026-09-27 10:23:01 INFO 收到客户端请求,开始处理订单,读取配置42项,回源187ms;"


def _long_single_cell(row_tag: str) -> str:
    head = f"{row_tag} 2026-09-27 10:23:01 收到客户端请求,开始处理订单;"
    return (head + _LOG_LINE_UNIT * 700)[:30_000]


def _long_single_cell_text(full: bool) -> str:
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer, lineterminator="\n")
    writer.writerow(("编号", "客户描述", "类别"))
    for i in range(1, 11):
        if not full and i == 2:  # 样例第 2 行带 3 万字符长格
            writer.writerow(("002", _long_single_cell("工单002"), "物流"))
        elif full and i == 5:  # 全量第 5 行带同形态长格(行号嵌入,内容不同)
            writer.writerow(("005", _long_single_cell("工单005"), "物流"))
        else:
            writer.writerow((f"{i:03d}", f"问题{i}", "质量" if i % 2 else "物流"))
    return buffer.getvalue()


_LONG_SINGLE_CELL_SAMPLE = "".join(
    _long_single_cell_text(full=False).splitlines(keepends=True)[:3]
).encode()
_LONG_SINGLE_CELL_FULL = _long_single_cell_text(full=True).encode()


def _long_line_text() -> str:
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer, lineterminator="\n")
    writer.writerow(("编号", "问题描述", "类别"))
    writer.writerow(("001", _LONG_CELL_PREFIX * 520, "质量"))
    writer.writerow(("002", "物流未更新", "物流"))
    for i in range(3, 11):
        writer.writerow(
            (f"{i:03d}", _LONG_CELL_PREFIX * (300 + i * 20), "质量" if i % 2 else "物流")
        )
    return buffer.getvalue()


# 样例只取表头 + 前两条(单元格内无换行,按行切分安全);全量含全部十条超长行。
_LONG_LINE_LINES = _long_line_text().splitlines(keepends=True)
_LONG_LINE_SAMPLE = "".join(_LONG_LINE_LINES[:3]).encode()
_LONG_LINE_FULL = "".join(_LONG_LINE_LINES).encode()

# 样例(非全量)中部混入重复表头行:与 _DUP_HEADER_FULL 方向相反的脏数据——
# 已实测定局(场景 30):样例侧不拦、按普通数据行读入,与全量侧硬拦构成不对称边界;
# 两侧同现的组合见场景 40(duplicate-header-rows-in-both,与 22 共用 _DUP_HEADER_FULL)。
_DUP_HEADER_IN_SAMPLE = (
    "编号,客户描述,类别\n001,杯子破损,质量\n编号,客户描述,类别\n002,物流未更新,物流\n"
).encode()

# 引号内含换行的多行单元格:输入与标签都按 CSV 规范加引号且引号内有换行,
# 读取/预览/对比核验/盲标是否原样保留、精确匹配行为如何,以实测为准。
_MULTILINE_CELLS_ROWS = (
    ("001", "杯子破损，\n附照片一张，杯身有明显裂纹。", "质量\n（外观破损）"),
    ("002", "物流未更新", "物流\n（一直未更新）"),
    ("003", "屏幕碎裂，\n无法正常显示。", "质量\n（外观破损）"),
    ("004", "快递丢失", "物流\n（一直未更新）"),
    ("005", "开不了机，\n按电源键无反应，\n充电指示灯也不亮。", "质量\n（外观破损）"),
    ("006", "地址填错", "物流\n（一直未更新）"),
    ("007", "异味", "质量\n（外观破损）"),
    ("008", "延迟送达", "物流\n（一直未更新）"),
    ("009", "无法充电，\n更换充电器与数据线后依旧。", "质量\n（外观破损）"),
    ("010", "包装破损", "物流\n（一直未更新）"),
)


def _multiline_cells_text(row_count: int) -> str:
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer, lineterminator="\n")
    writer.writerow(("编号", "客户描述", "类别"))
    for row in _MULTILINE_CELLS_ROWS[:row_count]:
        writer.writerow(row)
    return buffer.getvalue()


_MULTILINE_CELLS_SAMPLE = _multiline_cells_text(2).encode()
_MULTILINE_CELLS_FULL = _multiline_cells_text(10).encode()

# 目标列大小写变体:Yes/yes/YES 指同一业务取值,入口已有 strip;大小写是否归一、
# 标签变体检出是否覆盖,以实测为准。样例三种写法各一条(配对核验需要不同答案)。
_CASE_VARIANTS_SAMPLE = (
    "编号,客户描述,类别\n001,杯子破损,Yes\n002,物流未更新,yes\n003,屏幕碎裂,YES\n"
).encode()
_CASE_VARIANTS_FULL = (
    "编号,客户描述,类别\n"
    "001,杯子破损,Yes\n002,物流未更新,yes\n003,屏幕碎裂,YES\n004,快递丢失,yes\n"
    "005,开不了机,Yes\n006,地址填错,YES\n007,异味,yes\n008,延迟送达,YES\n"
    "009,无法充电,Yes\n010,包装破损,yes\n"
).encode()


# 目标列含义反转:样例确认「类别」=售后类别(质量/物流),全量却是另一份导出,
# 同名列装的是优先级(高/低)——输入行与样例不重叠,答案值整体是样例未覆盖的新类别。
_REVERSAL_FULL = (
    "编号,客户描述,类别\n"
    "101,地址填错,高\n102,快递丢失,低\n103,屏幕碎裂,高\n104,延迟送达,低\n"
    "105,异味,高\n106,包装破损,低\n107,开不了机,高\n108,无法充电,低\n"
    "109,外壳划痕,高\n110,配件缺失,低\n"
).encode()

# 单格 10 万字符:比 extreme-long-single-cell(精确 3 万字符)更极端的同形态夹具
# ——一条记录的输入单元格粘贴精确 100 000 字符的运行日志(含逗号,按 CSV 规范加引号)。
_100K_LOG_UNIT = "2026-09-27 10:23:01 INFO 收到客户端请求,开始处理订单,读取配置42项,回源187ms;"


def _hundred_k_cell(row_tag: str) -> str:
    head = f"{row_tag} 2026-09-27 10:23:01 收到客户端请求,开始处理订单;"
    return (head + _100K_LOG_UNIT * 2000)[:100_000]


def _hundred_k_cell_text(full: bool) -> str:
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer, lineterminator="\n")
    writer.writerow(("编号", "客户描述", "类别"))
    for i in range(1, 11):
        if not full and i == 2:  # 样例第 2 行带 10 万字符长格
            writer.writerow(("002", _hundred_k_cell("工单002"), "物流"))
        elif full and i == 5:  # 全量第 5 行带同形态长格(行号嵌入,内容不同)
            writer.writerow(("005", _hundred_k_cell("工单005"), "物流"))
        else:
            writer.writerow((f"{i:03d}", f"问题{i}", "质量" if i % 2 else "物流"))
    return buffer.getvalue()


_100K_SINGLE_CELL_SAMPLE = "".join(
    _hundred_k_cell_text(full=False).splitlines(keepends=True)[:3]
).encode()
_100K_SINGLE_CELL_FULL = _hundred_k_cell_text(full=True).encode()


# Excel「CSV UTF-8」导出(Windows Excel 2016+ 的默认 UTF-8 导出):UTF-8 带 BOM(EF BB BF)
# + CRLF 行尾。BOM 若不剥掉,首列名会变成「\ufeff编号」;CRLF 若按裸行切分会残留 \r。
def _bom_crlf(text: str) -> bytes:
    return b"\xef\xbb\xbf" + text.replace("\n", "\r\n").encode("utf-8")


_UTF8_BOM_SAMPLE = _bom_crlf("编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n")
_UTF8_BOM_FULL = _bom_crlf(
    "编号,客户描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
)


# 多 Sheet Excel:xlsx 含两个 sheet——第一个 sheet 是工单数据,第二个是完全不同的
# 员工表(一份工作簿装多个业务表的常见形态)。入口默认只读第一个 sheet
# (pd.read_excel 默认 sheet_name=0,读取行为不变),profile 如实标注读取范围
# 「该文件含 N 个 sheet,仅读取第一个(名称)」(增强提示已上线,曾为静默忽略);
# 数据在第二个 sheet 的反例见 excel-data-on-second-sheet,显式指定 sheet
# (--sheet,名称或 1 起始序号)后走通的正例见 excel-second-sheet-selected。
_XLSX_CACHE: dict[str, bytes] = {}


def _two_sheet_xlsx(
    key: str,
    rows: tuple[tuple[str, str, str], ...],
    *,
    staff_first: bool = False,
) -> bytes:
    """生成含两个 sheet 的工作簿:工单表(业务数据)+员工表(无关表)。

    staff_first=False 时数据在第一个 sheet(excel-multi-sheet 用);
    True 时数据在第二个 sheet、第一个 sheet 是员工表(excel-data-on-second-sheet 用)。
    """
    cached = _XLSX_CACHE.get(key)
    if cached is not None:
        return cached
    from datetime import datetime
    from io import BytesIO

    from openpyxl import Workbook

    staff_rows = (("员工号", "姓名", "部门"), ("E01", "张三", "质检"), ("E02", "李四", "物流"))
    ticket_rows = (("编号", "客户描述", "类别"), *rows)
    names = ("员工表", "工单表") if staff_first else ("工单表", "员工表")
    lead_rows, tail_rows = (staff_rows, ticket_rows) if staff_first else (ticket_rows, staff_rows)
    workbook = Workbook()
    lead = workbook.active
    lead.title = names[0]
    for row in lead_rows:
        lead.append(row)
    tail = workbook.create_sheet(names[1])
    for row in tail_rows:
        tail.append(row)
    stamp = datetime(2026, 9, 27, 8, 0, 0)  # 固定文档时间戳,同夹具字节可复现
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    _XLSX_CACHE[key] = data
    return data


# Excel 合并单元格:目标列(类别)纵向合并——Excel 里「同一类别只写一次然后下拉合并」
# 的常见形态,读取时除左上角锚点外均读为空串。read_source 检测与数据区相交的合并区,
# merged_note 如实点名范围与根因、说明没有自动填充(仅 xlsx;xls 引擎不提供合并范围)。
def _merged_xlsx(
    key: str, rows: tuple[tuple[str, str, str | None], ...], merges: tuple[str, ...]
) -> bytes:
    """生成目标列含合并单元格的单 sheet 工作簿;None 表示合并区非首格(openpyxl 写入即为空)。"""

    cached = _XLSX_CACHE.get(key)
    if cached is not None:
        return cached
    from datetime import datetime
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "客户描述", "类别"))
    for row in rows:
        sheet.append(row)
    for coord in merges:
        sheet.merge_cells(coord)
    stamp = datetime(2026, 9, 27, 8, 0, 0)  # 固定文档时间戳,同夹具字节可复现
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    _XLSX_CACHE[key] = data
    return data


# 干净 10 行(编号 001-010,质量/物流交替):多 sheet 与行序颠倒场景共用同源数据。
_CLEAN_TEN_ROWS = (
    ("001", "杯子破损", "质量"),
    ("002", "物流未更新", "物流"),
    ("003", "屏幕碎裂", "质量"),
    ("004", "快递丢失", "物流"),
    ("005", "开不了机", "质量"),
    ("006", "地址填错", "物流"),
    ("007", "异味", "质量"),
    ("008", "延迟送达", "物流"),
    ("009", "无法充电", "质量"),
    ("010", "包装破损", "物流"),
)


def _rows_as_csv(rows: tuple[tuple[str, str, str], ...], target_column: str = "类别") -> bytes:
    return (
        f"编号,客户描述,{target_column}\n" + "".join(f"{a},{b},{c}\n" for a, b, c in rows)
    ).encode()


# 样例:类别 C2:C4 合并 → 001 质量(锚点)保留,002/003 空读,004 物流另有一格。
_MERGED_SAMPLE = _merged_xlsx(
    "merged-sample",
    (
        ("001", "杯子破损", "质量"),
        ("002", "物流未更新", None),  # 合并区非首格:读取为空串
        ("003", "屏幕碎裂", None),
        ("004", "快递丢失", "物流"),
    ),
    ("C2:C4",),
)
# 全量:C6:C8 合并(openpyxl 合并时丢弃非首格的既有值)→ 005 质量(锚点)保留,
# 006/007 空读,其余 8 行干净——对照探针用它实测全量侧结局。
_MERGED_FULL = _merged_xlsx("merged-full", _CLEAN_TEN_ROWS, ("C6:C8",))


# 目标列数字型连续值:答案列是 1.0/2.5/3.7 这类连续测量值(回归形态,非离散类别),
# 全量含样例未覆盖的新测量值。value_kind 判定与旅程行为以实测为准。
_CONTINUOUS_SAMPLE = _rows_as_csv(
    (("001", "杯子破损", "1.0"), ("002", "物流未更新", "2.5")), target_column="处理时长"
)
_CONTINUOUS_FULL = _rows_as_csv(
    (
        ("001", "杯子破损", "1.0"),
        ("002", "物流未更新", "2.5"),
        ("003", "屏幕碎裂", "3.7"),
        ("004", "快递丢失", "4.2"),
        ("005", "开不了机", "0.8"),
        ("006", "地址填错", "5.1"),
        ("007", "异味", "2.9"),
        ("008", "延迟送达", "3.3"),
        ("009", "无法充电", "1.6"),
        ("010", "包装破损", "4.8"),
    ),
    target_column="处理时长",
)


def builtin_scenarios() -> list[ScenarioSpec]:
    return [
        ScenarioSpec(
            scenario_id="clean-classification",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full=_CLEAN_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            tags=("classification",),
        ),
        ScenarioSpec(
            scenario_id="dirty-missing-label-column",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full="编号,客户描述\n001,杯子破损\n002,物流未更新\n".encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
            expect_note="全量缺答案列必须拦在全量验证",
            tags=("dirty-data",),
        ),
        ScenarioSpec(
            scenario_id="ambiguous-labels-user-mismatch",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full=_CLEAN_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:blind_verification",
            user_answers=_contradictory,
            expect_note="用户盲标与标签矛盾必须拦在盲标关",
            tags=("ambiguous",),
        ),
        ScenarioSpec(
            scenario_id="too-few-groups",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full="编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n".encode(),
            target_column="类别",
            group_columns=("客户描述",),
            expect="blocked_at:materialize",
            expect_note="独立分组不足 3 个不能形成三分区",
            tags=("small-sample",),
        ),
        ScenarioSpec(
            scenario_id="forecast-shaped-goal",
            goal="根据披露文本预测未来20个交易日价格是否上涨",
            sample=("event_id,披露文本,方向\ne1,收入上升,上涨\ne2,收入持平,未上涨\n").encode(),
            sample_name="events.csv",
            full=(
                "event_id,披露文本,方向\n"
                "e1,收入上升,上涨\ne2,收入持平,未上涨\ne3,收入下降,未上涨\n"
                "e4,收入大增,上涨\ne5,收入持平,未上涨\ne6,收入上升,上涨\n"
            ).encode(),
            target_column="方向",
            group_columns=("event_id",),
            expect="passes",
            expect_note=(
                "通过但带泄漏预警(预测型目标在基础分析收到时间分区警告);"
                "零密钥路径尚不能建立时间方案,配置 Agent 后应走时间契约——已知局限,非缺陷"
            ),
            tags=("forecast", "known-limitation"),
        ),
        ScenarioSpec(
            scenario_id="high-cardinality-target",
            goal="根据描述生成对应工单编号",
            sample=("描述,工单编号\n杯子破损,GD-001\n物流未更新,GD-002\n").encode(),
            sample_name="tickets.csv",
            full=(
                "描述,工单编号\n"
                "杯子破损,GD-001\n物流未更新,GD-002\n屏幕碎裂,GD-003\n快递丢失,GD-004\n"
                "开不了机,GD-005\n地址填错,GD-006\n异味,GD-007\n延迟送达,GD-008\n"
            ).encode(),
            target_column="工单编号",
            group_columns=(),
            expect="passes",
            expect_note=(
                "通过且带「可能选错答案列」预警(高基数目标早期警告已上线,M5/task-001);"
                "抽取类任务的高基数是合法的,预警不阻断"
            ),
            tags=("high-cardinality",),
        ),
        ScenarioSpec(
            scenario_id="dirty-label-variants",
            goal="根据客户首次描述判断售后类别",
            # 样例同时含两组变体写法(带句号与不带),规范写法各占多数:
            # 归一草案从样例生成,因此草案能覆盖全量出现的全部四种写法。
            sample=(
                "编号,客户描述,类别\n"
                "001,杯子破损,质量\n002,物流未更新,物流。\n003,屏幕碎裂,质量。\n"
                "004,快递丢失,物流\n005,开不了机,质量\n006,地址填错,物流\n"
            ).encode(),
            sample_name="工单.csv",
            full=(
                "编号,客户描述,类别\n"
                "001,杯子破损,质量\n002,物流未更新,物流。\n003,屏幕碎裂,质量。\n"
                "004,快递丢失,物流\n005,开不了机,质量\n006,地址填错,物流。\n"
                "007,异味,质量\n008,延迟送达,物流\n009,无法充电,质量。\n010,包装破损,物流\n"
            ).encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "通过且带「标签多种写法」预警(变体检出已上线),且归一草案已预置:"
                "map_values 规则自动映射到各组出现次数最多的写法,预览答案只剩规范写法,"
                "全旅程以归一后的标签通过;采用哪种写法仍由用户在预览确认时裁决"
            ),
            tags=("dirty-data", "label-variants"),
        ),
        ScenarioSpec(
            scenario_id="long-text-inputs",
            goal="根据客户投诉详情判断严重程度",
            sample=(
                "编号,投诉详情,严重程度\n001," + "非常冗长的投诉描述。" * 80 + ",高\n"
                "002," + "等待多日仍未送达。" * 80 + ",低\n"
            ).encode(),
            sample_name="投诉.csv",
            full=(
                "编号,投诉详情,严重程度\n"
                + "".join(
                    f"{i:03d}," + "投诉内容细节。" * (60 + i % 40) + f",{'高' if i % 2 else '低'}\n"
                    for i in range(1, 11)
                )
            ).encode(),
            target_column="严重程度",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "数据层通过;长文本的截断风险由训练前检查(真实 tokenizer)负责,不在矩阵覆盖内——边界如实记录"
            ),
            tags=("long-text", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="duplicate-inputs",
            goal="根据客户首次描述判断售后类别",
            sample=(
                "编号,客户描述,类别\n001,杯子破损,质量\n001,杯子破损,质量\n002,物流未更新,物流\n"
            ).encode(),
            sample_name="工单.csv",
            full=(
                "编号,客户描述,类别\n"
                + "".join(
                    f"{'001' if i % 2 else f'{i:03d}'},{'杯子破损' if i % 2 else f'问题{i}'},质量\n"
                    for i in range(1, 13)
                )
            ).encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note="通过且带「输入完全重复」预警(重复检出已上线);重复行答案一致,仅样本量虚高",
            tags=("dirty-data", "duplicates"),
        ),
        ScenarioSpec(
            scenario_id="same-input-conflicting-labels",
            goal="根据客户首次描述判断售后类别",
            sample="编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n".encode(),
            sample_name="工单.csv",
            full=(
                "编号,客户描述,类别\n"
                "001,杯子破损,质量\n001,杯子破损,物流\n002,物流未更新,物流\n"
                "003,屏幕碎裂,质量\n004,快递丢失,物流\n005,开不了机,质量\n"
                "006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
            ).encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
            expect_note="同一输入对应不同答案(矛盾标签)必须被拦:模型无法同时满足两个答案",
            tags=("ambiguous", "conflicting-labels"),
        ),
        ScenarioSpec(
            scenario_id="multi-source-boundary",
            goal="用工单表和审核表联合判断售后类别",
            sample="描述,类别\n杯子破损,质量\n物流未更新,物流\n".encode(),
            sample_name="工单.csv",
            full=(
                "描述,类别\n"
                "杯子破损,质量\n物流未更新,物流\n屏幕碎裂,质量\n快递丢失,物流\n"
                "开不了机,质量\n地址填错,物流\n异味,质量\n延迟送达,物流\n"
                "无法充电,质量\n包装破损,物流\n"
            ).encode(),
            target_column="类别",
            group_columns=(),
            expect="passes",
            expect_note=(
                "边界如实记录:零密钥路径仅分析主资料;多源组合需配置 Agent——"
                "目标提及多表时基础分析不做组合,按已知边界通过"
            ),
            tags=("boundary", "multi-source"),
        ),
        ScenarioSpec(
            scenario_id="severe-class-imbalance",
            goal="根据客户首次描述判断是否需要人工复核",
            sample=("编号,描述,需复核\n001,普通问题,否\n002,投诉升级,是\n").encode(),
            sample_name="工单.csv",
            full=(
                "编号,描述,需复核\n"
                + "".join(
                    f"{i:03d},问题{i},{'是' if i % 10 == 0 else '否'}\n" for i in range(1, 21)
                )
            ).encode(),
            target_column="需复核",
            group_columns=("编号",),
            expect="passes",
            expect_note="通过且带「分布严重不均衡」预警(多数类 90%:全猜多数类即 90% 准确率,准确率会骗人)",
            tags=("imbalance",),
        ),
        ScenarioSpec(
            scenario_id="gbk-encoded-upload",
            goal="根据客户首次描述判断售后类别",
            sample=_GBK_SAMPLE,
            sample_name="工单.csv",
            full=_GBK_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:create 不被拦——utf-8-sig 解码失败后入口自动回退 gb18030"
                "(GBK 的超集)正确解码,全旅程无乱码;入口暂无需手动选择编码,"
                "需在入口选择编码的是 gb18030 也解不开的冷门编码——边界如实记录"
            ),
            tags=("encoding", "gbk", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="wide-table",
            goal="根据客户首次描述判断售后类别",
            sample=_WIDE_SAMPLE,
            sample_name="宽表.csv",
            full=_WIDE_FULL,
            target_column="类别",
            group_columns=("编号",),
            excluded_columns=_WIDE_JUNK_COLUMNS,
            expect="passes",
            expect_note=(
                "31 列宽表:用户在入口排除 28 个无关列,只留编号(分组)+客户描述(输入)+类别(答案);"
                "宽表不淹死基础分析——列角色由用户选择决定,被排除列仅作记录"
            ),
            tags=("wide-table", "column-selection"),
        ),
        ScenarioSpec(
            scenario_id="mixed-type-column",
            goal="根据客户首次描述判断售后类别",
            sample=_MIXED_SAMPLE,
            sample_name="工单.csv",
            full=_MIXED_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:工单号列混杂数字样(10023)与文本样(AB-1002)编号,"
                "当前零密钥路径不做类型区分,一律按文本(strip 后原样)进入输入;"
                "如需数字语义须先在数据侧规整——已知边界,非缺陷"
            ),
            tags=("mixed-type", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="chinese-punctuation-variants",
            goal="根据客户首次描述判断售后类别",
            sample=("编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n").encode(),
            sample_name="工单.csv",
            full=(
                "编号,客户描述,类别\n"
                "001,杯子破损。,质量\n002,物流未更新,物流\n003,屏幕碎裂！,质量\n"
                "004,快递丢失?,物流\n005,开不了机。,质量\n006,地址填错，地址错了,物流\n"
                "007,异味,质量\n008,延迟送达!,物流\n009,无法充电……,质量\n010,包装破损,物流\n"
            ).encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:零密钥路径不做标点归一,全角/半角标点与句尾标点差异"
                "原样进入训练;基础分析的变体检出只针对答案列,输入侧标点差异由"
                "真实预览由用户核对。语义是否受影响由用户判断。"
            ),
            tags=("boundary", "punctuation"),
        ),
        ScenarioSpec(
            scenario_id="utf16-excel-export",
            goal="根据客户首次描述判断售后类别",
            sample=_UTF16_SAMPLE,
            sample_name="工单.csv",
            full=_UTF16_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "Excel「Unicode 文本」导出(UTF-16 带 BOM + Tab 分隔)自动识别:"
                "带 BOM 的 UTF-16 按证据解码(FF FE/FE FF 开头不可能是合法 UTF-8 或 GBK),"
                "Tab 分隔由嗅探器识别;无 BOM 的 UTF-16 仍被明确拒绝——边界如实记录"
            ),
            tags=("encoding", "utf16", "excel"),
        ),
        ScenarioSpec(
            scenario_id="ultra-long-single-line",
            goal="根据客户粘贴的运行日志判断问题类别",
            sample=_LONG_LINE_SAMPLE,
            sample_name="日志工单.csv",
            full=_LONG_LINE_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:单个输入单元格约数万字符(单行数十 KB)不被数据层拦截,"
                "数据层只验证结构与答案保留;是否截断由训练前预检用真实 tokenizer 测量,"
                "语义是否受影响由用户在真实预览核对"
            ),
            tags=("long-text", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="empty-label-rows-in-full",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full=_EMPTY_LABEL_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
            expect_note=(
                "实测结局:样例有标签、全量混入 2 条空答案行(005/008)——既不被静默跳过,"
                "也不带病通过,而是被全量验证硬拦:blocking 问题「全量存在缺少监督答案的记录…」"
                "点名行号;用户须补标签或删行后重新验证"
            ),
            tags=("dirty-data", "empty-labels"),
        ),
        ScenarioSpec(
            scenario_id="spaced-header-names",
            goal="根据客户首次描述判断售后类别",
            sample=_SPACED_SAMPLE,
            sample_name="工单.csv",
            full=_SPACED_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:baseline_analysis",
            expect_note=(
                "实测结局:表头「 编号 / 类别」带前后空格,入口读表不剥空格,"
                "零密钥路径的答案列按精确名匹配——「类别」匹配不上,在基础分析即被拦;"
                "报错把带空格的真实列名原样列出(可用:[' 编号 ','客户描述',' 类别']),"
                "用户能看见差异。边界:若改选确切带空格列名可全程通过,空格随之进入"
                "指令与标签——如实记录,不做自动剥空格"
            ),
            tags=("dirty-data", "header-hygiene"),
        ),
        ScenarioSpec(
            scenario_id="jsonl-long-line",
            goal="根据客户粘贴的运行日志判断问题类别",
            sample=_JSONL_SAMPLE,
            sample_name="日志工单.jsonl",
            full=_JSONL_FULL,
            full_name="full.jsonl",
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:JSONL 单行数十 KB(单条记录的输入字段带数万字符日志)"
                "逐行 json.loads 不被数据层拦截,数据层只验证结构与答案保留;"
                "是否截断由训练前预检用真实 tokenizer 测量,语义是否受影响由用户在真实预览核对"
            ),
            tags=("long-text", "jsonl", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="duplicate-header-rows-in-full",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full=_DUP_HEADER_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
            expect_note=(
                "实测结局:导出拼接产生的重复表头行被全量验证硬拦(blocking 问题"
                "「与表头完全相同」点名行号);曾为已知缺口(静默成为训练样本),已清偿——"
                "用户删除表头行后可重新验证;没有自动删行"
            ),
            tags=("dirty-data", "header-hygiene"),
        ),
        ScenarioSpec(
            scenario_id="full-width-digits",
            goal="根据客户首次描述判断售后类别",
            sample=_FW_SAMPLE,
            sample_name="工单.csv",
            full=_FW_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:编号与描述中的全角数字(０１２３)不做全角/半角归一,"
                "按原样文本进入输入与分组;与标点变体同理,语义是否受影响由用户判断,"
                "需要数字语义时先在数据侧规整——已知边界,非缺陷"
            ),
            tags=("boundary", "full-width"),
        ),
        ScenarioSpec(
            scenario_id="single-row-sample",
            goal="根据客户首次描述判断售后类别",
            sample=_ONE_ROW_SAMPLE,
            sample_name="工单.csv",
            full=_CLEAN_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:contrast_check",
            expect_note=(
                "实测结局:样例只有 1 行数据(2 行含表头)——create 与基础分析都不拦"
                "(分析如实观察「答案列非空取值 1/1 行,共 1 类」),旅程在对比核验被拦:"
                "「对比核验需要至少两条答案不同的已标注行。」即使全量数据正常,"
                "一条样例也不足以让用户完成配对核验;用户须至少提供 2 条答案不同的已标注样例"
            ),
            tags=("small-sample", "negative-scenario"),
        ),
        ScenarioSpec(
            scenario_id="all-empty-target-column",
            goal="根据客户首次描述判断售后类别",
            sample=_ALL_EMPTY_TARGET_SAMPLE,
            sample_name="工单.csv",
            full=_ALL_EMPTY_TARGET_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:contrast_check",
            expect_note=(
                "实测结局:答案列表头存在但所有值为空(样例与全量皆空)——基础分析不拦,"
                "finding 如实观察「答案列非空取值 0/2 行,共 0 类」,预览逐行标 needs_label;"
                "旅程在对比核验被拦:「对比核验需要至少两条答案不同的已标注行。」"
                "不静默跳过、不自动补值;用户须先补标签再重走旅程"
            ),
            tags=("dirty-data", "empty-target"),
        ),
        ScenarioSpec(
            scenario_id="utf16-no-bom-rejected",
            goal="根据客户首次描述判断售后类别",
            sample=_UTF16_NOBOM_SAMPLE,
            sample_name="工单.csv",
            full=_UTF16_NOBOM_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:create",
            expect_note=(
                "实测结局:无 BOM 的 UTF-16 字节流(「Unicode 文本」导出经其他工具转存丢 BOM)"
                "utf-8-sig 解不开,gb18030 兜底也解不开(前导字节后跟交错的 NUL/换行字节),"
                "create 即明确拒绝:「无法解码文件，请明确指定编码；原始数据未修改。」;"
                "恢复路径已实测:入口显式指定编码(如 utf-16-le)即可正确解码并继续旅程;"
                "与带 BOM 的 utf16-excel-export(按 BOM 证据自动识别)构成编码家族的完整边界"
            ),
            tags=("encoding", "utf16", "excel"),
        ),
        ScenarioSpec(
            scenario_id="constant-target",
            goal="根据客户首次描述判断售后类别",
            sample=_CONSTANT_TARGET_SAMPLE,
            sample_name="工单.csv",
            full=_CONSTANT_TARGET_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:contrast_check",
            expect_note=(
                "实测结局:答案列只有一种取值——create 与基础分析都不拦,基础分析如实发出"
                "「分布严重不均衡:「质量」占 100%(10 行),全猜这一类就有 100% 准确率」的"
                "非阻断预警;旅程在对比核验被拦:「对比核验需要至少两条答案不同的已标注行。」"
                "单一类别连配对核验都组不成,模型没有可学习的区分边界;"
                "用户须让数据覆盖至少两个类别,或确认答案列本身选错"
            ),
            tags=("dirty-data", "single-class"),
        ),
        ScenarioSpec(
            scenario_id="whitespace-only-values",
            goal="根据客户首次描述判断售后类别",
            sample=_WHITESPACE_TARGET_SAMPLE,
            sample_name="工单.csv",
            full=_WHITESPACE_TARGET_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:contrast_check",
            expect_note=(
                "实测结局:空格不算真值——预览层按 strip 判空,逐行标 needs_label"
                "(「缺少监督答案，需要补充或确认标签；没有自动生成真值。」),与真空值同路;"
                "旅程在对比核验被拦:「对比核验需要至少两条答案不同的已标注行。」"
                "边界如实记录:基础分析的「非空取值」计数按不等于空串统计,空格被算作非空"
                "(finding 显示「非空取值 2/2 行,共 1 类:空格×2」)——两层判空口径不一致,"
                "但既不静默跳过也不带病通过,拦截关卡与报错同真空值完全一致;用户须补真标签"
            ),
            tags=("dirty-data", "whitespace"),
        ),
        ScenarioSpec(
            scenario_id="extreme-long-single-cell",
            goal="根据客户粘贴的运行日志判断问题类别",
            sample=_LONG_SINGLE_CELL_SAMPLE,
            sample_name="工单.csv",
            full=_LONG_SINGLE_CELL_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:单个输入单元格精确 3 万字符(按 CSV 规范加引号,引号内逗号"
                "不是分隔符),入口读取完整无截断,预览原样进入(零密钥数据层不截断),"
                "全量验证无 blocking,全旅程通过。已知边界:「2000 字符展示截断」只是 "
                "Agent 工具的展示层惯例(原始值保留本地,read_cell_content 按 offset/limit"
                " 分页可读全),不是数据层截断;训练期是否截断由训练前预检用真实 tokenizer"
                " 测量,语义是否受影响由用户在真实预览核对。对照事实:未按 CSV 规范加引号的"
                "长格(逗号裸奔)在 create 即按「列数与表头不一致」拒绝,规范内的完整通过"
            ),
            tags=("long-text", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="duplicate-header-row-in-sample",
            goal="根据客户首次描述判断售后类别",
            sample=_DUP_HEADER_IN_SAMPLE,
            sample_name="工单.csv",
            full=_CLEAN_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:样例(非全量)中部的重复表头行不被拦——入口按普通数据行读入"
                "(3 行数据),该行以「就绪」面目进入预览:输入是列名「客户描述」、答案是"
                "列名「类别」,分布 finding 如实把「类别」计为一类,对比核验也可能拿它"
                "出题(探针实测:配对项与选项都出现列名);与全量侧 "
                "duplicate-header-rows-in-full(验证硬拦、点名行号)构成不对称边界:"
                "样例侧没有对称检查,靠用户在预览逐行核对自行识别。该行不进入物化数据集"
                "(物化只消费全量预览),全量干净时全旅程通过"
            ),
            tags=("dirty-data", "header-hygiene"),
        ),
        ScenarioSpec(
            scenario_id="multiline-quoted-cells",
            goal="根据客户首次描述判断售后类别",
            sample=_MULTILINE_CELLS_SAMPLE,
            sample_name="工单.csv",
            full=_MULTILINE_CELLS_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:CSV 引号内的换行按规范完整读入——多行输入与多行标签原样保留,"
                "行号记录的是记录起始物理行,分隔符嗅探不受换行干扰;预览输入/答案逐字"
                "含换行,对比核验与盲标按精确原文匹配,全旅程通过。两个对照事实均实测:"
                "同样的换行不加引号(裸换行)在 create 即按「列数与表头不一致」诚实拒绝;"
                "标签首尾带换行(如「\\n质量\\n」)时盲标必拦——提交侧按 strip 比对、"
                "数据侧保留原文,用户照抄预览答案也对不上(探针实测,未入夹具);"
                "含内部换行的标签(本场景夹具形态)不受影响"
            ),
            tags=("csv-format", "multiline", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="target-case-variants",
            goal="根据客户首次描述判断售后类别",
            sample=_CASE_VARIANTS_SAMPLE,
            sample_name="工单.csv",
            full=_CASE_VARIANTS_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:Yes/yes/YES 被当成 3 个不同类别原样通过——入口已有 strip,"
                "大小写无归一;标签变体检出只覆盖 strip/末尾标点归一后相同的写法,"
                "大小写差异不在归一范围,因此不发「标签多种写法」预警、不生成 map_values"
                " 草案(与 dirty-label-variants 的句号变体明确分界:那里检出并预置草案,"
                "这里原样通过);分布 finding 如实列出 3 类,全量无新增类别,全旅程通过。"
                "边界如实记录:大小写变体会被模型当不同答案学习,如需归一须用户手工加"
                " map_values 规则或在数据侧统一写法"
            ),
            tags=("dirty-data", "case-variants", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="target-meaning-reversal",
            goal="根据客户首次描述判断售后类别",
            sample=_CLEAN_SAMPLE,
            sample_name="工单.csv",
            full=_REVERSAL_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:blind_verification",
            user_answers=_reversal_semantics_user,
            expect_note=(
                "实测结局:全量与样例输入不重叠,答案值(高/低)整体是样例未覆盖的新类别——"
                "答案语义漂移的既有守卫在 validate_full 触发但不硬拦:new_categories 是 "
                "review 级预警,「类别字段 类别 出现样例未覆盖的 2 种答案,请核对是否属于"
                "目标类别」点名全部行号,new_target_values 如实记录 ['高','低'],全量预览"
                "原样携带反转后的标签,确认全量不被 review 拦;含义反转最终拦在盲标关:"
                "按样例确认的业务语义(售后类别)作答的用户复现不出 高/低,盲标 0/5 "
                "不一致。对照事实(探针实测):若全量还含与样例同输入的行,"
                "sample_answer_disagreement(blocking)会在 validate_full 更早硬拦。"
                "边界:照抄数据标签的全知用户会全程通过——自动检出=review 预警,"
                "语义裁决靠用户在预览与盲标两处人工核对"
            ),
            tags=("dirty-data", "semantic-drift", "negative-scenario"),
        ),
        ScenarioSpec(
            scenario_id="100k-single-cell",
            goal="根据客户粘贴的运行日志判断问题类别",
            sample=_100K_SINGLE_CELL_SAMPLE,
            sample_name="工单.csv",
            full=_100K_SINGLE_CELL_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "边界如实记录:单格精确 10 万字符(extreme-long-single-cell 3 万字符的 3.3 倍)"
                "入口读取完整无截断——读入口按文件长度抬高 csv 字段上限,不依赖 128KB 默认值;"
                "预览原样进入(输入 100 006 字符=单元格+「客户描述: 」前缀),全量验证无 "
                "blocking,全旅程通过,与 3 万字符场景同结论:零密钥数据层不设长度上限,"
                "训练期截断风险由训练前预检用真实 tokenizer 测量"
            ),
            tags=("long-text", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="excel-utf8-bom-csv",
            goal="根据客户首次描述判断售后类别",
            sample=_UTF8_BOM_SAMPLE,
            sample_name="工单.csv",
            full=_UTF8_BOM_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:Excel「CSV UTF-8」导出(UTF-8 带 BOM + CRLF 行尾,Windows Excel "
                "2016+ 的默认 UTF-8 导出形态)全程无碍——utf-8-sig 解码剥掉 BOM,CRLF 由 "
                "csv 规范消化,列名与单元格值都不残留 \\ufeff/\\r,全旅程通过。对照事实"
                "(探针实测):同样的字节用纯 utf-8 解码,首列名是「\\ufeff编号」——若入口"
                "不按 utf-8-sig 兜底,用户按业务口径选「编号」作分组列就会像 "
                "spaced-header-names 一样在基础分析被拦;与 gbk-encoded-upload(编码回退)、"
                "utf16-excel-export(BOM 证据识别)、utf16-no-bom-rejected(明确拒绝)"
                "共同构成入口编码家族边界"
            ),
            tags=("encoding", "utf8", "excel", "bom"),
        ),
        ScenarioSpec(
            scenario_id="excel-multi-sheet",
            goal="根据客户首次描述判断售后类别",
            sample=_two_sheet_xlsx("sample", _CLEAN_TEN_ROWS[:2]),
            sample_name="工单.xlsx",
            full=_two_sheet_xlsx("full", _CLEAN_TEN_ROWS),
            full_name="full.xlsx",
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:xlsx 含两个 sheet(工单表+员工表,表结构完全不同)——入口按 "
                "pd.read_excel 默认 sheet_name=0 只读第一个 sheet,列名/数据/全旅程全部只"
                "来自工单表,八关全过。多 Sheet 读取范围提示已上线:profile 如实标注"
                "「该文件含 2 个 sheet,仅读取第一个『工单表』,其余 1 个(员工表)未读取」"
                "——此前第二个 sheet 被静默忽略(不报错、不提示),现读取即可见。边界如实记录:"
                "不带 --sheet 时默认仍只读第一个 sheet,sheet 选择已上线(服务层/CLI --sheet);"
                "数据在第二个 sheet 的反例与指定后的走通实测见 excel-data-on-second-sheet、"
                "excel-second-sheet-selected"
            ),
            tags=("excel", "multi-sheet", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="numeric-continuous-target",
            goal="根据客户描述预估处理时长(小时)",
            sample=_CONTINUOUS_SAMPLE,
            sample_name="工单.csv",
            full=_CONTINUOUS_FULL,
            target_column="处理时长",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:答案列是 1.0/2.5/3.7 这类连续测量值(回归形态,非离散类别)。"
                "value_kind 判成 numeric_continuous(全部非空取值均为带小数的可解析数值)"
                "——回归式目标不再被静默当作分类:分析即给出诚实 finding(逐字字符串学习,"
                "非数值回归,需数值容差请先离散化),对照/验收按开放任务口径(无自动严格"
                "评分),全量 8 个样例未覆盖的新测量值触发 numeric_new_values review 预警"
                "(连续目标新值是常态,不阻断),八关全过。剩余边界如实记录:训练与预览"
                "仍按逐字字符串匹配(「1.0」≠「1.00」),整数测量值/编码(无小数点)仍走"
                "类别路径;需要真正数值误差评分时须在数据侧离散化答案列"
            ),
            tags=("numeric", "regression", "boundary-note"),
        ),
        ScenarioSpec(
            scenario_id="row-order-reversed-full",
            goal="根据客户首次描述判断售后类别",
            sample=_rows_as_csv(_CLEAN_TEN_ROWS[:2]),
            sample_name="工单.csv",
            full=_rows_as_csv(tuple(reversed(_CLEAN_TEN_ROWS))),
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:全量与样例同源但行序完全颠倒——内容级守卫全部按内容而非行号工作:"
                "sample_answer_disagreement 按输入文本比对、new_categories 按答案值比对、"
                "冲突/重复检出按 canonical 内容,行序不影响任何判断,全量验证无 blocking,"
                "八关全过。血缘/版本(探针实测重传路径):行 ID 按物理行序重新编号"
                "(r000001 从 001 变成 010)——行身份是位置的,不是内容的;同一会话重传后旧"
                "盲标核验按 full_source_digest 失效(stale 如实标记,不静默复用),full/ 目录"
                "按内容寻址同时保留正序与倒序两份原始字节,重新确认+重新盲标(对新绑定从"
                "第 1 轮开始)后照常 verified,物化产生新版本并绑定重传文件的摘要"
            ),
            tags=("lineage", "row-order", "full-data"),
        ),
        ScenarioSpec(
            scenario_id="excel-data-on-second-sheet",
            goal="根据客户首次描述判断售后类别",
            sample=_two_sheet_xlsx("staff-first-sample", _CLEAN_TEN_ROWS[:2], staff_first=True),
            sample_name="工单.xlsx",
            full=_two_sheet_xlsx("staff-first-full", _CLEAN_TEN_ROWS, staff_first=True),
            full_name="full.xlsx",
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:baseline_analysis",
            expect_note=(
                "实测结局:数据在第二个 sheet(第一个 sheet 是完全不同的员工表)——入口仍按 "
                "sheet_name=0 把员工表当数据读入(列名 员工号/姓名/部门,2 行员工记录),"
                "create 不拦;旅程在基础分析被拦:「答案列『类别』不在数据字段中(可用:"
                "['员工号', '姓名', '部门'])」——被读入的列名原样列出,用户能看出读到的不是"
                "工单数据。多 Sheet 读取范围提示已上线(profile sheet_note):会话在 create 即"
                "如实标注「该文件含 2 个 sheet,仅读取第一个『员工表』,其余 1 个(工单表)"
                "未读取」,数据放错 sheet 不再无声。边界如实记录:不带 --sheet 时仍只读第一个 "
                "sheet;sheet 选择已上线(服务层/CLI create 与 full-validate 的 --sheet,数据在"
                "第二个 sheet 指定后旅程可走通,见 excel-second-sheet-selected),页面入口待接入"
            ),
            tags=("excel", "multi-sheet", "negative-scenario"),
        ),
        ScenarioSpec(
            scenario_id="duplicate-header-rows-in-both",
            goal="根据客户首次描述判断售后类别",
            sample=_DUP_HEADER_IN_SAMPLE,
            sample_name="工单.csv",
            full=_DUP_HEADER_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
            expect_note=(
                "实测结局:重复表头行两侧同现(导出拼接在样例与全量各留了一条表头行)——"
                "样例侧行为与 duplicate-header-row-in-sample(场景 30)一致:中部表头行按"
                "普通数据行读入、不拦,以「就绪」面目进预览并参与对比核验,样例确认照常通过;"
                "全量侧行为与 duplicate-header-rows-in-full(场景 22)一致:validate_full "
                "硬拦「1 条记录与表头完全相同(通常是导出拼接产生的重复表头行),会变成无意义"
                "的训练样本…」,组合场景的结局由全量侧决定。任务初衷「样例中部混入重复表头"
                "是否与全量侧对称」已由场景 30 测定:不对称——样例侧至今没有对称检查,靠用户"
                "预览逐行核对;40 号覆盖此前未测的两侧同现组合,钉住「全量硬拦兜底,两侧同时"
                "脏也不会带病物化」"
            ),
            tags=("dirty-data", "header-hygiene"),
        ),
        ScenarioSpec(
            scenario_id="excel-second-sheet-selected",
            goal="根据客户首次描述判断售后类别",
            sample=_two_sheet_xlsx("staff-first-sample", _CLEAN_TEN_ROWS[:2], staff_first=True),
            sample_name="工单.xlsx",
            full=_two_sheet_xlsx("staff-first-full", _CLEAN_TEN_ROWS, staff_first=True),
            full_name="full.xlsx",
            sample_sheet=2,
            full_sheet=2,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
            expect_note=(
                "实测结局:同一份「数据在第二个 sheet」的工作簿(与 excel-data-on-second-sheet "
                "同源字节,摘要相同),创建入口 --sheet 2 指定后旅程走通——入口按序号读取第二个 "
                "sheet「工单表」(1 表示第一个),列名/数据/预览全部来自工单表,基础分析、对比核验、"
                "样例确认照常;全量验证同样 --sheet 2 按指定读取,10 行工单记录无 blocking/"
                "review,八关全过,盲标 5/5 一致,物化 train 8/validation 1/test 1。读取范围如实"
                "标注:样例与全量 profile sheet_note 均为「该文件含 2 个 sheet,按指定读取『工单表』"
                ";其余 1 个(员工表)未读取」;同一份字节不带 --sheet 时标注仍是「仅读取第一个"
                "『员工表』」——标注随来源对象生成并持久化,同摘要不同选择不串味。边界如实记录:"
                "sheet 选择当前在服务层与 CLI(create/full-validate 的 --sheet,名称或 1 起始序号,"
                "不指定默认读第一个),页面入口待接入"
            ),
            tags=("excel", "multi-sheet", "sheet-selection"),
        ),
        ScenarioSpec(
            scenario_id="merged-cells-in-target-column",
            goal="根据客户首次描述判断售后类别",
            sample=_MERGED_SAMPLE,
            sample_name="工单.xlsx",
            full=_MERGED_FULL,
            full_name="full.xlsx",
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:confirm_sample",
            expect_note=(
                "实测结局:目标列纵向合并(C2:C4)——create 不拦,合并区除左上角外读为"
                "空串(001 质量/004 物流保留,002/003 空读,不自动填充),merged_note "
                "已上线且 create 即如实点名:「该 sheet 含 1 处合并单元格(类别 C2:C4):"
                "合并区除左上角外均读为空值…请取消合并并逐行填写受影响的值;没有自动填充」"
                "——用户在「缺少监督答案」被拦前就能看到根因是 Excel 合并。对比核验不拦"
                "(两条已标注行 001/004 答案不同足以配对);旅程在样例确认被拦:"
                "「当前方案仍有业务问题、缺标签或转换问题,不能确认数据就绪。」——样例含"
                " needs_label 行不能确认(通用缺标签门,矩阵首个 confirm_sample 关场景;"
                "对照探针实测:同样空答案但不合并的文件拦在同一关同一条报错——拦截本身"
                "与合并无关,合并的独有价值是 merged_note 点名根因)。全量侧合并(C6:C8)"
                "对照实测:样例干净时旅程走到全量验证被硬拦「全量存在缺少监督答案的记录,"
                "需要补充标签;没有自动生成真值。(2 条)」(与场景 19 同守卫,根因不同);"
                "本场景取样例侧合并形态定局,两侧同现时样例侧更早拦截。边界如实记录:"
                "取消合并并逐行填写由用户决定,没有自动填充;xls 引擎不提供合并范围,不检测"
            ),
            tags=("excel", "merged-cells", "negative-scenario"),
        ),
    ]


def _contradictory(session):
    answers = {row.row_id: row.target for row in session.full_data.preview.rows}
    first = sorted(answers)[0]
    answers[first] = "和所有标签都不同的答案"
    return answers


def _reversal_semantics_user(session):
    """按样例确认的业务语义(售后类别)作答的用户:目标列含义反转后复现不出 高/低。

    样例教会用户的是「类别=质量/物流」;全量同名列实际是优先级(高/低)。
    忠实的用户按输入文本给出售后类别答案,与反转后的数据标签必然不一致。
    """
    category_by_text = {
        "地址填错": "物流",
        "快递丢失": "物流",
        "屏幕碎裂": "质量",
        "延迟送达": "物流",
        "异味": "质量",
        "包装破损": "物流",
        "开不了机": "质量",
        "无法充电": "质量",
        "外壳划痕": "质量",
        "配件缺失": "物流",
    }
    return {
        row.row_id: category_by_text[row.original["客户描述"]]
        for row in session.full_data.preview.rows
    }
