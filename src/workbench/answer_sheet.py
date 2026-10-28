"""待补答案清单:把「由了解业务的人填写」变成可交接的文件(介入点 10 编码)。

缺答案行此前只是页面上的一段提示——填写人拿不到只含待补行的清单,
预览编号也无法直接定位原文件。这里把行筛选、填写表与交接规则收敛为
单一来源:页面下载按钮与 CLI answer-sheet 命令渲染同一份行与规则,
needs_labels 状态判定也复用同一个行筛选。
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from zipfile import ZipFile

from src.workbench.intake_models import IntakeSession, PreviewRow, SampleSource
from src.workbench.sources import source_view_digest
from src.workbench.temporal_split import TemporalSplitPolicy


def missing_answer_rows(
    rows: list[PreviewRow], policy: TemporalSplitPolicy | None = None
) -> list[PreviewRow]:
    """筛选当前需要人工补答案的行(单一来源)。

    needs_label 且不是时间方案的未成熟标签——未成熟行是「标签窗口未结束」
    的既定保留,不是漏填,不应混进交给填写人的清单。next_action 的
    needs_labels 判定、页面提示与 CLI 清单共用这一个筛选。
    """
    # 调用时再导入:测试按 temporal_split 模块属性替换 is_pending_label 时,
    # 这里的引用随调用解析(顶层 from-import 会把旧引用绑进本模块)。
    from src.workbench.temporal_split import is_pending_label

    return [
        row for row in rows if row.status == "needs_label" and not is_pending_label(row, policy)
    ]


def answer_sheet_source(session: IntakeSession, *, full: bool = False) -> SampleSource | None:
    """Return a direct source only when preview IDs still identify its rows."""
    view = session.full_data if full else session
    if view is None or view.composition_report is not None:
        return None
    if view.sources:
        return view.sources.get("main")
    # Legacy adapted snapshots may not retain their original source metadata.
    return None if view.adapter_report is not None else view.source


def answer_sheet_lineage(
    session: IntakeSession, *, full: bool = False, rows: list[PreviewRow] | None = None
) -> dict[str, str]:
    """Resolve recorded composition origins without guessing ambiguous source views."""
    view = session.full_data if full else session
    if view is None or view.composition_report is None or view.preview is None:
        return {}
    if view.preview.source_digest != view.source.digest:
        return {}
    origins = view.composition_report.get("origins")
    if not isinstance(origins, dict):
        return {}
    lookup: dict[tuple[str, str], list[tuple[str, SampleSource, str, int]]] = {}
    for alias, source in view.sources.items():
        digest = source_view_digest(source)
        for source_row in source.rows:
            lookup.setdefault((source.digest, source_row.row_id), []).append(
                (alias, source, digest, source_row.line)
            )
    processed = {row.row_id: row for row in view.source.rows}
    if len(processed) != len(view.source.rows):
        return {}
    locations = {}
    for row in view.preview.rows if rows is None else rows:
        current = processed.get(row.row_id)
        refs = origins.get(row.row_id)
        if (
            current is None
            or current.values != row.original
            or not isinstance(refs, list)
            or not refs
        ):
            continue
        resolved = []
        for ref in refs:
            if not isinstance(ref, dict):
                break
            origin_digest, origin_row = ref.get("source_digest"), ref.get("row_id")
            if not isinstance(origin_digest, str) or not isinstance(origin_row, str):
                break
            candidates = lookup.get((origin_digest, origin_row), [])
            if "source_alias" in ref:
                candidates = [item for item in candidates if item[0] == ref["source_alias"]]
            if "source_view_digest" in ref:
                candidates = [item for item in candidates if item[2] == ref["source_view_digest"]]
            if len(candidates) != 1:
                break
            alias, source, _, line = candidates[0]
            sheet = f" · 工作表 {source.sheet}" if source.sheet else ""
            resolved.append(f"{alias}：文件 {source.name}{sheet} · 第 {line} 行")
        else:
            locations[row.row_id] = "关联来源（需确认答案字段归属）：" + "；".join(resolved)
    return locations


def _answer_sheet_rows(
    rows: list[PreviewRow],
    field_names: list[str],
    *,
    source: SampleSource | None = None,
    lineage: dict[str, str] | None = None,
) -> Iterator[list[str]]:
    yield [
        "行ID",
        "题目输入（模型将看到的内容）",
        *[f"待填答案（{name}）" for name in field_names],
        "原始记录位置",
    ]
    source_rows = {row.row_id: row for row in source.rows} if source else {}
    if source and len(source_rows) != len(source.rows):
        source_rows = {}
    for row in rows:
        original = source_rows.get(row.row_id)
        location = (lineage or {}).get(row.row_id, "未记录或不可核验，请查看来源血缘")
        if (
            source is not None
            and original is not None
            and all(
                key in row.original and row.original[key] == value
                for key, value in original.values.items()
            )
        ):
            sheet = f" · 工作表 {source.sheet}" if source.sheet else ""
            location = f"文件 {source.name}{sheet} · 第 {original.line} 行（记录起始行）"
        yield [row.row_id, row.input, *[""] * len(field_names), location]


def answer_sheet_csv(
    rows: list[PreviewRow],
    field_names: list[str],
    *,
    source: SampleSource | None = None,
    lineage: dict[str, str] | None = None,
) -> bytes:
    """Export the original values as CSV for programmatic consumption."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(_answer_sheet_rows(rows, field_names, source=source, lineage=lineage))
    return buffer.getvalue().encode("utf-8-sig")


def answer_sheet_xlsx(
    rows: list[PreviewRow],
    field_names: list[str],
    *,
    source: SampleSource | None = None,
    lineage: dict[str, str] | None = None,
) -> bytes:
    """Keep every cell literal, including text that starts with a formula marker."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    if len(rows) + 1 > 1_048_576 or len(field_names) + 3 > 16_384:
        raise ValueError("清单超过 Excel 的行列限制，请使用原始 CSV。")
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "待补答案"
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 56
    for column in range(3, len(field_names) + 3):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.column_dimensions[get_column_letter(len(field_names) + 3)].width = 60
    wrapped = Alignment(vertical="top", wrap_text=True)
    has_carriage_return = False
    for index, values in enumerate(
        _answer_sheet_rows(rows, field_names, source=source, lineage=lineage), 1
    ):
        for value in values:
            has_carriage_return |= "\r" in value
            try:
                units = len(value.encode("utf-16-le")) // 2
            except UnicodeEncodeError:
                raise ValueError("清单包含无法写入 Excel 的字符，请核对原始资料。") from None
            if units > 32_767:
                raise ValueError("清单包含超过 Excel 单元格长度限制的内容，请使用原始 CSV。")
            if any(
                not (
                    char in "\t\n\r"
                    or 0x20 <= ord(char) <= 0xD7FF
                    or 0xE000 <= ord(char) <= 0xFFFD
                    or 0x10000 <= ord(char) <= 0x10FFFF
                )
                for char in value
            ):
                raise ValueError("清单包含无法写入 Excel 的字符，请使用原始 CSV。")
        sheet.append(values)
        for column in range(1, len(values) + 1):
            cell = sheet.cell(index, column)
            cell.data_type = "s"
            if index == 1:
                cell.font = Font(bold=True)
            if column in {2, len(values)}:
                cell.alignment = wrapped
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(field_names) + 3)}{len(rows) + 1}"
    output = io.BytesIO()
    workbook.save(output)
    if has_carriage_return:
        # XML parsers normalize literal CR to LF. Character references retain
        # the exact original text on round-trip through the workbook.
        escaped = io.BytesIO()
        with (
            ZipFile(io.BytesIO(output.getvalue())) as source_zip,
            ZipFile(escaped, "w") as target_zip,
        ):
            for entry in source_zip.infolist():
                content = source_zip.read(entry)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(b"\r", b"&#13;")
                target_zip.writestr(entry, content)
        return escaped.getvalue()
    return output.getvalue()


def answer_sheet_lines(count: int) -> list[str]:
    """交接与填写规则(单一来源):页面提示与 CLI stderr 输出同一份行。

    这三行是介入点 10 的判断编码:怎么交接(只填答案列、手动回填原文件)、
    填写依据(业务事实,不凭猜测——答案直接成为监督信号)、输入不足时
    怎么办(留空并说明,是任务定义问题,不是编答案的理由)。
    """
    return [
        f"共 {count} 条待补：交给填写人后只填「待填答案」列，行ID 与题目输入列保持原样。行ID 是当前预览编号；依据原始记录位置手动回填预览时的原文件，组合来源先核对来源血缘。文件已修改或重排时请重新预览；清单不能直接上传回填。",
        "填写依据是业务事实（查业务系统或档案都可以）；不要凭猜测填——答案会直接成为模型的监督信号，错一条教错一条。",
        "填写人认为题目输入不足以判断答案时，把该行留空并在回传时说明——输入信息不足是任务定义问题（与盲标核验、可学性探针的同类提示同方向），比编一个答案更有价值。",
    ]
