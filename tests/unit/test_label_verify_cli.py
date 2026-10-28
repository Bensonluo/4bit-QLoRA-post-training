"""label-verify / label-verify-submit CLI:人读输出的统计局限说明与判定下界。

CLI 进程内调用（monkeypatch sys.argv，参考 test_probe_cli 的调用模式），
数据用 _full 确认后的全量预览，全部落在 tmp_path，不碰 outputs/。
"""

import csv
import io
import json
import sys

import pytest
from openpyxl import load_workbook

from src.workbench.intake_service import IntakeService
from tests.unit.test_data_materialize import _full


@pytest.fixture()
def store(tmp_path):
    service = IntakeService(tmp_path / "intake")
    session = _full(service)
    return service, session


def _run_verify(monkeypatch, service, session, *extra):
    from scripts import data_intake

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "data_intake.py",
            "--store",
            str(service.root),
            "label-verify",
            session.session_id,
            "--revision",
            str(session.revision),
            *map(str, extra),
        ],
    )
    return data_intake.main()


def _run_submit(monkeypatch, service, session, verification_id, answers: dict):
    from scripts import data_intake

    argv = [
        "data_intake.py",
        "--store",
        str(service.root),
        "label-verify-submit",
        session.session_id,
        "--verification-id",
        verification_id,
    ]
    for row_id, answer in answers.items():
        argv += ["--answer", f"{row_id}={answer}"]
    monkeypatch.setattr(sys, "argv", argv)
    return data_intake.main()


def _answers_for(pending: dict, wrong_row_id: str | None = None) -> dict:
    """构造提交答案：wrong_row_id 之外的行都给正确答案（全量数据的标签都是 yes）。"""
    return {
        row_id: ("故意答错的答案" if row_id == wrong_row_id else "yes")
        for row_id in pending["row_ids"]
    }


def test_cli_label_verify_previews_sample_lower_bound(store, monkeypatch, capsys):
    """抽题时 stderr 如实预告下界：即使全部一致，小样本的真实一致率下界也远低于 100%。"""
    service, session = store
    assert _run_verify(monkeypatch, service, session) == 0
    out, err = capsys.readouterr()
    payload = json.loads(out)
    assert payload["sample_size"] == 5
    assert "即使全部一致" in err, "抽题预告必须说明全对时下界也不到 100%"
    assert "下界" in err, "统计局限说明必须给出可依赖的下界口径"
    attached = service.load(session.session_id).label_verification
    assert attached["verification_id"] == payload["verification_id"]
    assert attached["status"] == "pending"


def test_cli_label_verify_discloses_shortfall_when_labelled_rows_fewer(store, monkeypatch, capsys):
    """已标注行不足所选条数时，stderr 如实说明缩水，不冒充按请求数抽题。"""
    service, session = store
    assert _run_verify(monkeypatch, service, session, "--size", 50) == 0
    out, err = capsys.readouterr()
    payload = json.loads(out)
    assert payload["sample_size"] == 7, "FULL 全量 9 行只有 7 个不同模型输入"
    assert "只有 7" in err and "不足你选择的 50 条" in err
    assert "本轮 7 条即使全部一致" in err, "证据说明按独立题目数量口径"


def test_cli_submit_verdict_line_shows_lower_bound_when_verified(store, monkeypatch, capsys):
    """判定行补上下界（通过态）：5/5 一致的下界远低于 100%，观测一致率不许被当成真实水平。"""
    service, session = store
    assert _run_verify(monkeypatch, service, session) == 0
    pending = json.loads(capsys.readouterr().out)
    capsys.readouterr()

    from src.workbench.intake_service import wilson_lower_bound

    assert (
        _run_submit(
            monkeypatch, service, session, pending["verification_id"], _answers_for(pending)
        )
        == 0
    )
    _, err = capsys.readouterr()
    assert "verified" in err
    assert "5/5 一致" in err
    expected = f"{wilson_lower_bound(5, 5):.0%}"
    assert f"95% 置信下界约 {expected}" in err, "通过判定也必须亮出下界数字"
    assert "下界约 100%" not in err, "小样本全对的下界不许冒充 100%"


def test_cli_submit_verdict_line_shows_lower_bound_when_failed(store, monkeypatch, capsys):
    """判定行补上下界（未通过态）：存在不一致时判定行同样给出下界数字。"""
    service, session = store
    assert _run_verify(monkeypatch, service, session) == 0
    pending = json.loads(capsys.readouterr().out)
    capsys.readouterr()

    from src.workbench.intake_service import wilson_lower_bound

    wrong_row = pending["row_ids"][0]
    assert (
        _run_submit(
            monkeypatch,
            service,
            session,
            pending["verification_id"],
            _answers_for(pending, wrong_row_id=wrong_row),
        )
        == 0
    )
    _, err = capsys.readouterr()
    assert "insufficient_agreement" in err
    assert "4/5 一致" in err
    expected = f"{wilson_lower_bound(4, 5):.0%}"
    assert f"95% 置信下界约 {expected}" in err
    assert "训练不会开始" in err


def test_cli_submit_hint_is_copyable_verbatim(store, monkeypatch, capsys):
    """submit_hint 照抄就能用：提示里的参数与解析器一致，不含会被拒的 --revision。"""
    service, session = store
    assert _run_verify(monkeypatch, service, session) == 0
    payload = json.loads(capsys.readouterr().out)
    hint = payload["submit_hint"]
    assert hint.startswith("data_intake.py label-verify-submit")
    assert "--revision" not in hint, "提交路径不收 --revision，提示不得把用户引向报错"
    assert "--verification-id" in hint and "--answer" in hint


def _open_store(tmp_path):
    """开放生成目标:长文本答案(≥40 字触发逐条审阅模式)。"""
    from src.workbench.baseline_analysis import propose_baseline_analysis

    service = IntakeService(tmp_path / "open-intake")
    long_q = (
        "客户在收货后发现商品存在质量问题并且多次联系客服没有得到满意答复情绪激动要求立即处理。"
    )
    long_a = "根据您的反馈我们已安排专员跟进处理并将持续关注解决进度直到问题闭环。" * 2
    sample = (
        "编号,客户问题,标准答复\n"
        + "".join(f"{i:03d},{long_q}（{i}）,{long_a}（{i}）\n" for i in range(1, 9))
    ).encode()
    full = (
        "编号,客户问题,标准答复\n"
        + "".join(f"{i:03d},{long_q}（{i}）,{long_a}（{i}）\n" for i in range(1, 11))
    ).encode()
    session = service.create("根据客户问题生成标准答复", "qa.csv", sample)
    session = service.apply_analysis(
        session,
        propose_baseline_analysis(session, target_column="标准答复", group_columns=("编号",)),
        model="cli-open",
    )
    session = service.confirm(session.session_id, session.revision)
    session = service.validate_full_data(session.session_id, session.revision, "f.csv", full)
    session = service.confirm_full_data(session.session_id, session.revision)
    return service, session


def test_cli_label_verify_review_mode_shows_answers_and_accepts_yes_no(
    tmp_path, capsys, monkeypatch
):
    """审阅模式 CLI:出示数据答复,是/否裁决可提交,判定行附下界。"""
    service, session = _open_store(tmp_path)

    from scripts import data_intake

    captured_errs = []

    def _run(*extra):
        monkeypatch.setattr(
            sys,
            "argv",
            [
                "data_intake.py",
                "--store",
                str(service.root),
                "label-verify",
                session.session_id,
                "--revision",
                str(session.revision),
                *map(str, extra),
            ],
        )
        capsys.readouterr()
        assert data_intake.main() == 0
        captured = capsys.readouterr()
        captured_errs.append(captured.err)
        return json.loads(captured.out)

    target = tmp_path / "review-questions.csv"
    workbook_path = tmp_path / "review-questions.xlsx"
    pending = _run("--size", "2", "--export-csv", target, "--export-xlsx", workbook_path)
    assert pending["mode"] == "review"
    first_err = captured_errs[-1]
    assert "逐条审阅" in first_err and "是/否" in first_err
    assert "数据中的答复：" in first_err
    assert all("answer" in item for item in pending["items"])
    exported = list(csv.reader(io.StringIO(target.read_text(encoding="utf-8-sig"))))
    assert exported[0] == ["行ID", "题目输入", "数据中的答复（供审阅）", "审阅结论（填写是或否）"]
    assert exported[1:] == [
        [item["row_id"], item["input"], item["answer"], ""] for item in pending["items"]
    ]
    sheet = load_workbook(workbook_path, data_only=False).active
    assert [sheet.cell(2, column).value for column in (1, 2, 3)] == [
        pending["items"][0][key] for key in ("row_id", "input", "answer")
    ]
    assert sheet["C2"].data_type == "s" and sheet["D2"].value is None
    assert "不屏蔽电子表格公式" in first_err

    # 提交:第一条认可,第二条否决 → 审阅未通过
    row_yes = pending["items"][0]["row_id"]
    row_no = pending["items"][1]["row_id"]
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "data_intake.py",
            "--store",
            str(service.root),
            "label-verify-submit",
            session.session_id,
            "--verification-id",
            pending["verification_id"],
            "--answer",
            f"{row_yes}=是",
            "--answer",
            f"{row_no}=否",
        ],
    )
    assert data_intake.main() == 0  # submit 本身成功,verdict 在 JSON 里
    err = capsys.readouterr().err
    assert "insufficient_agreement" in err


def test_cli_label_verify_export_csv_writes_blind_question_list(
    store, tmp_path, monkeypatch, capsys
):
    """--export-csv 抽题后落盘：题目清单含行ID与题目输入，绝不含数据答案。

    与 learnability-probe 的候选导出同款（BOM 表头、父目录自动
    创建）；差别是盲标清单不许泄露标签——导出文件泄露答案，核验就失效。
    """
    service, session = store
    csv_path = tmp_path / "exports" / "questions.csv"
    assert _run_verify(monkeypatch, service, session, "--export-csv", csv_path) == 0
    out, err = capsys.readouterr()
    payload = json.loads(out)
    assert csv_path.exists()
    text = csv_path.read_bytes().decode("utf-8-sig")
    assert text.startswith("行ID")  # 表头：BOM 之后第一行就是列名
    assert "行ID,题目输入（仅根据此列作答）,你的盲标答案（填写后提交）" in text
    inputs = {
        row.row_id: row.input for row in service.load(session.session_id).full_data.preview.rows
    }
    for row_id in payload["row_ids"]:
        assert row_id in text
        assert inputs[row_id] in text, "题目列必须是隐藏答案后的真实输入"
    assert "yes" not in text, "数据标签 yes 不得出现在题目清单里，否则盲标失效"
    listed = [line for line in text.splitlines() if line.strip()]
    assert len(listed) == payload["sample_size"] + 1, "每条抽样题一行，外加表头"
    assert "题目清单已导出" in err


def test_cli_label_verify_xlsx_hides_labels(store, tmp_path, monkeypatch, capsys):
    service, session = store
    workbook_path = tmp_path / "questions.xlsx"
    assert _run_verify(monkeypatch, service, session, "--export-xlsx", workbook_path) == 0
    payload = json.loads(capsys.readouterr().out)
    sheet = load_workbook(workbook_path, data_only=False).active
    assert sheet.max_row == payload["sample_size"] + 1
    for row in list(sheet.values)[1:]:
        assert row[0] in payload["row_ids"]
        assert row[2] is None
        assert "yes" not in str(row)


def test_label_verify_rejects_shared_export_path_before_sampling(
    store, tmp_path, monkeypatch, capsys
):
    service, session = store
    path = tmp_path / "same"
    assert (
        _run_verify(monkeypatch, service, session, "--export-csv", path, "--export-xlsx", path) == 2
    )
    assert "导出路径不能相同" in capsys.readouterr().err
    assert service.load(session.session_id).label_verification is None
    assert not path.exists()


def test_failed_label_workbook_keeps_pending_questions_for_retry(
    store, tmp_path, monkeypatch, capsys
):
    from src.workbench import label_sheet

    service, session = store
    path, raw = tmp_path / "questions.xlsx", tmp_path / "questions.csv"

    def fail(*args, **kwargs):
        raise ValueError("超过 Excel 单元格长度限制")

    monkeypatch.setattr(label_sheet, "verification_questions_to_xlsx", fail)
    assert (
        _run_verify(monkeypatch, service, session, "--export-xlsx", path, "--export-csv", raw) == 2
    )
    assert "核验题目已保存" in capsys.readouterr().err
    previous = service.load(session.session_id).label_verification["verification_id"]
    assert not path.exists() and not raw.exists()
    assert _run_verify(monkeypatch, service, session, "--export-csv", raw) == 0
    assert json.loads(capsys.readouterr().out)["verification_id"] == previous
    assert raw.exists()


def test_cli_label_verify_export_csv_not_written_when_sampling_rejected(
    store, tmp_path, monkeypatch, capsys
):
    """抽题被拒（版本过期）时不写文件，也不装作导出成功。"""
    service, session = store
    csv_path = tmp_path / "questions.csv"
    assert (
        _run_verify(
            monkeypatch,
            service,
            session,
            "--revision",
            session.revision + 1,
            "--export-csv",
            csv_path,
        )
        == 2
    )
    _, err = capsys.readouterr()
    assert not csv_path.exists()
    assert "任务已更新" in err


def test_reproduce_sheet_never_exports_answer_even_if_item_contains_one():
    from scripts.data_intake import _verification_questions_to_csv

    items = [{"row_id": "r000001", "input": "题目", "answer": "不应泄露的标签"}]
    data = _verification_questions_to_csv(items, mode="reproduce").decode("utf-8-sig")
    assert "不应泄露的标签" not in data
    assert list(csv.reader(io.StringIO(data)))[1] == ["r000001", "题目", ""]


def test_review_sheet_refuses_missing_answer_or_unknown_mode():
    from scripts.data_intake import _verification_questions_to_csv

    items = [{"row_id": "r000001", "input": "题目"}]
    with pytest.raises(ValueError, match="缺少数据中的答复"):
        _verification_questions_to_csv(items, mode="review")
    with pytest.raises(ValueError, match="不支持的核验方式"):
        _verification_questions_to_csv(items, mode="unknown")


def test_cli_label_verify_submit_failure_prints_triage_lines(store, monkeypatch, capsys):
    """未通过时 stderr 逐行输出三因分辨(记录的 mismatch_triage 键),通过态不输出。"""
    service, session = store
    targets = {row.row_id: row.target for row in session.full_data.preview.rows}

    def _submit_with(wrong_answer):
        pending = service.start_label_verification(session.session_id, session.revision)
        answers = {item["row_id"]: targets[item["row_id"]] for item in pending["items"]}
        first = pending["items"][0]["row_id"]
        answers[first] = wrong_answer
        assert _run_submit(monkeypatch, service, session, pending["verification_id"], answers) == 0
        return capsys.readouterr()

    _, err = _submit_with("明显不同的答案")
    assert "存在不一致" in err
    assert "全部标签里没有出现过" in err
    assert any(line.startswith("对号修正") for line in err.splitlines())
    # 防背题事实随尾行一起到达 CLI 用户
    assert "换一组题" in err

    # 通过的一轮不输出三因分辨(通过态记录没有该键)
    pending = service.start_label_verification(session.session_id, session.revision)
    answers = {item["row_id"]: targets[item["row_id"]] for item in pending["items"]}
    assert _run_submit(monkeypatch, service, session, pending["verification_id"], answers) == 0
    _, ok_err = capsys.readouterr()
    assert "对号修正" not in ok_err
    assert "存在不一致" not in ok_err
