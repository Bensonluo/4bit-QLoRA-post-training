"""场景矩阵:干净分类场景应一路通过;脏数据与歧义标签应被对应关卡诚实拦下。"""

import csv
import io

import pytest

from src.workbench.scenario_matrix import ScenarioSpec, run_matrix, run_scenario
from src.workbench.scenario_specs import builtin_scenarios

CLEAN_SAMPLE = ("编号,客户描述,类别\n001,杯子破损,质量\n002,物流未更新,物流\n").encode()
CLEAN_FULL = (
    "编号,客户描述,类别\n"
    "001,杯子破损,质量\n002,物流未更新,物流\n003,屏幕碎裂,质量\n004,快递丢失,物流\n"
    "005,开不了机,质量\n006,地址填错,物流\n007,异味,质量\n008,延迟送达,物流\n"
    "009,无法充电,质量\n010,包装破损,物流\n"
).encode()


def test_clean_classification_scenario_passes_every_gate(tmp_path):
    spec = ScenarioSpec(
        scenario_id="clean-classification",
        goal="根据客户首次描述判断售后类别",
        sample=CLEAN_SAMPLE,
        sample_name="工单.csv",
        full=CLEAN_FULL,
        target_column="类别",
        group_columns=("编号",),
        expect="passes",
    )
    result = run_scenario(spec, tmp_path)
    assert result.verdict == "as_expected", result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values())


def test_dirty_full_data_is_blocked_at_validation(tmp_path):
    dirty_full = (
        "编号,客户描述\n001,杯子破损\n002,物流未更新\n"  # 缺答案列
    ).encode()
    spec = ScenarioSpec(
        scenario_id="dirty-missing-label",
        goal="根据客户首次描述判断售后类别",
        sample=CLEAN_SAMPLE,
        sample_name="工单.csv",
        full=dirty_full,
        target_column="类别",
        group_columns=("编号",),
        expect="blocked_at:validate_full",
        expect_note="全量缺答案列,必须在全量验证被拦,不能带病进入训练",
    )
    result = run_scenario(spec, tmp_path)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "validate_full"


def test_blind_mismatch_is_blocked_at_verification(tmp_path):
    """用户盲标与数据标签不一致(标签确有歧义):必须在盲标关被拦。"""

    def contradictory_answers(session):
        answers = {row.row_id: row.target for row in session.full_data.preview.rows}
        first = sorted(answers)[0]
        answers[first] = "和所有标签都不同的答案"
        return answers

    spec = ScenarioSpec(
        scenario_id="ambiguous-labels",
        goal="根据客户首次描述判断售后类别",
        sample=CLEAN_SAMPLE,
        sample_name="工单.csv",
        full=CLEAN_FULL,
        target_column="类别",
        group_columns=("编号",),
        expect="blocked_at:blind_verification",
        user_answers=contradictory_answers,
    )
    result = run_scenario(spec, tmp_path)
    assert result.verdict == "as_expected", result.to_dict()
    assert "盲标不一致" in result.blocked_message


def test_matrix_summary_counts_truthfully(tmp_path):
    specs = [
        ScenarioSpec(
            scenario_id="ok",
            goal="g",
            sample=CLEAN_SAMPLE,
            sample_name="s.csv",
            full=CLEAN_FULL,
            target_column="类别",
            group_columns=("编号",),
            expect="passes",
        ),
        ScenarioSpec(
            scenario_id="dirty",
            goal="g",
            sample=CLEAN_SAMPLE,
            sample_name="s.csv",
            full="编号,客户描述\n1,x\n".encode(),
            target_column="类别",
            group_columns=("编号",),
            expect="blocked_at:validate_full",
        ),
    ]
    report = run_matrix(specs, tmp_path)
    assert report["summary"]["total"] == 2
    assert report["summary"]["as_expected"] == 2
    assert "不测模型效果" in report["summary"]["note"]


def test_builtin_new_scenarios_match_expected_verdicts(tmp_path):
    """GBK 编码/宽表/混合类型三个新场景:期望结局以实测为准(全旅程通过)。"""

    wanted = {"gbk-encoded-upload", "wide-table", "mixed-type-column"}
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    missing = wanted - set(specs)
    assert not missing, f"缺少场景:{sorted(missing)}"

    wide = specs["wide-table"]
    assert len(wide.excluded_columns) >= 28, "宽表场景应把大部分列交给用户排除"

    for scenario_id in sorted(wanted):
        result = run_scenario(specs[scenario_id], tmp_path / scenario_id)
        assert result.verdict == "as_expected", result.to_dict()
        assert result.blocked_at is None, result.to_dict()


def test_excel_data_on_second_sheet_blocked_at_baseline_analysis(tmp_path):
    """场景 39:数据在第二个 sheet——入口把第一个 sheet(员工表)当数据读入,基础分析拦截;
    读取范围标注已上线,create 即如实告知「只读了员工表、工单表未读取」。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "excel-data-on-second-sheet" in specs, "缺少场景 excel-data-on-second-sheet"

    spec = specs["excel-data-on-second-sheet"]

    # 夹具真实性:第一个 sheet 是员工表,工单数据在第二个 sheet
    from io import BytesIO

    from openpyxl import load_workbook

    for blob in (spec.sample, spec.full):
        assert load_workbook(BytesIO(blob), read_only=True).sheetnames == ["员工表", "工单表"]

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "baseline_analysis", result.to_dict()
    assert "答案列" in result.blocked_message and "不在数据字段中" in result.blocked_message
    # 报错把被读入的员工表列名原样列出——用户能看出读到的不是工单数据
    assert "员工号" in result.blocked_message
    # create 不拦:入口照旧把第一个 sheet 当数据读(读取行为不变),拦在列选择这一步
    assert result.stages["create"] == "passed"

    # 读取范围标注(探针实测):profile 在 create 即如实呈现「读的是员工表、工单表没读」
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "second-sheet-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert session.source.columns == ["员工号", "姓名", "部门"], session.source.columns
    assert session.profile["sheet_note"] == (
        "该文件含 2 个 sheet，仅读取第一个「员工表」；其余 1 个（工单表）未读取。"
    )


def test_excel_second_sheet_selected_passes_full_journey(tmp_path):
    """场景 41:数据在第二个 sheet + --sheet 指定——入口按指定读取,八关走通,标注如实。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "excel-second-sheet-selected" in specs, "缺少场景 excel-second-sheet-selected"

    spec = specs["excel-second-sheet-selected"]

    # 夹具真实性:与场景 39 同源字节(员工表在前,工单表在第二个),按序号指定 2
    from io import BytesIO

    from openpyxl import load_workbook

    assert spec.sample_sheet == 2 and spec.full_sheet == 2
    for blob in (spec.sample, spec.full):
        assert load_workbook(BytesIO(blob), read_only=True).sheetnames == ["员工表", "工单表"]

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 标注实测:按指定读取「工单表」;同一份字节不带 --sheet(场景 39 路径)时标注仍是
    # 「仅读取第一个『员工表』」——标注随来源对象生成并持久化,同摘要不同选择不串味
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "selected-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample, sheet=spec.sample_sheet)
    assert session.source.sheet == "工单表"
    assert session.source.columns == ["编号", "客户描述", "类别"], session.source.columns
    assert [row.values["编号"] for row in session.source.rows] == ["001", "002"]
    assert session.profile["sheet_note"] == (
        "该文件含 2 个 sheet，按指定读取「工单表」；其余 1 个（员工表）未读取。"
    )
    default_session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert default_session.source.digest == session.source.digest
    assert default_session.source.sheet == "员工表"
    assert default_session.profile["sheet_note"] == (
        "该文件含 2 个 sheet，仅读取第一个「员工表」；其余 1 个（工单表）未读取。"
    )

    # 全量侧同样按指定读取:验证标注与结论(旅程内的 validate_full 已按 sheet=2 走过)
    from src.workbench.baseline_analysis import propose_baseline_analysis

    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    pending = service.start_contrast_check(session.session_id, session.revision)
    targets = {row.row_id: row.target for row in session.preview.rows}
    service.submit_contrast_check(
        session.session_id,
        pending["check_id"],
        {item["row_id"]: targets[item["row_id"]] for item in pending["items"]},
    )
    session = service.confirm(session.session_id, session.revision)
    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, spec.full, sheet=spec.full_sheet
    )
    assert session.full_data.source.sheet == "工单表"
    assert session.full_data.profile["record_count"] == 10
    assert not [issue for issue in session.full_data.issues if issue.severity == "blocking"]
    assert session.full_data.profile["sheet_note"] == (
        "该文件含 2 个 sheet，按指定读取「工单表」；其余 1 个（员工表）未读取。"
    )


def test_duplicate_header_rows_in_both_sides_blocked_at_validate_full(tmp_path):
    """场景 40:重复表头行两侧同现——样例侧不拦(同场景 30 的不对称边界),全量侧硬拦
    (同场景 22),组合结局由全量侧决定:validate_full 拦下,不会带病物化。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "duplicate-header-rows-in-both" in specs, "缺少场景 duplicate-header-rows-in-both"

    spec = specs["duplicate-header-rows-in-both"]

    # 夹具真实性:样例与全量的数据中部各含一条与表头完全相同的行
    header = spec.sample.split(b"\n")[0]
    assert header in spec.sample.split(b"\n")[1:-1], "样例应在数据中部含重复表头行"
    assert header in spec.full.split(b"\n")[1:-1], "全量应在数据中部含重复表头行"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "validate_full", result.to_dict()
    assert "与表头完全相同" in result.blocked_message, result.to_dict()
    # 样例侧四关照常通过:中部表头行按普通数据行读入并参与对比核验——
    # 样例侧没有对称检查的不对称边界与场景 30 实测结论一致
    assert all(
        result.stages[stage] == "passed"
        for stage in ("create", "baseline_analysis", "contrast_check", "confirm_sample")
    ), result.to_dict()


def test_builtin_matrix_all_scenarios_as_expected(tmp_path):
    """内置场景全集跑台:无论多少个,全部必须 as_expected(意外=产品缺陷)。"""
    report = run_matrix(builtin_scenarios(), tmp_path)
    total = report["summary"]["total"]
    assert total >= 41, f"内置场景应随 known-gap 清偿持续增长,当前 {total}"
    assert report["summary"]["as_expected"] == total
    assert report["summary"]["unexpected_pass"] == 0
    assert report["summary"]["unexpected_block"] == 0
    assert report["summary"]["error"] == 0


def test_label_variants_scenario_normalizes_variants_in_preview(tmp_path):
    """场景 dirty-label-variants 端到端:归一草案预置后,预览答案只剩规范写法,全旅程通过。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "dirty-label-variants" in specs, "缺少场景 dirty-label-variants"
    spec = specs["dirty-label-variants"]

    # 夹具真实性:样例同时含两组变体写法(带句号与不带),且规范写法各占多数
    sample_labels = [line.rsplit(",", 1)[1] for line in spec.sample.decode().splitlines()[1:]]
    assert sample_labels.count("质量") > sample_labels.count("质量。"), sample_labels
    assert sample_labels.count("物流") > sample_labels.count("物流。"), sample_labels

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 探针:草案已预置进方案,预览输入/答案与归一规则一致(变体写法不再进入答案)
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "variants-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    assert any("已生成归一规则草案" in finding.message for finding in analysis.findings)
    draft = analysis.recipe.targets[0].transforms[0]
    assert draft.operation == "map_values"
    assert draft.mapping == {
        "质量": "质量",
        "质量。": "质量",
        "物流": "物流",
        "物流。": "物流",
    }
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert session.preview is not None
    assert {row.target for row in session.preview.rows} == {"质量", "物流"}
    assert not any(row.issues for row in session.preview.rows)


def test_utf16_excel_export_and_ultra_long_single_line_scenarios(tmp_path):
    """场景 17/18:Excel UTF-16 导出自动识别;超长单行边界如实记录。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    for scenario_id in ("utf16-excel-export", "ultra-long-single-line"):
        assert scenario_id in specs, f"缺少场景 {scenario_id}"
        result = run_scenario(specs[scenario_id], tmp_path / scenario_id)
        assert result.verdict == "as_expected", result.to_dict()
        assert result.blocked_at is None, result.to_dict()

    # 超长单行场景确实覆盖「单行数十 KB」的量级,而不是普通长文本
    long_full = specs["ultra-long-single-line"].full
    assert max(len(line) for line in long_full.split(b"\n")) >= 40_000


def test_empty_label_rows_in_full_blocked_at_validation(tmp_path):
    """场景 19:样例有标签、全量混入空答案行——必须被全量验证硬拦,不能跳过或带病通过。"""

    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "empty-label-rows-in-full" in specs, "缺少场景 empty-label-rows-in-full"

    spec = specs["empty-label-rows-in-full"]
    # 夹具真实性:全量确实混有空标签行,且不止一条
    empty_label_rows = [line for line in spec.full.split(b"\n") if line.endswith(b",")]
    assert len(empty_label_rows) >= 2, "全量应混入若干空标签行"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "validate_full", result.to_dict()
    assert "缺少监督答案" in result.blocked_message
    # 样例阶段(分析/对比/确认)全部先通过,拦截发生在全量验证这一关
    assert all(
        result.stages[stage] == "passed"
        for stage in ("create", "baseline_analysis", "contrast_check", "confirm_sample")
    )


def test_spaced_header_names_blocked_at_baseline_analysis(tmp_path):
    """场景 20:表头带前后空格——精确名匹配在基础分析即失败,报错如实列出带空格列名。"""

    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "spaced-header-names" in specs, "缺少场景 spaced-header-names"

    spec = specs["spaced-header-names"]
    # 夹具真实性:表头确实带前后空格,而用户按业务口径选「类别」
    header = spec.sample.split(b"\n")[0].decode()
    assert header != header.strip(), "表头应带前后空格"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "baseline_analysis", result.to_dict()
    assert "答案列" in result.blocked_message and "不在数据字段中" in result.blocked_message
    # 报错把带空格的真实列名原样列出,用户能看见差异
    assert " 编号 " in result.blocked_message
    # 入口读表本身不拦空格表头,失败发生在列选择这一步
    assert result.stages["create"] == "passed"


def test_jsonl_long_line_duplicate_header_and_full_width_scenarios(tmp_path):
    """场景 21-23:JSONL 超长行/重复表头(全量验证硬拦)/全角数字——期望结局以实测为准。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    for scenario_id in ("jsonl-long-line", "full-width-digits"):
        assert scenario_id in specs, f"缺少场景 {scenario_id}"
        result = run_scenario(specs[scenario_id], tmp_path / scenario_id)
        assert result.verdict == "as_expected", result.to_dict()
        assert result.blocked_at is None, result.to_dict()

    # 重复表头行必须被全量验证拦下,不能静默成为训练样本(曾经的已知缺口,已清偿)
    dup_spec = specs["duplicate-header-rows-in-full"]
    dup_result = run_scenario(dup_spec, tmp_path / "duplicate-header-rows-in-full")
    assert dup_result.verdict == "as_expected", dup_result.to_dict()
    assert dup_result.blocked_at == "validate_full", dup_result.to_dict()
    assert "与表头完全相同" in dup_result.blocked_message, dup_result.to_dict()
    # 样例阶段(分析/对比/确认)先通过,拦截发生在全量验证这一关
    assert all(
        dup_result.stages[stage] == "passed"
        for stage in ("create", "baseline_analysis", "contrast_check", "confirm_sample")
    )

    # JSONL 场景确实覆盖「单行数十 KB」量级,且走的是 jsonl 入口而非 CSV
    jsonl_spec = specs["jsonl-long-line"]
    assert jsonl_spec.sample_name.endswith(".jsonl")
    assert jsonl_spec.full_name.endswith(".jsonl")
    assert max(len(line) for line in jsonl_spec.full.split(b"\n")) >= 40_000

    # 重复表头场景的夹具确实在数据中部含一条与表头相同的行
    dup_lines = dup_spec.full.split(b"\n")
    assert dup_lines[0] == "编号,客户描述,类别".encode()
    assert dup_lines[0] in dup_lines[1:], "全量应含重复表头行"

    # 全角数字场景的夹具确实含全角数字
    assert "００１".encode() in specs["full-width-digits"].full


def test_single_row_sample_blocked_at_contrast_check(tmp_path):
    """场景 24:样例只有 1 行数据——配对核验需要 2 条不同答案,在哪一关拦以实测为准。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "single-row-sample" in specs, "缺少场景 single-row-sample"

    spec = specs["single-row-sample"]
    # 夹具真实性:样例确实只有表头+1 行数据,且该行本身有标签(拦截纯因样例数量,不因缺标签)
    lines = spec.sample.splitlines()
    assert len(lines) == 2, "样例应为表头 + 1 行数据"
    assert not lines[1].endswith(b","), "唯一的数据行应有非空答案"

    result = run_scenario(spec, tmp_path / "single-row-sample")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "contrast_check", result.to_dict()
    assert "对比核验需要至少两条答案不同的已标注行" in result.blocked_message
    # create 与基础分析都不拦(一条有标签的行足以生成真实预览),拦截发生在配对核验
    assert result.stages["create"] == "passed"
    assert result.stages["baseline_analysis"] == "passed"


def test_all_empty_target_column_blocked_at_contrast_check(tmp_path):
    """场景 25:答案列所有值为空——分析阶段如实观察 0 类答案,旅程在对比核验被拦。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "all-empty-target-column" in specs, "缺少场景 all-empty-target-column"

    spec = specs["all-empty-target-column"]
    # 夹具真实性:样例与全量每一条数据行的答案字段都为空(表头保留「类别」列名)
    data_lines = spec.sample.splitlines()[1:] + spec.full.splitlines()[1:]
    assert data_lines and all(line.endswith(b",") for line in data_lines), "所有数据行答案应为空"
    assert spec.sample.splitlines()[0].endswith(",类别".encode())

    result = run_scenario(spec, tmp_path / "all-empty-target-column")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "contrast_check", result.to_dict()
    assert "对比核验需要至少两条答案不同的已标注行" in result.blocked_message
    # 基础分析阶段既不静默通过也不拦:如实生成「0 类答案」的观察与逐行 needs_label 预览
    assert result.stages["create"] == "passed"
    assert result.stages["baseline_analysis"] == "passed"


def test_utf16_no_bom_rejected_at_create(tmp_path):
    """场景 26:无 BOM 的 UTF-16 在入口即被明确拒绝、原始数据未修改;显式指定编码可恢复。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "utf16-no-bom-rejected" in specs, "缺少场景 utf16-no-bom-rejected"

    spec = specs["utf16-no-bom-rejected"]
    # 夹具真实性:确无 BOM,且 utf-8-sig 与 gb18030 两个自动兜底都解不开——
    # 若夹具漂移成碰巧可解码的形态,场景就不再覆盖「入口明确拒绝」这条路
    assert spec.sample[:2] not in (b"\xff\xfe", b"\xfe\xff")
    for fallback in ("utf-8-sig", "gb18030"):
        try:
            spec.sample.decode(fallback)
        except UnicodeDecodeError:
            continue
        raise AssertionError(f"夹具意外可被 {fallback} 解码,拒绝路径将不复存在")
    # 字节本身是合法 UTF-16LE:显式指定编码是真实可用的恢复路径
    assert spec.sample.decode("utf-16-le").startswith("编号")

    result = run_scenario(spec, tmp_path / "utf16-no-bom-rejected")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "create", result.to_dict()
    assert "无法解码文件" in result.blocked_message
    assert "原始数据未修改" in result.blocked_message
    # 旅程在第一关即停,没有任何后续阶段被记录
    assert set(result.stages) == {"create"}


def test_constant_target_blocked_at_contrast_check(tmp_path):
    """场景 27:答案列单一取值——不均衡预警先行(不拦),旅程在对比核验被拦。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "constant-target" in specs, "缺少场景 constant-target"

    spec = specs["constant-target"]
    # 夹具真实性:样例与全量的每一条数据行都是同一类别「质量」
    for blob in (spec.sample, spec.full):
        data_lines = blob.splitlines()[1:]
        assert data_lines, "样例与全量都应有数据行"
        assert all(line.endswith(",质量".encode()) for line in data_lines), (
            "所有行答案应同为「质量」"
        )

    result = run_scenario(spec, tmp_path / "constant-target")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "contrast_check", result.to_dict()
    assert "对比核验需要至少两条答案不同的已标注行" in result.blocked_message
    # create 与基础分析都不拦;有价值的提示发生在基础分析(不均衡预警),拦截发生在配对核验
    assert result.stages["create"] == "passed"
    assert result.stages["baseline_analysis"] == "passed"

    # 基础分析对单一取值有如实、不阻断的关卡提示:100% 多数类预警
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "constant-target-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    messages = [finding.message for finding in analysis.findings]
    assert any("分布严重不均衡" in message and "100%" in message for message in messages), messages


def test_whitespace_only_values_treated_as_missing(tmp_path):
    """场景 28:答案列全是空格——不等价于真值:预览层判空标 needs_label,旅程在对比核验被拦。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "whitespace-only-values" in specs, "缺少场景 whitespace-only-values"

    spec = specs["whitespace-only-values"]
    # 夹具真实性:表头保留「类别」,每条数据行的答案字段都是 3 个空格
    assert spec.sample.splitlines()[0].endswith(",类别".encode())
    for blob in (spec.sample, spec.full):
        data_lines = blob.splitlines()[1:]
        assert data_lines and all(line.endswith(b",   ") for line in data_lines), (
            "答案字段应全为空格"
        )

    result = run_scenario(spec, tmp_path / "whitespace-only-values")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "contrast_check", result.to_dict()
    assert "对比核验需要至少两条答案不同的已标注行" in result.blocked_message
    # create 与基础分析都不拦;空格不被当成真值放行,与真空值同一关卡同一报错
    assert result.stages["create"] == "passed"
    assert result.stages["baseline_analysis"] == "passed"

    # 预览层事实:空格答案按 strip 判空,逐行标 needs_label、target 为 None,不自动补值
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "whitespace-only-values-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert session.preview.counts["ready"] == 0
    assert session.preview.counts["needs_label"] == len(session.preview.rows)
    for row in session.preview.rows:
        assert row.target is None, row
        assert any("缺少监督答案" in issue for issue in row.issues), row


def test_extreme_long_single_cell_passes_full_journey(tmp_path):
    """场景 29:单格精确 3 万字符——入口与预览完整无截断,全旅程通过;截断风险边界如实钉住。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "extreme-long-single-cell" in specs, "缺少场景 extreme-long-single-cell"

    spec = specs["extreme-long-single-cell"]

    # 夹具真实性:样例与全量各含一条精确 30 000 字符的输入单元格(按 CSV 规范加引号)
    def cell_lengths(blob: bytes) -> list[int]:
        rows = list(csv.reader(io.StringIO(blob.decode())))
        return [len(row[1]) for row in rows[1:]]

    assert 30_000 in cell_lengths(spec.sample), "样例应含一条 3 万字符单元格"
    assert 30_000 in cell_lengths(spec.full), "全量应含一条 3 万字符单元格"
    assert max(cell_lengths(spec.full)) == 30_000, "全量最长单元格就是这 3 万字符"

    result = run_scenario(spec, tmp_path / "extreme-long-single-cell")
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口与预览事实:单元格原样保留(零密钥数据层不截断),预览完整进入
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "extreme-long-single-cell-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    long_cell = max((row.values["客户描述"] for row in session.source.rows), key=len)
    assert len(long_cell) == 30_000, "入口读取不应截断单元格"
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    long_row = max(session.preview.rows, key=lambda row: len(row.input))
    assert len(long_row.input) >= 30_000, "预览应原样携带 3 万字符输入"
    assert long_row.status == "ready", long_row


def test_duplicate_header_row_in_sample_passes_without_sample_side_check(tmp_path):
    """场景 30:样例(非全量)中部混入重复表头行——样例侧不拦、按普通数据行读入,
    与全量侧硬拦(duplicate-header-rows-in-full)构成不对称边界,期望以实测为准。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "duplicate-header-row-in-sample" in specs, "缺少场景 duplicate-header-row-in-sample"

    spec = specs["duplicate-header-row-in-sample"]
    # 夹具真实性:重复表头行混在样例数据中部;全量是干净数据(表头只出现一次)
    lines = spec.sample.split(b"\n")
    header = lines[0]
    assert header in lines[1:-1], "样例应在数据中部含重复表头行"
    assert spec.full.split(b"\n").count(header) == 1, "全量不应含重复表头行"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口与预览事实:重复表头行按普通数据行读入(3 行数据),原样成为「就绪」预览行
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "dup-header-sample-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert len(session.source.rows) == 3, "重复表头行应被当作一条普通数据行"
    assert session.source.rows[1].values == {
        "编号": "编号",
        "客户描述": "客户描述",
        "类别": "类别",
    }
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    distribution = next(
        finding.message for finding in analysis.findings if finding.message.startswith("答案列")
    )
    assert "共 3 类" in distribution and "类别×1" in distribution, distribution
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    junk = session.preview.rows[1]
    assert junk.status == "ready", junk
    assert junk.input == "客户描述: 客户描述", junk
    assert junk.target == "类别", junk

    # 该行只污染样例确认旅程,不进入物化数据集(物化只消费全量预览,全量为干净 10 行)
    session = service.confirm(session.session_id, session.revision)
    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, spec.full
    )
    assert len(session.full_data.preview.rows) == 10
    assert all(row.original.get("编号") != "编号" for row in session.full_data.preview.rows)


def test_multiline_quoted_cells_kept_verbatim_through_journey(tmp_path):
    """场景 31:引号内含换行的多行单元格——读取/预览/核验原样保留,精确匹配以实测为准。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "multiline-quoted-cells" in specs, "缺少场景 multiline-quoted-cells"

    spec = specs["multiline-quoted-cells"]

    # 夹具真实性:按 CSV 规范解析后,输入与标签单元格都真含引号内换行
    def data_rows(blob: bytes) -> list[list[str]]:
        return list(csv.reader(io.StringIO(blob.decode())))[1:]

    assert any("\n" in row[1] for row in data_rows(spec.sample)), "样例输入应含引号内换行"
    assert all("\n" in row[2] for row in data_rows(spec.sample)), "样例标签应含引号内换行"
    assert {row[2] for row in data_rows(spec.full)} == {
        "质量\n（外观破损）",
        "物流\n（一直未更新）",
    }

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口事实:多行单元格原样读入,行号是记录起始物理行,分隔符嗅探不受换行干扰
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "multiline-cells-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert session.source.delimiter == ","
    first, second = session.source.rows
    assert first.values["客户描述"] == "杯子破损，\n附照片一张，杯身有明显裂纹。"
    assert first.values["类别"] == "质量\n（外观破损）"
    assert (first.line, second.line) == (2, 5), (first.line, second.line)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert session.preview.counts["ready"] == 2
    assert session.preview.rows[0].input == "客户描述: 杯子破损，\n附照片一张，杯身有明显裂纹。"
    assert session.preview.rows[0].target == "质量\n（外观破损）"

    # 对照事实(实测):同样的换行不加引号(裸换行),入口按列数不一致诚实拒绝
    bare = "编号,客户描述,类别\n001,杯子破损,\n质量\n002,物流未更新,物流\n".encode()
    with pytest.raises(ValueError, match="第 3 行有 1 列"):
        service.create(spec.goal, spec.sample_name, bare)


def test_target_case_variants_pass_without_case_normalization(tmp_path):
    """场景 32:Yes/yes/YES 大小写变体——strip 已有、大小写无归一,变体检出不覆盖,原样通过。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "target-case-variants" in specs, "缺少场景 target-case-variants"

    spec = specs["target-case-variants"]

    # 夹具真实性:三种大小写写法在样例与全量都出现(同一业务取值的变体)
    def labels(blob: bytes) -> list[str]:
        return [line.rsplit(",", 1)[1] for line in blob.decode().splitlines()[1:]]

    assert set(labels(spec.sample)) == {"Yes", "yes", "YES"}
    assert set(labels(spec.full)) == {"Yes", "yes", "YES"}

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 基础分析事实:大小写差异不算变体——不发「标签多种写法」预警、不生成 map_values 草案
    # (与 dirty-label-variants 的句号变体形成明确分界:那里检出并预置草案,这里原样通过)
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "case-variants-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    messages = [finding.message for finding in analysis.findings]
    distribution = next(message for message in messages if message.startswith("答案列"))
    assert "共 3 类" in distribution, distribution
    assert "Yes×1" in distribution and "yes×1" in distribution and "YES×1" in distribution
    assert not any("多种写法" in message for message in messages), messages
    assert analysis.recipe.targets[0].transforms == [], analysis.recipe.targets[0]
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert {row.target for row in session.preview.rows} == {"Yes", "yes", "YES"}


def test_target_meaning_reversal_warned_then_blocked_at_blind_verification(tmp_path):
    """场景 33:样例与全量目标列含义反转——validate_full 发 review 预警不硬拦,盲标关拦住。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "target-meaning-reversal" in specs, "缺少场景 target-meaning-reversal"

    spec = specs["target-meaning-reversal"]

    # 夹具真实性:全量答案值(高/低)整体不在样例答案值(质量/物流)里,且输入行不重叠
    sample_lines = spec.sample.decode().splitlines()
    full_lines = spec.full.decode().splitlines()
    assert {line.rsplit(",", 1)[1] for line in sample_lines[1:]} == {"质量", "物流"}
    assert {line.rsplit(",", 1)[1] for line in full_lines[1:]} == {"高", "低"}
    sample_inputs = {line.split(",")[1] for line in sample_lines[1:]}
    full_inputs = {line.split(",")[1] for line in full_lines[1:]}
    assert not sample_inputs & full_inputs, "输入不重叠,只考 new_categories 这一道守卫"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at == "blind_verification", result.to_dict()
    assert "盲标不一致" in result.blocked_message
    # validate_full 不拦:语义漂移的既有守卫是 review 级预警,旅程继续走到盲标
    assert result.stages["validate_full"] == "passed", result.to_dict()

    # 预警事实(探针实测):new_categories 触发、点名全部行号、new_target_values 如实记录
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "reversal-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    pending = service.start_contrast_check(session.session_id, session.revision)
    targets = {row.row_id: row.target for row in session.preview.rows}
    service.submit_contrast_check(
        session.session_id,
        pending["check_id"],
        {item["row_id"]: targets[item["row_id"]] for item in pending["items"]},
    )
    session = service.confirm(session.session_id, session.revision)
    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, spec.full
    )
    report = session.full_data
    assert report.status == "review", report.status
    assert not [issue for issue in report.issues if issue.severity == "blocking"], report.issues
    warning = next(issue for issue in report.issues if issue.code == "new_categories")
    assert "样例未覆盖的 2 种答案" in warning.message
    assert "请核对是否属于目标类别" in warning.message
    assert len(warning.row_ids) == 10, warning.row_ids
    assert report.new_target_values == {"类别": ["高", "低"]}
    assert {row.target for row in report.preview.rows} == {"高", "低"}
    # review 级预警不拦全量确认:数据照常进入人工复核,由用户裁决语义
    session = service.confirm_full_data(session.session_id, session.revision)
    assert session.full_data.status == "confirmed"


def test_100k_single_cell_passes_full_journey(tmp_path):
    """场景 34:单格精确 10 万字符——入口与预览完整无截断,全旅程通过(3 万字符的 3.3 倍)。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "100k-single-cell" in specs, "缺少场景 100k-single-cell"

    spec = specs["100k-single-cell"]

    # 夹具真实性:样例与全量各含一条精确 100 000 字符的输入单元格(按 CSV 规范加引号)
    def cell_lengths(blob: bytes) -> list[int]:
        rows = list(csv.reader(io.StringIO(blob.decode())))
        return [len(row[1]) for row in rows[1:]]

    assert 100_000 in cell_lengths(spec.sample), "样例应含一条 10 万字符单元格"
    assert 100_000 in cell_lengths(spec.full), "全量应含一条 10 万字符单元格"
    assert max(cell_lengths(spec.full)) == 100_000, "全量最长单元格就是这 10 万字符"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口与预览事实:单元格原样保留,预览输入=单元格+「客户描述: 」前缀,状态就绪
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "100k-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    long_cell = max((row.values["客户描述"] for row in session.source.rows), key=len)
    assert len(long_cell) == 100_000, "入口读取不应截断单元格"
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    long_row = max(session.preview.rows, key=lambda row: len(row.input))
    assert len(long_row.input) == 100_000 + len("客户描述: ")
    assert long_row.status == "ready", long_row


def test_excel_utf8_bom_csv_passes_full_journey(tmp_path):
    """场景 35:Excel「CSV UTF-8」导出(BOM+CRLF)——BOM 被剥、CRLF 不残留,全旅程通过。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "excel-utf8-bom-csv" in specs, "缺少场景 excel-utf8-bom-csv"

    spec = specs["excel-utf8-bom-csv"]

    # 夹具真实性:样例与全量都带 UTF-8 BOM 且使用 CRLF 行尾(Excel 默认导出形态)
    for blob in (spec.sample, spec.full):
        assert blob.startswith(b"\xef\xbb\xbf"), "应带 UTF-8 BOM"
        assert b"\r\n" in blob, "应使用 CRLF 行尾"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口与预览事实:utf-8-sig 剥 BOM、csv 规范消化 CRLF,列名与值都不残留 \ufeff/\r
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService
    from src.workbench.sources import read_source

    service = IntakeService(tmp_path / "bom-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert session.source.encoding == "utf-8-sig"
    assert session.source.columns == ["编号", "客户描述", "类别"], session.source.columns
    assert session.source.rows[0].values == {
        "编号": "001",
        "客户描述": "杯子破损",
        "类别": "质量",
    }
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert session.preview.rows[0].input == "客户描述: 杯子破损"
    assert session.preview.rows[0].target == "质量"

    # 对照事实(实测):同样的字节用纯 utf-8 解码,首列名带 \ufeff——
    # 若入口不做 utf-8-sig 兜底,业务口径的列名会匹配不上(spaced-header-names 同款拦截)
    plain = read_source("工单.csv", spec.sample, scope="sample", encoding="utf-8")
    assert plain.columns[0] == "\ufeff编号"


def test_excel_multi_sheet_reads_first_sheet_only(tmp_path):
    """场景 36:xlsx 含两个 sheet——入口只读第一个 sheet(读取行为不变),
    profile 如实标注读取范围(多 Sheet 提示已上线,曾为静默忽略)。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "excel-multi-sheet" in specs, "缺少场景 excel-multi-sheet"

    spec = specs["excel-multi-sheet"]

    # 夹具真实性:工作簿确实有两个 sheet,且第二个 sheet 是完全不同的表(员工表)
    from io import BytesIO

    from openpyxl import load_workbook

    for blob in (spec.sample, spec.full):
        assert load_workbook(BytesIO(blob), read_only=True).sheetnames == ["工单表", "员工表"]

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 入口事实:只读第一个 sheet——列名与数据全部来自工单表,员工表的列不可见
    from src.workbench.sources import profile_source, read_source

    source = read_source(spec.sample_name, spec.sample, scope="sample")
    assert source.format == "xlsx"
    assert source.columns == ["编号", "客户描述", "类别"], source.columns
    assert [row.values["客户描述"] for row in source.rows] == ["杯子破损", "物流未更新"]
    # 多 Sheet 读取范围提示已上线:profile 如实标注含 2 个 sheet、只读了第一个
    note = profile_source(source)["sheet_note"]
    assert "2 个 sheet" in note and "仅读取第一个「工单表」" in note, note
    assert "员工表" in note, note
    # 旅程侧同样可见:会话 profile 在 create 即携带该标注
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "multi-sheet-note-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    assert "仅读取第一个「工单表」" in session.profile["sheet_note"]

    # 对照事实(实测):数据在第二个 sheet(第一个是员工表)时,入口把员工表
    # 当数据读入,同样不报错——读取范围标注此时是用户唯一的「读错了 sheet」线索
    from openpyxl import Workbook

    workbook = Workbook()
    staff = workbook.active
    staff.append(("员工号", "姓名", "部门"))
    staff.append(("E01", "张三", "质检"))
    tickets = workbook.create_sheet("工单表")
    tickets.append(("编号", "客户描述", "类别"))
    tickets.append(("001", "杯子破损", "质量"))
    buffer = BytesIO()
    workbook.save(buffer)
    swapped = read_source("混合.xlsx", buffer.getvalue(), scope="sample")
    assert swapped.columns == ["员工号", "姓名", "部门"], swapped.columns


def test_numeric_continuous_target_treated_as_categorical(tmp_path):
    """场景 37:答案列连续数值——numeric_continuous 如实标注,逐字学习边界诚实告知。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "numeric-continuous-target" in specs, "缺少场景 numeric-continuous-target"

    spec = specs["numeric-continuous-target"]

    # 夹具真实性:答案列全为小数点数值(非离散类别),且全量含样例未覆盖的新测量值
    def numbers(blob: bytes) -> list[str]:
        return [line.rsplit(",", 1)[1] for line in blob.decode().splitlines()[1:]]

    assert all("." in value for value in numbers(spec.sample) + numbers(spec.full)), "应为连续数值"
    assert set(numbers(spec.full)) - set(numbers(spec.sample)), "全量应含样例未覆盖的新测量值"

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 判定事实:value_kind=numeric_continuous(不再静默当分类),并给出逐字学习边界的诚实 finding
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "numeric-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="处理时长", group_columns=("编号",))
    assert analysis.recipe.targets[0].value_kind == "numeric_continuous", analysis.recipe.targets[0]
    assert analysis.recipe.targets[0].transforms == []
    distribution = next(f.message for f in analysis.findings if f.message.startswith("答案列"))
    assert "共 2 类" in distribution and "1.0×1" in distribution and "2.5×1" in distribution
    honest = next(f for f in analysis.findings if "numeric_continuous" in f.message)
    assert "不是数值回归" in honest.message and "逐字" in honest.message
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    assert {row.target for row in session.preview.rows} == {"1.0", "2.5"}

    # 全量事实:8 个新测量值触发 new_categories review 预警(不阻断,旅程继续)
    pending = service.start_contrast_check(session.session_id, session.revision)
    targets = {row.row_id: row.target for row in session.preview.rows}
    service.submit_contrast_check(
        session.session_id,
        pending["check_id"],
        {item["row_id"]: targets[item["row_id"]] for item in pending["items"]},
    )
    session = service.confirm(session.session_id, session.revision)
    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, spec.full
    )
    blocking = [issue for issue in session.full_data.issues if issue.severity == "blocking"]
    assert not blocking, [issue.to_dict() for issue in blocking]
    warning = next(
        issue for issue in session.full_data.issues if issue.code == "numeric_new_values"
    )
    assert "8 个样例未覆盖的新测量值" in warning.message
    assert "逐字" in warning.message
    assert session.full_data.new_target_values == {
        "处理时长": ["3.7", "4.2", "0.8", "5.1", "2.9", "3.3", "1.6", "4.8"]
    }


def test_row_order_reversed_full_passes_and_reupload_stays_content_stable(tmp_path):
    """场景 38:全量行序完全颠倒——内容级守卫按内容稳定;重传后血缘/版本/盲标照常。"""
    specs = {spec.scenario_id: spec for spec in builtin_scenarios()}
    assert "row-order-reversed-full" in specs, "缺少场景 row-order-reversed-full"

    spec = specs["row-order-reversed-full"]

    # 夹具真实性:全量与样例同源(编号 001-010),但行序完全颠倒
    def ids(blob: bytes) -> list[str]:
        return [line.split(",")[0] for line in blob.decode().splitlines()[1:]]

    assert ids(spec.full) == [f"{i:03d}" for i in range(10, 0, -1)], "全量应为 010→001 倒序"
    assert ids(spec.sample) == ["001", "002"]

    # 同源:样例每行的答案与全量同编号行完全一致(只是全量把行序颠倒了)
    def label_of(blob: bytes) -> dict[str, str]:
        return {
            line.split(",")[0]: line.rsplit(",", 1)[1] for line in blob.decode().splitlines()[1:]
        }

    full_labels = label_of(spec.full)
    assert all(full_labels[row_id] == label for row_id, label in label_of(spec.sample).items())

    result = run_scenario(spec, tmp_path / spec.scenario_id)
    assert result.verdict == "as_expected", result.to_dict()
    assert result.blocked_at is None, result.to_dict()
    assert all(stage == "passed" for stage in result.stages.values()), result.to_dict()

    # 重传实测(探针路径):同一会话先走正序全量旅程,再用倒序字节重传全量。
    # 正序夹具由场景倒序全量反推(同一批行,只还原行序),保证两份文件内容同源。
    from src.workbench.baseline_analysis import propose_baseline_analysis
    from src.workbench.intake_service import IntakeService
    from src.workbench.sources import canonical

    lines = spec.full.decode().splitlines()
    ordered_full = ("\n".join([lines[0], *reversed(lines[1:])]) + "\n").encode()

    service = IntakeService(tmp_path / "roworder-probe")
    session = service.create(spec.goal, spec.sample_name, spec.sample)
    analysis = propose_baseline_analysis(session, target_column="类别", group_columns=("编号",))
    session = service.apply_analysis(session, analysis, model="scenario-matrix")
    pending = service.start_contrast_check(session.session_id, session.revision)
    targets = {row.row_id: row.target for row in session.preview.rows}
    service.submit_contrast_check(
        session.session_id,
        pending["check_id"],
        {item["row_id"]: targets[item["row_id"]] for item in pending["items"]},
    )
    session = service.confirm(session.session_id, session.revision)

    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, ordered_full
    )
    digest_original = session.full_data.source.digest
    rows_original = {
        canonical({"编号": row.original["编号"], "答案": row.target})
        for row in session.full_data.preview.rows
    }
    session = service.confirm_full_data(session.session_id, session.revision)
    pending = service.start_label_verification(session.session_id, session.revision)
    answers = {
        item["row_id"]: next(
            r.target for r in session.full_data.preview.rows if r.row_id == item["row_id"]
        )
        for item in pending["items"]
    }
    first = service.submit_label_verification(
        session.session_id, pending["verification_id"], answers
    )
    assert first["verdict"] == "verified", first
    session = service.materialize_dataset(
        session.session_id, session.revision, independent_rows_confirmed=False
    )
    version_original = session.dataset.version

    # 倒序重传:全量验证不拦,内容集合不变,行 ID 按物理行序重新编号
    session = service.validate_full_data(
        session.session_id, session.revision, spec.full_name, spec.full
    )
    report = session.full_data
    digest_reversed = report.source.digest
    assert digest_reversed != digest_original, "行序颠倒后文件摘要必须不同"
    assert not [issue for issue in report.issues if issue.severity == "blocking"], report.issues
    assert not [issue for issue in report.issues if issue.code == "sample_answer_disagreement"]
    rows_reversed = {
        canonical({"编号": row.original["编号"], "答案": row.target}) for row in report.preview.rows
    }
    assert rows_reversed == rows_original, "行序颠倒不改变内容身份"
    by_position = {row.row_id: row.original["编号"] for row in report.preview.rows}
    assert by_position["r000001"] == "010" and by_position["r000010"] == "001", by_position
    # 血缘:full/ 目录按内容寻址同时保留正序与倒序两份原始字节
    stored = {
        path.name for path in (tmp_path / "roworder-probe" / session.session_id / "full").iterdir()
    }
    assert f"{digest_original}.csv" in stored and f"{digest_reversed}.csv" in stored, stored

    # 版本:重传后旧盲标核验按摘要失效(stale 如实标记,不静默复用),须重新确认+重新盲标
    reloaded = service.load(session.session_id)
    assert reloaded.label_verification == {
        "stale": True,
        "previous_verdict": "verified",
        "previous_created_at": reloaded.label_verification["previous_created_at"],
    }
    session = service.confirm_full_data(session.session_id, session.revision)
    pending = service.start_label_verification(session.session_id, session.revision)
    answers = {
        item["row_id"]: next(
            r.target for r in session.full_data.preview.rows if r.row_id == item["row_id"]
        )
        for item in pending["items"]
    }
    second = service.submit_label_verification(
        session.session_id, pending["verification_id"], answers
    )
    assert second["verdict"] == "verified", second
    session = service.materialize_dataset(
        session.session_id, session.revision, independent_rows_confirmed=False
    )
    assert session.dataset.version != version_original, "重传后必须物化新版本"
    assert session.dataset.source_digest == digest_reversed, "新版本绑定重传文件摘要"
