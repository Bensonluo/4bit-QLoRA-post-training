"""sources 入口的来源说明:多 Sheet Excel 必须如实标注读取范围,且不改变读取行为。

sheet 选择(按名称或 1 起始序号指定工作表)随多 Sheet 读取一并覆盖:默认读第一个
sheet 的行为不变,显式指定的读取与标注以实测为准;标注随来源对象携带,不共用摘要。
"""

import json
import sys
from io import BytesIO

import pytest
from openpyxl import Workbook

from src.workbench.sources import profile_source, read_source


def _workbook_bytes(*sheets: tuple[str, tuple[str, ...], list[tuple[str, ...]]]) -> bytes:
    """sheets: (名称, 表头行, 数据行) 三元组,按出现顺序写入同一工作簿。"""

    workbook = Workbook()
    for index, (name, header, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_TWO_SHEET_DATA = _workbook_bytes(
    ("工单表", ("编号", "类别"), [("001", "质量")]),
    ("员工表", ("员工号", "部门"), [("E01", "质检")]),
)


def test_single_sheet_excel_profile_has_no_sheet_note():
    """单 sheet 工作簿读取行为不变,profile 不携带 sheet_note;CSV 同样没有。"""
    data = _workbook_bytes(("工单表", ("编号", "类别"), [("001", "质量")]))
    source = read_source("工单.xlsx", data, scope="sample")
    assert source.format == "xlsx"
    assert source.columns == ["编号", "类别"], source.columns
    assert [row.values for row in source.rows] == [{"编号": "001", "类别": "质量"}]
    assert "sheet_note" not in profile_source(source), profile_source(source)

    csv_source = read_source("工单.csv", "编号,类别\n001,质量\n".encode(), scope="sample")
    assert "sheet_note" not in profile_source(csv_source)


def test_multi_sheet_excel_profile_annotates_first_sheet_only():
    """多 sheet 工作簿:读取行为不变(仍只读第一个 sheet),profile 如实标注读取范围。"""
    data = _workbook_bytes(
        ("工单表", ("编号", "类别"), [("001", "质量"), ("002", "物流")]),
        ("员工表", ("员工号", "部门"), [("E01", "质检")]),
    )
    source = read_source("工作簿.xlsx", data, scope="sample")
    # 读取行为不变:列与行全部来自第一个 sheet,第二个 sheet 的列与行不可见
    assert source.columns == ["编号", "类别"], source.columns
    assert [row.values for row in source.rows] == [
        {"编号": "001", "类别": "质量"},
        {"编号": "002", "类别": "物流"},
    ]
    note = profile_source(source)["sheet_note"]
    assert "2 个 sheet" in note, note
    assert "仅读取第一个「工单表」" in note, note
    assert "员工表" in note, note  # 未读取的 sheet 名称如实列出


def test_multi_sheet_note_covers_full_scope_and_caps_long_sheet_lists():
    """全量侧同样标注;其余 sheet 超过 5 个时不逐一罗列,以「等」收尾。"""
    data = _workbook_bytes(
        ("主表", ("编号", "类别"), [("001", "质量")]),
        *[(f"表{i}", ("列",), [("值",)]) for i in range(1, 9)],
    )
    source = read_source("full.xlsx", data, scope="full")
    assert source.scope == "full"
    note = profile_source(source)["sheet_note"]
    assert "9 个 sheet" in note, note
    assert "仅读取第一个「主表」" in note, note
    assert "其余 8 个" in note, note
    assert "表5" in note and "等" in note, note  # 只列前 5 个
    assert "表8" not in note, note


def test_explicit_sheet_by_name_reads_designated_sheet():
    """按名称指定第二个 sheet:列与行全部来自员工表,标注如实写「按指定读取」。"""
    source = read_source("工作簿.xlsx", _TWO_SHEET_DATA, scope="sample", sheet="员工表")
    assert source.sheet == "员工表", source.sheet
    assert source.columns == ["员工号", "部门"], source.columns
    assert [row.values for row in source.rows] == [{"员工号": "E01", "部门": "质检"}]
    note = profile_source(source)["sheet_note"]
    assert "2 个 sheet" in note, note
    assert "按指定读取「员工表」" in note, note
    assert "工单表" in note and "未读取" in note, note  # 未读取的 sheet 名称如实列出


def test_explicit_sheet_by_ordinal_one_based_and_digit_string():
    """序号 1 起始:2=第二个 sheet;CLI 的字符串数字同义;1=第一个(数据同默认,标注写明按指定)。"""
    by_ordinal = read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet=2)
    by_string = read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet="2")
    assert by_ordinal.sheet == by_string.sheet == "员工表"
    assert by_ordinal.columns == by_string.columns == ["员工号", "部门"]
    assert [row.values for row in by_string.rows] == [{"员工号": "E01", "部门": "质检"}]

    first = read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet=1)
    assert first.sheet == "工单表"
    assert first.columns == ["编号", "类别"]
    assert "按指定读取「工单表」" in profile_source(first)["sheet_note"]


def test_explicit_sheet_errors_list_available_sheets():
    """名称不存在/序号越界都明确拒绝,报错列出全部 sheet 名与序号口径,不静默回退。"""
    with pytest.raises(ValueError, match="找不到 sheet「销售表」") as by_name:
        read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet="销售表")
    assert "工单表、员工表" in str(by_name.value)
    assert "序号从 1 开始" in str(by_name.value)
    with pytest.raises(ValueError, match="sheet 序号 3 超出范围"):
        read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet=3)
    with pytest.raises(ValueError, match="sheet 序号 9 超出范围"):
        read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet="9")


def test_sheet_selection_requires_excel():
    """CSV/JSONL 不接受 sheet 选择:明确拒绝而不是静默忽略。"""
    with pytest.raises(ValueError, match="sheet 选择仅对 Excel 文件有效；当前文件是 csv"):
        read_source("工单.csv", "编号\n001\n".encode(), sheet="工单表")
    with pytest.raises(ValueError, match="当前文件是 jsonl"):
        read_source("工单.jsonl", '{"编号": "001"}\n'.encode(), sheet=1)


def test_single_sheet_explicit_selection_reads_it_without_note():
    """单 sheet 工作簿显式指定:正常读取该表;没有未读取的 sheet,profile 无 sheet_note。"""
    data = _workbook_bytes(("工单表", ("编号", "类别"), [("001", "质量")]))
    source = read_source("工单.xlsx", data, sheet="工单表")
    assert source.sheet == "工单表"
    assert source.columns == ["编号", "类别"]
    assert "sheet_note" not in profile_source(source), profile_source(source)


def test_same_workbook_selections_annotated_independently():
    """同一份字节两种指定,标注各自如实——标注随来源对象走,不按文件摘要共用。"""
    first = read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet=1)
    second = read_source("工作簿.xlsx", _TWO_SHEET_DATA, sheet=2)
    assert first.digest == second.digest  # 同一份字节
    assert "按指定读取「工单表」" in profile_source(first)["sheet_note"]
    assert "按指定读取「员工表」" in profile_source(second)["sheet_note"]


def test_service_create_honors_sheet_and_survives_roundtrip(tmp_path):
    """创建入口(服务层)透传 sheet 选择:会话建在指定的 sheet 上,存档回读后标注仍在。"""
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "intake")
    session = service.create(
        "根据客户首次描述判断售后类别", "工作簿.xlsx", _TWO_SHEET_DATA, sheet="员工表"
    )
    assert session.source.sheet == "员工表"
    assert session.source.columns == ["员工号", "部门"], session.source.columns
    assert "按指定读取「员工表」" in session.profile["sheet_note"]
    reloaded = service.load(session.session_id)
    assert reloaded.source.sheet == "员工表"
    assert "按指定读取「员工表」" in reloaded.profile["sheet_note"]


def test_cli_create_sheet_flag_reads_designated_sheet(tmp_path, monkeypatch, capsys):
    """CLI create --sheet 接线冒烟:--sheet 2 建出的会话读到第二个 sheet 员工表数据。

    详细 CLI 行为回归归 test_data_intake 域;这里只钉住参数透传到服务层的链路。
    """
    from scripts import data_intake as cli

    path = tmp_path / "工作簿.xlsx"
    path.write_bytes(_TWO_SHEET_DATA)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "data_intake.py",
            "--store",
            str(tmp_path / "store"),
            "create",
            "--input",
            str(path),
            "--goal",
            "根据客户首次描述判断售后类别",
            "--sheet",
            "2",
        ],
    )
    assert cli.main() == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["source"]["sheet"] == "员工表"
    assert payload["source"]["columns"] == ["员工号", "部门"]
    assert "按指定读取「员工表」" in payload["profile"]["sheet_note"]


def _merged_workbook_bytes(merges: list[str], *, sheet_name: str = "工单表") -> bytes:
    """生成类别列含合并单元格的工作簿:merges 是 openpyxl 坐标列表(如 "C2:C4")。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(("编号", "客户描述", "类别"))
    for number, text, label in (
        ("001", "杯子破损", "质量"),
        ("002", "物流未更新", None),  # 合并区非首格:openpyxl 写入时即为空
        ("003", "屏幕碎裂", None),
        ("004", "快递丢失", "物流"),
    ):
        sheet.append((number, text, label))
    for coord in merges:
        sheet.merge_cells(coord)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_merged_cells_read_as_empty_and_annotated():
    """合并单元格如实点名:非首格读为空串(不自动填充),merged_note 说出根因与修法。"""
    data = _merged_workbook_bytes(["C2:C4"])
    source = read_source("工单.xlsx", data, scope="sample")
    # 读取行为不变:合并区除左上角外均读为空串,绝不猜业务语义去填充
    assert [row.values["类别"] for row in source.rows] == ["质量", "", "", "物流"]
    note = source.merged_note
    assert "1 处合并单元格" in note, note
    assert "类别 C2:C4" in note, note  # 受影响列与范围点名
    assert "除左上角外均读为空值" in note, note
    assert "没有自动填充" in note, note
    assert profile_source(source)["merged_note"] == note  # profile 同步如实呈现


def test_merged_note_skips_out_of_region_and_unread_sheet_merges():
    """数据区之外的合并不涉及本次读取不点名;未读取 sheet 的合并同样不点名;
    干净文件与 CSV 的 profile 形状不变。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "类别"))
    sheet.append(("001", "质量"))
    sheet.append(("002", "物流"))
    sheet.merge_cells("D9:D11")  # 完全在数据区之外(无内容格)
    other = workbook.create_sheet("备注表")
    other.append(("备注",))
    other.append(("正常",))
    other.merge_cells("A2:A3")  # 合并在另一个 sheet
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()

    default = read_source("工作簿.xlsx", data)
    assert default.merged_note == ""
    assert "merged_note" not in profile_source(default)
    second = read_source("工作簿.xlsx", data, sheet="备注表")
    assert "备注 A2:A3" in second.merged_note, second.merged_note  # 只看实际读取的 sheet

    csv_source = read_source("工单.csv", "编号,类别\n001,质量\n".encode())
    assert csv_source.merged_note == ""
    assert "merged_note" not in profile_source(csv_source)


def test_merged_note_lists_first_five_and_caps_long_lists():
    """合并区超过 5 处时不逐一罗列,以「等」收尾——与 sheet_note 同款口径。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "备注"))
    for i in range(1, 13):
        sheet.append((f"{i:03d}", f"备注{i}"))
    for start in range(2, 13, 2):  # B2:B3 … B12:B13 共 6 处
        sheet.merge_cells(f"B{start}:B{start + 1}")
    buffer = BytesIO()
    workbook.save(buffer)
    source = read_source("工作簿.xlsx", buffer.getvalue())
    note = source.merged_note
    assert "6 处合并单元格" in note, note
    assert "B10:B11" in note, note  # 只列前 5 处
    assert "B12:B13" not in note, note
    assert note.rstrip("。").endswith("等）") or "等" in note


def test_service_create_persists_merged_note(tmp_path):
    """创建入口(服务层)透传:会话建在含合并单元格的文件上,存档回读后标注仍在。"""
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "intake")
    session = service.create(
        "根据客户首次描述判断售后类别", "工单.xlsx", _merged_workbook_bytes(["C2:C4"])
    )
    assert "类别 C2:C4" in session.source.merged_note
    assert "merged_note" in session.profile
    reloaded = service.load(session.session_id)
    assert "类别 C2:C4" in reloaded.source.merged_note
    assert "merged_note" in reloaded.profile


def _formula_workbook_bytes(formulas: dict[str, str]) -> bytes:
    """生成类别列含无缓存公式格的工作簿:坐标 → 公式(openpyxl 写公式即无缓存计算结果)。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "客户描述", "类别"))
    for number, text, label in (
        ("001", "杯子破损", "质量"),
        ("002", "物流未更新", "物流"),
        ("003", "屏幕碎裂", "质量"),
        ("004", "快递丢失", "物流"),
    ):
        sheet.append((number, text, label))
    for coord, formula in formulas.items():
        sheet[coord] = formula
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cached_formula_workbook_bytes(cached: dict[str, str]) -> bytes:
    """把无缓存公式格补上缓存计算结果,模拟真实 Excel 打开并保存过的文件形态。

    openpyxl 写出的公式格 XML 是 <c r="C2"><f>…</f><v /></c>(空缓存);直接改 zip 里的
    sheet XML,把空 <v /> 换成 t="str" 与缓存值——与 Excel 保存后的形态一致。找不到
    目标格即断言失败,防止 openpyxl 输出漂移让夹具静默失效。
    """
    import re
    import zipfile

    data = _formula_workbook_bytes(
        {coord: f'=IF(LEN({coord[0]}{coord[1:]})>0,"质量","物流")' for coord in cached}
    )
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries = [(name, archive.read(name)) for name in archive.namelist()]
    patched: list[tuple[str, bytes]] = []
    for name, blob in entries:
        if name == "xl/worksheets/sheet1.xml":
            xml = blob.decode("utf-8")
            for coord, value in cached.items():
                pattern = re.compile(rf'<c r="{coord}">(.*?)<v\s*/></c>')
                xml, count = pattern.subn(
                    rf'<c r="{coord}" t="str">\1<v>{value}</v></c>', xml, count=1
                )
                assert count == 1, f"公式格 {coord} 未找到,openpyxl 输出形态漂移"
            blob = xml.encode("utf-8")
        patched.append((name, blob))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in patched:
            archive.writestr(name, blob)
    return buffer.getvalue()


def test_formula_cells_read_as_empty_and_annotated():
    """公式格无缓存计算结果如实点名:读为空串(不自动计算),formula_note 说出根因与修法。"""
    data = _formula_workbook_bytes(
        {"C2": '=IF(LEN(B2)>0,"物流","质量")', "C3": '=IF(LEN(B3)>0,"质量","物流")'}
    )
    source = read_source("工单.xlsx", data, scope="sample")
    # 读取行为不变:无缓存公式格读为空串,绝不自动计算去猜值
    # (C2/C3 即前两条数据行,因此空值在前的形态是公式夹具与合并夹具的天然差异)
    assert [row.values["类别"] for row in source.rows] == ["", "", "质量", "物流"]
    note = source.formula_note
    assert "2 个没有缓存计算结果的公式单元格" in note, note
    assert "类别 C2" in note and "类别 C3" in note, note  # 受影响列与坐标点名
    assert "这些公式读为空值" in note, note
    assert "没有自动计算" in note, note
    assert profile_source(source)["formula_note"] == note  # profile 同步如实呈现


def test_cached_formula_cells_read_values_without_false_positive():
    """真实 Excel 保存过的公式格带缓存值:按缓存值正常读取,formula_note 不误报。"""
    data = _cached_formula_workbook_bytes({"C2": "物流", "C3": "质量"})
    source = read_source("工单.xlsx", data, scope="sample")
    # 带缓存值的公式格按缓存值读取(C2 缓存=物流、C3 缓存=质量),
    # 不因「存在公式」而列入 formula_note
    assert [row.values["类别"] for row in source.rows] == ["物流", "质量", "质量", "物流"]
    assert source.formula_note == ""
    assert "formula_note" not in profile_source(source), profile_source(source)


def test_formula_note_skips_out_of_region_and_unread_sheet():
    """数据区之外的公式格不涉及本次读取不点名;未读取 sheet 的公式同样不点名;
    干净文件与 CSV 的 profile 形状不变。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "类别"))
    sheet.append(("001", "质量"))
    sheet.append(("002", "物流"))
    sheet["D9"] = "=1+1"  # 完全在数据区之外(无内容格)
    other = workbook.create_sheet("备注表")
    other.append(("备注", "标记"))
    other.append(("正常", "=1+1"))  # 公式在另一个 sheet 的数据区内
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()

    default = read_source("工作簿.xlsx", data)
    assert default.formula_note == ""
    assert "formula_note" not in profile_source(default)
    second = read_source("工作簿.xlsx", data, sheet="备注表")
    assert "标记 B2" in second.formula_note, second.formula_note  # 只看实际读取的 sheet

    csv_source = read_source("工单.csv", "编号,类别\n001,质量\n".encode())
    assert csv_source.formula_note == ""
    assert "formula_note" not in profile_source(csv_source)


def test_formula_note_lists_first_five_and_caps_long_lists():
    """公式格超过 5 个时不逐一罗列,以「等」收尾——与 merged_note/sheet_note 同款口径。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "类别"))
    for i in range(1, 13):
        sheet.append((f"{i:03d}", "质量" if i % 2 else "物流"))
    for row_no in range(2, 8):  # 类别列 B2..B7 共 6 个无缓存公式格
        sheet[f"B{row_no}"] = f'=IF(A{row_no}>0,"质量","物流")'
    buffer = BytesIO()
    workbook.save(buffer)
    source = read_source("工作簿.xlsx", buffer.getvalue())
    note = source.formula_note
    assert "6 个没有缓存计算结果的公式单元格" in note, note
    assert "类别 B6" in note, note  # 只列前 5 个
    assert "类别 B7" not in note, note
    assert "等" in note, note


def test_service_create_persists_formula_note(tmp_path):
    """创建入口(服务层)透传:会话建在含无缓存公式格的文件上,存档回读后标注仍在。"""
    from src.workbench.intake_service import IntakeService

    service = IntakeService(tmp_path / "intake")
    session = service.create(
        "根据客户首次描述判断售后类别",
        "工单.xlsx",
        _formula_workbook_bytes({"C2": '=IF(LEN(B2)>0,"物流","质量")'}),
    )
    assert "类别 C2" in session.source.formula_note
    assert "formula_note" in session.profile
    reloaded = service.load(session.session_id)
    assert "类别 C2" in reloaded.source.formula_note
    assert "formula_note" in reloaded.profile
