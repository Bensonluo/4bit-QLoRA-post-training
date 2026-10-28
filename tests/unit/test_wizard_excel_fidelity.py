"""Wizard Excel imports preserve literal text and reject spreadsheet errors."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from src.data.wizard.importers import import_table
from src.data.wizard.spec import WizardError


def workbook_path(tmp_path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "原始资料"
    for row in rows:
        sheet.append(row)
    path = tmp_path / "source.xlsx"
    workbook.save(path)
    return path


def test_excel_preserves_na_like_strings_and_real_blanks(tmp_path):
    path = workbook_path(
        tmp_path,
        [
            ["标准名", "别名", "编码"],
            ["NA", "NULL", "001"],
            ["None", "N/A", "002"],
            ["有效", "", "003"],
        ],
    )
    table = import_table(path)
    assert table.rows == [
        {"标准名": "NA", "别名": "NULL", "编码": "001"},
        {"标准名": "None", "别名": "N/A", "编码": "002"},
        {"标准名": "有效", "别名": None, "编码": "003"},
    ]


@pytest.mark.parametrize("error", ["#DIV/0!", "#N/A", "#REF!"])
def test_excel_error_cell_is_not_silently_turned_into_empty_value(tmp_path, error):
    path = workbook_path(tmp_path, [["标准名", "别名"], ["甲", error]])
    with pytest.raises(WizardError, match=r"原始资料.*B2") as exc:
        import_table(path)
    assert error in str(exc.value)


def test_excel_literal_error_text_is_preserved(tmp_path):
    path = workbook_path(tmp_path, [["标准名", "别名"], ["甲", "#N/A"]])
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook.active["B2"].data_type = "s"
    workbook.save(path)
    assert import_table(path).rows[0]["别名"] == "#N/A"


def test_error_in_header_is_not_replaced_by_an_unnamed_column(tmp_path):
    path = workbook_path(tmp_path, [["标准名", "#REF!"], ["甲", "别名"]])
    with pytest.raises(WizardError, match=r"原始资料.*B1.*#REF!"):
        import_table(path)


def test_cached_formula_error_is_rejected(tmp_path):
    path = workbook_path(tmp_path, [["标准名", "别名"], ["甲", "=1/0"]])
    output = BytesIO()
    with ZipFile(path) as original, ZipFile(output, "w") as modified:
        for info in original.infolist():
            content = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                old = b'<c r="B2"><f>1/0</f><v></v></c>'
                if old not in content:
                    old = b'<c r="B2"><f>1/0</f><v /></c>'
                assert old in content
                content = content.replace(old, b'<c r="B2" t="e"><f>1/0</f><v>#DIV/0!</v></c>')
            modified.writestr(info, content)
    path.write_bytes(output.getvalue())
    with pytest.raises(WizardError, match=r"原始资料.*B2.*#DIV/0!"):
        import_table(path)


def test_legacy_xls_preserves_text_and_rejects_error_cells(tmp_path):
    pytest.importorskip("xlrd")
    xlwt = pytest.importorskip("xlwt")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("旧资料", cell_overwrite_ok=True)
    for column, value in enumerate(("标准名", "别名", "编码")):
        sheet.write(0, column, value)
    for column, value in enumerate(("NA", "NULL", "001")):
        sheet.write(1, column, value)
    path = tmp_path / "source.xls"
    workbook.save(str(path))
    assert import_table(path).rows == [{"标准名": "NA", "别名": "NULL", "编码": "001"}]
    sheet.row(1).set_cell_error(1, "#DIV/0!")
    workbook.save(str(path))
    with pytest.raises(WizardError, match=r"旧资料.*B2.*#DIV/0!"):
        import_table(path)
