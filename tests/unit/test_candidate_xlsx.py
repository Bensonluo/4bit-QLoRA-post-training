"""Candidate spreadsheets preserve evidence as text, never as formulas."""

import csv
import io
import zipfile

import pytest
from openpyxl import load_workbook

from src.workbench.learnability_probe import candidates_to_csv, candidates_to_xlsx


@pytest.mark.parametrize(
    "value", ["=1+2", "+1", "-1", "@SUM(A1)", "\t=1+2", "＝1+2", "#N/A", "😀\r\n原文"]
)
def test_candidate_values_stay_literal_and_match_original_csv(value):
    candidate = {
        "row_id": "r000001",
        "data_label": value,
        "base_zero_shot": value,
        "user_blind_answer": value,
        "evidence": value,
    }
    payload = candidates_to_xlsx([candidate])
    sheet = load_workbook(io.BytesIO(payload), data_only=False).active
    assert [sheet.cell(2, column).value for column in range(2, 6)] == [value] * 4
    assert all(sheet.cell(2, column).data_type == "s" for column in range(2, 6))
    assert sheet.freeze_panes == "B2"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")
    raw = candidates_to_csv([candidate]).decode("utf-8-sig")
    assert list(csv.reader(io.StringIO(raw)))[1] == ["r000001", *([value] * 4)]


def test_candidate_xlsx_contains_every_row_and_preserves_missing_values():
    candidates = [
        {"row_id": f"r{index:06d}", "data_label": "质量", "user_blind_answer": None}
        for index in range(25)
    ]
    sheet = load_workbook(io.BytesIO(candidates_to_xlsx(candidates)), data_only=False).active
    assert sheet.max_row == 26
    assert sheet["A26"].value == "r000024"
    assert sheet["D2"].value is None


def test_candidate_xlsx_rejects_too_many_rows_before_iterating():
    class TooManyRows(list):
        def __len__(self):
            return 1_048_576

        def __iter__(self):
            pytest.fail("oversized export should be rejected before reading rows")

    with pytest.raises(ValueError, match="行列限制"):
        candidates_to_xlsx(TooManyRows())


@pytest.mark.parametrize("value", ["a" * 32768, "😀" * 16384, "bad\x00text", "bad\ud800text"])
def test_candidate_xlsx_rejects_content_excel_cannot_preserve(value):
    with pytest.raises(ValueError, match="Excel"):
        candidates_to_xlsx([{"row_id": "r1", "data_label": value}])
