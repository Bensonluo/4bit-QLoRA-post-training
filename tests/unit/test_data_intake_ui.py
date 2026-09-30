"""UI flows run without GPU/MLflow and use the same production intake service."""

import hashlib
import io
import json
import shutil
from pathlib import Path

import pytest

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest

from src.workbench.business_evaluation import EvaluationReport
from src.workbench.demo_task import (
    DEMO_FULL_BUTTON,
    DEMO_FULL_NAME,
    DEMO_GOAL,
    DEMO_SAMPLE_BUTTON,
    DEMO_SAMPLE_ENTRY_LABEL,
    DEMO_SAMPLE_NAME,
    demo_full,
    demo_sample,
    is_demo_session,
)
from src.workbench.intake_service import IntakeService, next_action
from tests.unit.test_data_intake import CSV, analysis, model_for
from tests.unit.test_full_data import FULL

PAGE = Path(__file__).resolve().parents[2] / "ui/pages/07_Data_Intake.py"
REPO_EXAMPLES = Path(__file__).resolve().parents[2] / "data" / "custom" / "examples"


def stage_demo_files(project_root: Path) -> None:
    """把仓库自带的演示文件复制到该 PROJECT_ROOT，模拟文件在场的完整安装。"""
    target = project_root / "data" / "custom" / "examples"
    target.mkdir(parents=True, exist_ok=True)
    for name in (DEMO_SAMPLE_NAME, DEMO_FULL_NAME):
        shutil.copy(REPO_EXAMPLES / name, target / name)


@pytest.fixture()
def data_page(tmp_path, monkeypatch):
    import ui.config

    for name in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv(f"TUNESMITH_AGENT_{name}", raising=False)
    monkeypatch.setattr(ui.config, "PROJECT_ROOT", tmp_path)
    service = IntakeService(tmp_path / "outputs/workbench/intake")
    session = service.create("根据客户首次描述预测类别", "工单.csv", CSV)
    return service, session, AppTest.from_file(str(PAGE), default_timeout=20)


def button(page, label):
    return next(b for b in page.button if b.label == label)


def test_open_task_analyze_confirm_and_return_to_new(data_page, monkeypatch):
    import src.agent.intake

    service, session, page = data_page
    monkeypatch.setattr(
        src.agent.intake, "CompatibleChatClient", lambda *a, **kw: model_for(analysis())
    )
    page.run()
    assert not page.exception
    assert any(area.label == "希望模型完成什么业务工作？" for area in page.text_area)
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    button(page, "联合分析目标与数据").click().run()
    assert not page.exception
    assert any(s.value == "真实转换预览" for s in page.subheader)
    assert len(page.metric) == 4
    assert button(page, "确认当前转换含义").disabled
    next(c for c in page.checkbox if c.label.startswith("已核对预览")).check().run()
    button(page, "确认当前转换含义").click().run()
    assert not page.exception
    assert next_action(service.load(session.session_id)) == "awaiting_full_data"
    assert any("尚未认定可以正式训练" in message.value for message in page.success)
    button(page, "新建数据任务").click().run()
    assert not page.exception
    assert any(area.label == "希望模型完成什么业务工作？" for area in page.text_area)


def test_full_segment_banners_render_single_source_phrase_verbatim(data_page):
    """全量段横幅三态的页面等值钉(R91):页面选择逐字显示 next_action_phrase 的三个
    状态——awaiting_full_data/awaiting_full_validation/needs_data_revision——渲染文本
    必须逐字包含单一来源短语,R85/R90 的人工同步从此有测试锁定,页面词汇不再可能
    悄悄漂移。awaiting_full_data 横幅以「；尚未认定可以正式训练。」接边界句替换短语
    句号(页面惯例),子串断言剥去句号尾,其余词汇仍逐字锁定。其余状态的页面呈现是
    动作本体(review_preview 的配对表单、awaiting_dataset_split 的分区小节等),不显
    示命令行词汇,由页面旅程测试覆盖 cues,不在本钉范围。"""
    from src.workbench.intake_service import next_action_phrase

    service, session, page = data_page

    def rendered_text(chosen):
        page.run()
        page.selectbox(key="intake_select").select(chosen.session_id).run()
        assert not page.exception, [e.message for e in page.exception]
        return "\n".join(
            [m.value for m in page.markdown]
            + [w.value for w in page.warning]
            + [s.value for s in page.success]
            + [i.value for i in page.info]
            + [c.value for c in page.caption]
        )

    confirmed = service.apply_analysis(session, analysis())
    confirmed = service.confirm(confirmed.session_id, confirmed.revision)
    assert next_action(confirmed) == "awaiting_full_data"
    assert next_action_phrase("awaiting_full_data").rstrip("。") in rendered_text(confirmed)

    full_session = service.create("根据客户首次描述预测类别", "工单全量.csv", CSV, scope="full")
    full_session = service.apply_analysis(full_session, analysis())
    full_session = service.confirm(full_session.session_id, full_session.revision)
    assert next_action(full_session) == "awaiting_full_validation"
    assert next_action_phrase("awaiting_full_validation") in rendered_text(full_session)

    conflict_csv = (
        "编号,客户描述,类别,处理结果\n001,杯子破损,质量,补发\n002,杯子破损,物流,补发\n".encode()
    )
    broken = service.create("根据客户首次描述预测类别", "冲突.csv", conflict_csv)
    broken = service.apply_analysis(broken, analysis())
    assert next_action(broken) == "needs_data_revision"
    assert next_action_phrase("needs_data_revision") in rendered_text(broken)


def test_recipe_expander_renders_shared_tool_trace_line(data_page, monkeypatch):
    """处理规则与工具记录 expander:JSON 轨迹之外还有与评测解读同格式的人话行。"""
    import src.agent.intake

    service, session, page = data_page
    monkeypatch.setattr(
        src.agent.intake, "CompatibleChatClient", lambda *a, **kw: model_for(analysis())
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    button(page, "联合分析目标与数据").click().run()
    assert not page.exception
    assert any(expander.label == "处理规则与工具记录" for expander in page.expander)
    # model_for 夹具走 profile_data/inspect_rows/preview_recipe/submit_analysis 共 4 次成功调用
    assert any("工具核查轨迹：4 次调用，成功 4 次。" in m.value for m in page.markdown)


def test_iteration_revision_expander_renders_tool_trace_line(data_page):
    """修订记录 expander:tool_trace 的 JSON 之外还有同一格式的人话行（含失败计数）。"""
    import ui.config
    from src.workbench.iterations import IterationService

    service, session, page = data_page
    # 页面只读改进轮次记录；这里按 propose()+confirm()+revise_data_for_iteration()
    # 落盘的生产形状 (src/workbench/iterations.py:179-202、src/agent/revisions.py:42-53)
    # 直接写一条 confirmed 且带 data_revision 的记录，免拉起父轮训练链。
    iteration_service = IterationService(
        ui.config.PROJECT_ROOT / "outputs/workbench/iterations",
        ui.config.PROJECT_ROOT / "outputs/workbench/training",
        ui.config.PROJECT_ROOT / "outputs/workbench/evaluations",
    )
    suite_id = "es-" + "5c9d" * 8
    suite_root = ui.config.PROJECT_ROOT / "outputs/workbench/iterations/suites"
    iteration_service._save(
        {
            "iteration_id": "it-" + "3f1c" * 8,
            "session_id": session.session_id,
            "goal": session.goal,
            "parent_run_id": "tr-" + "0d2a" * 8,
            "parent_evaluation_id": "ev-" + "7b4e" * 8,
            "parent_report_digest": "digest",
            "evaluation_protocol": {"name": "fixed_suite"},
            "parent_evidence_binding": "content_verified",
            "evaluation_suite": {
                "suite_id": suite_id,
                "root": str(suite_root),
                "manifest_path": str(suite_root / f"{suite_id}.json"),
                "case_counts": {"validation": 1, "test": 1},
                "cases_digest": "digest",
            },
            "hypothesis": "补充坏例提升同题对照",
            "expected_outcome": "同题对照分数提高",
            "changes": "重新生成数据处理规则",
            "data_change": True,
            "training_start": "base",
            "model_path": str(ui.config.PROJECT_ROOT / "models/base"),
            "options": {},
            "parent_dataset": {},
            "proposal_dataset_digest": "digest",
            "new_run_id": None,
            "evaluation_id": None,
            "status": "confirmed",
            "created_at": "2026-09-28T00:00:00+00:00",
            "confirmed_session_revision": session.revision,
            "data_revision": {
                "from_revision": session.revision,
                "to_revision": session.revision + 1,
                "analysis_digest": "digest",
                "changed_components": ["recipe"],
                "before": {"recipe": None},
                "after": {"recipe": {"targets": []}},
                "next_action": "review_preview",
                "model": "stub-model",
                "tool_trace": [
                    {"tool": "profile_data", "ok": True},
                    {"tool": "inspect_rows", "ok": True},
                    {"tool": "preview_recipe", "ok": False, "error": "样例行不足"},
                ],
                "scope": "executed_preview_requires_business_confirmation",
            },
        }
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert any(expander.label.startswith("本轮 ") for expander in page.expander), [
        item.label for item in page.expander
    ]
    assert any(
        "工具核查轨迹：3 次调用，成功 2 次、失败 1 次"
        "——失败的调用没有取到证据，修订只依赖成功的调用。" in m.value
        for m in page.markdown
    ), [m.value for m in page.markdown if "工具核查轨迹" in m.value]


def test_plan_expander_renders_tool_trace_line(data_page):
    """训练方案 expander:trace 的 JSON 之外还有同一格式的人话行（含失败计数）。"""
    import sqlite3
    from datetime import datetime, timezone

    import ui.config
    from src.workbench.sources import canonical
    from src.workbench.training_plans import SCOPE, TrainingPlanService
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    # 方案区在 ready_for_training_preflight / preflight_passed 双态渲染(R94)：走既有夹具链
    # 分析→确认→全量校验→确认全量→物化 (同 test_preflight_only_loads_tokenizer…)。
    session = approved(service)
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session = service.confirm_full_data(session.session_id, session.revision)
    session = service.materialize_dataset(session.session_id, session.revision)
    assert next_action(session) == "ready_for_training_preflight"

    plans = TrainingPlanService(
        ui.config.PROJECT_ROOT / "outputs/workbench/training-plans",
        ui.config.PROJECT_ROOT / "outputs/workbench/training",
    )
    # 按 save() 落盘的生产记录形状 (src/workbench/training_plans.py:423-437) 直插一条
    # needs_data 方案：save() 的 ready 分支需要真实本地模型与 tokenizer 预检。
    plan_id = "tp-" + "9d61" * 8
    record = {
        "plan_id": plan_id,
        "session_id": session.session_id,
        "session_revision": session.revision,
        "dataset": session.dataset.model_dump(),
        "model_identity": None,
        "proposal": {
            "model_path": "",
            "max_length": 512,
            "rationale": ["样例里目标列缺唯一监督来源，先补数据再训练。"],
            "limitations": ["尚未在全量数据上验证该方案。"],
            "business_questions": [],
            "status": "needs_data",
            "training_options": {
                "num_epochs": 3,
                "batch_size": 1,
                "gradient_accumulation_steps": 8,
                "learning_rate": 0.0002,
            },
            "lora_options": {"r": 8},
            "model_options": {"quantization_bits": 4},
        },
        "status": "needs_data",
        "context": {"session_id": session.session_id},
        "trace": [
            {"tool": "inspect_context", "ok": True},
            {"tool": "probe", "ok": False, "error": "tokenizer 文件缺失"},
        ],
        "probe": None,
        "run_id": None,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "scope_note": SCOPE,
    }
    with sqlite3.connect(plans.database) as connection:
        connection.execute("INSERT INTO plans VALUES (?, ?)", (plan_id, canonical(record)))

    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    # 旧就绪句在场钉(r94-reviewer nit-2):未跑预检的 ready 态走 else 分支渲染原句,
    # 回归成两态同句时此钉先红——与新态的退场钉(test_preflight_passed…:1385)互为对偶。
    assert any(
        "独立数据分区已生成，可继续训练前检查" in message.value for message in page.success
    ), "ready_for_training_preflight 态的旧就绪句必须在场"
    assert any(expander.label.startswith("方案 ") for expander in page.expander), [
        item.label for item in page.expander
    ]
    assert any(
        "工具核查轨迹：2 次调用，成功 1 次、失败 1 次"
        "——失败的调用没有取到证据，方案只依赖成功的调用。" in m.value
        for m in page.markdown
    ), [m.value for m in page.markdown if "工具核查轨迹" in m.value]


def test_business_question_answer_survives_analysis_and_clears_widget(data_page, monkeypatch):
    import src.agent.intake

    service, session, page = data_page
    pending = analysis(
        recipe=None,
        questions=[
            {"question_id": "q", "question": "类别由谁审核？", "why": "需要可信的监督答案来源。"}
        ],
    )
    service.apply_analysis(session, pending)
    monkeypatch.setattr(
        src.agent.intake, "CompatibleChatClient", lambda *a, **kw: model_for(analysis())
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert any("类别由谁审核" in w.value for w in page.warning)
    next(t for t in page.text_area if t.label == "回答问题或修正理解").input(
        "类别由专门质检人员审核。"
    )
    button(page, "根据补充说明重新分析").click().run()
    assert not page.exception
    assert service.load(session.session_id).answers[-1]["answer"] == "类别由专门质检人员审核。"
    assert next(t for t in page.text_area if t.label == "回答问题或修正理解").value == ""


def test_remote_service_requires_data_consent_in_ui(data_page):
    service, session, page = data_page
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    next(t for t in page.text_input if t.label == "模型服务 API 地址").input(
        "https://example.test/v1"
    ).run()
    next(t for t in page.text_input if t.label == "支持工具调用的模型名称").input("model")
    button(page, "联合分析目标与数据").click().run()
    assert not page.exception
    assert any("需先允许" in error.value for error in page.error)
    assert service.load(session.session_id).analysis is None


def probe_client(monkeypatch):
    import src.agent.intake

    requests = []

    class Probe:
        def __init__(self, base_url, model, api_key, **kwargs):
            requests.append({"base_url": base_url, "model": model, "api_key": api_key, **kwargs})

        def complete(self, messages, tools):
            requests[-1]["messages"] = messages
            return {
                "tool_calls": [
                    {"function": {"name": "connection_check", "arguments": '{"ok": true}'}}
                ]
            }

    monkeypatch.setattr(src.agent.intake, "CompatibleChatClient", Probe)
    return requests


def test_configure_before_task_and_save_without_credentials(data_page, monkeypatch, tmp_path):
    _, _, page = data_page
    requests = probe_client(monkeypatch)
    page.run()
    page.selectbox(key="agent_provider").select("glm-coding").run()
    assert not page.exception
    assert "coding" in page.text_input(key="agent_base_url").value
    page.text_input(key="agent_api_key").input("test-session-secret").run()
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == "test-session-secret"
    assert "客户" not in json.dumps(requests[-1]["messages"], ensure_ascii=False)
    assert not page.checkbox(key="agent_remote_consent").value
    button(page, "保存模型配置").click().run()
    path = tmp_path / "outputs/workbench/agent-settings.json"
    saved = json.loads(path.read_text())
    assert set(saved) == {"provider", "base_url", "model"}
    assert "test-session-secret" not in path.read_text()
    reopened = AppTest.from_file(str(PAGE), default_timeout=20).run()
    assert not reopened.exception
    assert reopened.selectbox(key="agent_provider").value == "glm-coding"
    assert reopened.text_input(key="agent_api_key").value == ""
    assert not reopened.checkbox(key="agent_remote_consent").value


def test_switch_provider_and_endpoint_clear_key_and_consent(data_page, monkeypatch):
    _, _, page = data_page
    requests = probe_client(monkeypatch)
    page.run()
    page.selectbox(key="agent_provider").select("glm-coding").run()
    page.text_input(key="agent_api_key").input("first-provider-secret").run()
    page.checkbox(key="agent_remote_consent").check().run()
    page.selectbox(key="agent_provider").select("glm").run()
    assert page.text_input(key="agent_api_key").value == ""
    assert not page.checkbox(key="agent_remote_consent").value
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == ""
    page.text_input(key="agent_api_key").input("second-provider-secret").run()
    page.checkbox(key="agent_remote_consent").check().run()
    page.text_input(key="agent_base_url").input("https://different.example/v1").run()
    assert page.text_input(key="agent_api_key").value == ""
    assert not page.checkbox(key="agent_remote_consent").value
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == ""
    assert not page.exception


def test_environment_key_is_bound_to_initial_provider_and_endpoint(data_page, monkeypatch):
    _, _, page = data_page
    monkeypatch.setenv("TUNESMITH_AGENT_PROVIDER", "glm-coding")
    monkeypatch.setenv("TUNESMITH_AGENT_API_KEY", "test-environment-secret")
    requests = probe_client(monkeypatch)
    page.run()
    assert page.text_input(key="agent_api_key").value == ""
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == "test-environment-secret"
    page.text_input(key="agent_base_url").input("https://another.example/v1").run()
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == ""
    page.selectbox(key="agent_provider").select("glm").run()
    button(page, "测试模型连接").click().run()
    assert requests[-1]["api_key"] == ""
    assert not page.exception


def upload_full_file(monkeypatch, contents):
    import streamlit

    original = streamlit.file_uploader
    uploaded = io.BytesIO(contents)
    uploaded.name = "full.csv"

    def uploader(label, *args, **kwargs):
        return uploaded if label == "提供本次任务的全量文件" else original(label, *args, **kwargs)

    monkeypatch.setattr(streamlit, "file_uploader", uploader)


def test_full_upload_review_and_confirm(data_page, monkeypatch):
    from tests.unit.test_full_data import FULL

    service, session, page = data_page
    session = service.apply_analysis(session, analysis())
    session = service.confirm(session.session_id, session.revision)
    upload_full_file(monkeypatch, FULL)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    button(page, "按已确认方案验证全量数据").click().run()
    assert not page.exception
    current = service.load(session.session_id)
    assert next_action(current) == "review_full_data"
    assert current.source == session.source
    assert button(page, "确认全量数据含义").disabled
    next(c for c in page.checkbox if c.label.startswith("已核对全量报告")).check().run()
    button(page, "确认全量数据含义").click().run()
    assert next_action(service.load(session.session_id)) == "awaiting_dataset_split"
    assert any("独立训练与评测分区" in message.value for message in page.success)


def test_full_upload_blockers_cannot_be_confirmed(data_page, monkeypatch):
    service, session, page = data_page
    session = service.apply_analysis(session, analysis())
    session = service.confirm(session.session_id, session.revision)
    upload_full_file(monkeypatch, "编号,客户描述,类别,处理结果\n10,破损,,补发\n".encode())
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    button(page, "按已确认方案验证全量数据").click().run()
    assert not page.exception
    assert next_action(service.load(session.session_id)) == "needs_full_data_revision"
    assert any("缺少监督答案" in message.value for message in page.error)
    assert not any(b.label == "确认全量数据含义" for b in page.button)


def test_full_initial_source_reuse_and_business_feedback_invalidation(data_page):
    service, _, page = data_page
    session = service.create("分类", "full.csv", CSV, scope="full")
    session = service.apply_analysis(session, analysis())
    session = service.confirm(session.session_id, session.revision)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    button(page, "验证首次上传的全量文件").click().run()
    assert next_action(service.load(session.session_id)) == "review_full_data"
    next(t for t in page.text_area if t.label == "回答问题或修正理解").input(
        "类别含义需要修正"
    ).run()
    button(page, "保存业务补充，稍后分析").click().run()
    assert not page.exception
    assert service.load(session.session_id).full_data.status == "stale"
    assert any("全量报告已失效" in message.value for message in page.warning)
    assert not any(b.label == "确认全量数据含义" for b in page.button)


@pytest.mark.parametrize("groups", [["编号"], []])
def test_materialize_actual_partitions_after_full_confirmation(data_page, groups):
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    session = approved(service, group_columns=groups)
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session = service.confirm_full_data(session.session_id, session.revision)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    if not groups:
        assert button(page, "生成数据集版本").disabled
        next(c for c in page.checkbox if c.label.startswith("已确认每行是独立")).check().run()
    button(page, "生成数据集版本").click().run()
    assert not page.exception
    current = service.load(session.session_id)
    assert next_action(current) == "ready_for_training_preflight"
    assert current.dataset.statistics["row_counts"] == {"train": 1, "validation": 1, "test": 1}
    assert current.dataset.data_config["validation_split"] == 0
    assert any("独立数据分区已生成" in message.value for message in page.success)
    # 答案覆盖披露(第 30 轮):默认 seed 42 下 质量 整组落在验证/测试、训练只见 物流,
    # 统计点名缺口且摘要行真的渲染到页面——披露不埋进 JSON
    assert current.dataset.statistics["train_missing_answers"] == {
        "质量": {"validation": 1, "test": 1}
    }
    assert any("从未出现在训练集" in item.value for item in page.markdown), [
        item.value for item in page.markdown
    ]


def test_split_settings_guidance_renders_when_ratios_apply(data_page):
    """分区设置 expander:比例输入真正生效的分支渲染「何时该改」引导行(单一来源)。"""
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    session = approved(service, group_columns=["编号"])
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session = service.confirm_full_data(session.session_id, session.revision)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert any(expander.label == "分区设置" for expander in page.expander)
    captions = "\n".join(caption.value for caption in page.caption)
    assert "验证集比例＝" in captions
    assert "独立测试集比例＝" in captions
    assert "可复现分区种子＝" in captions
    # 引导行挂在比例真正生效的分支上:三个输入框与引导同时在场
    for label in ("验证集比例", "独立测试集比例", "可复现分区种子"):
        next(field for field in page.number_input if field.label == label)


SUFFICIENT_FULL = (
    "编号,客户描述,类别,处理结果\n".encode()
    + "".join(
        f"{index},客户描述{index},{'质量' if index % 2 else '物流'},补发\n" for index in range(400)
    ).encode()
)


@pytest.mark.parametrize(
    ("payload", "expect_caution"),
    [(FULL, True), (SUFFICIENT_FULL, False)],
    ids=["small-test-set", "sufficient-test-set"],
)
def test_dataset_artifacts_small_test_set_caution(data_page, payload, expect_caution):
    """分区产物区:独立测试集 <30 条时给出每题权重与偶然性提醒;条数够时不制造噪音。"""
    from tests.unit.test_full_data import approved

    service, _, page = data_page
    session = approved(service, group_columns=["编号"])
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", payload)
    session = service.confirm_full_data(session.session_id, session.revision)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    button(page, "生成数据集版本").click().run()
    assert not page.exception
    counts = service.load(session.session_id).dataset.statistics["row_counts"]
    if expect_caution:
        assert counts["test"] < 30
        assert any(f"独立测试集共 {counts['test']} 条" in m.value for m in page.markdown), [
            m.value for m in page.markdown
        ]
        assert any("结论偶然性大" in m.value for m in page.markdown)
        assert not any("独立测试集共" in c.value for c in page.caption)
    else:
        assert counts["test"] >= 30
        assert not any("独立测试集共" in m.value for m in page.markdown)
        assert not any("结论偶然性大" in m.value for m in page.markdown)


def test_add_original_source_preserves_business_description(data_page, monkeypatch):
    import streamlit

    service, session, page = data_page
    original = streamlit.file_uploader
    upload = io.BytesIO("编号,审核类别\n001,质量\n".encode())
    upload.name = "labels.csv"
    monkeypatch.setattr(
        streamlit,
        "file_uploader",
        lambda label, *a, **kw: (
            upload if label == "上传补充原始资料" else original(label, *a, **kw)
        ),
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    next(t for t in page.text_input if t.label == "补充资料名称").input("labels")
    next(t for t in page.text_area if t.label == "这份资料的用途和关联关系").input(
        "人工审核标签，通过编号关联。"
    )
    button(page, "保存补充资料").click().run()
    assert not page.exception
    saved = service.load(session.session_id)
    assert set(saved.sources) == {"main", "labels"}
    assert saved.sources["main"].digest == session.source.digest
    assert "人工审核标签" in saved.answers[-1]["answer"]


def test_combined_sources_show_lineage_and_accept_separate_full_uploads(data_page, monkeypatch):
    import streamlit

    from tests.unit.test_multisource_cli import FULL_LABELS, FULL_MAIN, composition_session

    service, _, page = data_page
    session = composition_session(service)
    original = streamlit.file_uploader
    uploads = {}
    for alias, contents in (("main", FULL_MAIN), ("labels", FULL_LABELS)):
        value = io.BytesIO(contents)
        value.name = f"{alias}.csv"
        uploads[f"全量原始资料：{alias}"] = value
    monkeypatch.setattr(
        streamlit,
        "file_uploader",
        lambda label, *a, **kw: uploads[label] if label in uploads else original(label, *a, **kw),
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert any(item.value == "资料组合与处理结果" for item in page.subheader)
    assert not any(b.label == "验证首次上传的全量文件" for b in page.button)
    button(page, "按组合方案验证全部全量资料").click().run()
    assert not page.exception
    saved = service.load(session.session_id)
    assert set(saved.full_data.sources) == {"main", "labels"}
    assert saved.full_data.preview.counts["ready"] == 3
    assert any(item.label == "全量资料组合过程与原始来源" for item in page.expander)


def test_preflight_only_loads_tokenizer_on_explicit_button_and_shows_row_failures(
    data_page, monkeypatch
):
    import src.workbench.training_preflight as preflight
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    session = approved(service)
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session = service.confirm_full_data(session.session_id, session.revision)
    session = service.materialize_dataset(session.session_id, session.revision)
    calls = []
    token = object()

    def load(path, *, local_files_only):
        calls.append((path, local_files_only))
        return token

    def report(current, tokenizer, max_length):
        assert tokenizer is token
        assert max_length == 6
        return {
            "status": "blocked",
            "scope_note": "UI protocol fixture",
            "issues": [
                {
                    "code": "answer_lost",
                    "severity": "blocking",
                    "message": "答案被截断",
                    "split": "train",
                    "row_ids": ["r000001"],
                }
            ],
            "splits": {"train": {"answer_lost_rows": 1}},
            "rows": [{"row_id": "r000001", "answer_supervised_tokens": 0}],
            "tokenizer": {"name_or_path": "local-fixture"},
        }

    monkeypatch.setattr(preflight, "load_local_tokenizer", load)
    monkeypatch.setattr(preflight, "preflight_dataset", report)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert calls == []
    next(t for t in page.text_input if t.label == "本地 tokenizer 目录或已缓存标识").input(
        "/tmp/local-fixture"
    )
    next(n for n in page.number_input if n.label == "训练最大 token 长度").set_value(6)
    button(page, "检查实际截断与答案保留").click().run()
    assert not page.exception
    assert calls == [("/tmp/local-fixture", True)]
    assert any("答案被截断" in entry.value for entry in page.error)
    # 三态横幅单一来源(R95):blocked 状态句与 CLI summarize_preflight 同源,
    # 手抄短句「存在阻断问题，请先修正数据或训练长度」退场;边界句同源渲染。
    assert any("训练前检查发现阻断问题，训练不能开始" in entry.value for entry in page.error), (
        "blocked 横幅必须渲染单源状态句"
    )
    assert not any("存在阻断问题，请先修正数据或训练长度" in entry.value for entry in page.error)
    assert any("不代表训练效果或业务达标" in entry.value for entry in page.caption)
    # 中段建议行同源铺出(r95-reviewer nit-1 采纳):丢答案引导在场,问题行不重复铺
    assert any("有 1 行的答案在截断后完全丢失" in entry.value for entry in page.markdown), (
        "blocked 报告的建议行必须在场"
    )
    assert not any("[blocking]" in entry.value for entry in page.markdown), "问题行走分级告警"
    assert any("r000001" in entry.value for entry in page.caption)
    page.run()
    assert len(calls) == 1


def test_adapter_evidence_shows_failed_full_validation_without_executing_source(data_page):
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    session = approved(service)
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session.analysis.adapter = {
        "source_code": 'raise AssertionError("UI must never execute adapter source")',
        "new_columns": ["parsed_category"],
        "config": {},
        "examples": [],
    }
    session.adapter_report = {
        "validation": {
            "status": "passed",
            "backend": "macos",
            "source_digest": "code-digest",
            "cases_digest": "cases-digest",
            "limits": {"memory_enforcement": "sampled_rss"},
            "cases": [
                {"name": "真实样例", "kind": "business", "passed": True},
                {"name": "缺失字段反例", "kind": "counterexample", "passed": True},
            ],
        },
        "spec_digest": "spec-digest",
        "origins": {"r000001": [{"source_digest": session.source.digest, "row_id": "r000001"}]},
    }
    session.full_data.adapter_report = {
        "validation": {
            "status": "failed",
            "backend": "macos",
            "cases": [
                {
                    "name": "全量新增格式",
                    "kind": "business",
                    "passed": False,
                    "error": "无法解析新格式",
                }
            ],
        }
    }
    service._save(session, session.revision)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert any(entry.value == "受限适配验证结果" for entry in page.subheader)
    assert any(entry.value == "全量受限适配验证结果" for entry in page.subheader)
    assert any("隔离测试未通过" in entry.value for entry in page.error)
    assert any("实际隔离后端：macos" in entry.value for entry in page.caption)
    assert any("UI must never execute" in entry.value for entry in page.code)
    rows = [frame.value for frame in page.dataframe]
    assert any("结果" in frame and "未通过" in frame["结果"].values for frame in rows)


def _preview_csv(count, *, missing_from=None):
    return (
        "编号,客户描述,类别,处理结果\n"
        + "".join(
            f"{index:03d},独立问题{index},{'' if missing_from is not None and index >= missing_from else '质量'},补发\n"
            for index in range(1, count + 1)
        )
    ).encode()


def test_sample_problem_after_twenty_rows_is_visible_and_has_concrete_label_next_step(data_page):
    service, _, page = data_page
    session = service.create("判断类别", "sample.csv", _preview_csv(25, missing_from=25))
    session = service.apply_analysis(session, analysis())
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert next_action(service.load(session.session_id)) == "needs_labels"
    assert any(item.label == "r000025 · 缺少答案" for item in page.expander)
    assert any("独立问题25" in item.value for item in page.code)
    assert any(
        "答案（字段：类别）" in item.value and "原始资料与补充文件" in item.value
        for item in page.info
    )
    selector = next(item for item in page.selectbox if item.label == "样例记录筛选")
    selector.select("缺少答案").run()
    row_labels = [item.label for item in page.expander if item.label.startswith("r000")]
    assert row_labels == ["r000025 · 缺少答案"]
    assert not any(item.label == "确认当前转换含义" for item in page.button)


def test_missing_answer_handoff_sheet_renders_rules_and_download(data_page):
    """缺答案时页面给交接物:下载按钮+三条规则与 CLI answer-sheet 同源。"""
    service, _, page = data_page
    session = service.create("判断类别", "sample.csv", _preview_csv(25, missing_from=25))
    session = service.apply_analysis(session, analysis())
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    # AppTest 没有 download_button 专用访问器(不在 page.button 里),按元素类型
    # 从消息树取节点、按标签过滤(页面还有分区下载等常驻下载按钮);CSV 内容
    # 本身由单元/CLI 测试钉住,这里钉住按钮在缺答案时出现。
    downloads = [node.proto.label for node in page.get("download_button")]
    assert "导出待补答案清单（交给填写人）" in downloads
    captions = " ".join(item.value for item in page.caption)
    assert "只填「待填答案」列" in captions
    assert "不要凭猜测填" in captions
    assert "监督信号" in captions
    assert "比编一个答案更有价值" in captions


def test_missing_answer_rules_absent_when_no_missing_rows(data_page):
    """没有缺答案行:规则与下载按钮都不渲染——不拿空清单冒充交接物。"""
    service, _, page = data_page
    session = service.create("判断类别", "sample.csv", _preview_csv(25))
    session = service.apply_analysis(session, analysis())
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    downloads = [node.proto.label for node in page.get("download_button")]
    assert "导出待补答案清单（交给填写人）" not in downloads
    captions = " ".join(item.value for item in page.caption)
    assert "只填「待填答案」列" not in captions
    assert "不要凭猜测填" not in captions


def test_sample_pagination_confirms_only_visible_rows_and_resets_acknowledgment(data_page):
    service, _, page = data_page
    session = service.create("判断类别", "sample.csv", _preview_csv(25))
    session = service.apply_analysis(session, analysis())
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    next(item for item in page.checkbox if item.label.startswith("已核对预览")).check().run()
    next(item for item in page.number_input if item.label == "样例预览页码").set_value(2).run()
    assert not page.exception
    assert button(page, "确认当前转换含义").disabled
    assert any(item.label.startswith("r000025") for item in page.expander)
    assert not any(item.label.startswith("r000001") for item in page.expander)
    next(item for item in page.checkbox if item.label.startswith("已核对预览")).check().run()
    button(page, "确认当前转换含义").click().run()
    assert not page.exception
    assert {row.row_id for row in service.load(session.session_id).confirmed_examples} == {
        f"r{index:06d}" for index in range(21, 26)
    }


def test_full_problem_pagination_reaches_all_issue_rows_and_points_to_full_reupload(data_page):
    from tests.unit.test_full_data import approved

    service, _, page = data_page
    session = approved(service)
    session = service.validate_full_data(
        session.session_id, session.revision, "full.csv", _preview_csv(45, missing_from=21)
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    next(item for item in page.selectbox if item.label == "全量记录筛选").select("缺少答案").run()
    next(item for item in page.number_input if item.label == "全量预览页码").set_value(2).run()
    assert not page.exception
    assert any(item.label == "全量 r000045 · 缺少答案" for item in page.expander)
    assert any("独立问题45" in item.value for item in page.code)
    assert any("全量数据验证」重新上传" in item.value for item in page.info)
    assert not any(item.label == "确认全量数据含义" for item in page.button)


SHEET_INPUT_LABEL = "Excel 工作表（留空读第一个）"


def _workbook_bytes(*sheets):
    """sheets: (名称, 表头行, 数据行) 三元组；openpyxl 现场生成多 Sheet 工作簿。"""
    from openpyxl import Workbook

    workbook = Workbook()
    for index, (name, header, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def uploaded_bytes(contents, name):
    payload = io.BytesIO(contents)
    payload.name = name
    return payload


def patch_uploader(monkeypatch, label, holder):
    """按标签替换 file_uploader 的返回值；holder["file"] 可在测试中途更换。"""
    import streamlit

    original = streamlit.file_uploader

    def uploader(rendered, *args, **kwargs):
        return holder["file"] if rendered == label else original(rendered, *args, **kwargs)

    monkeypatch.setattr(streamlit, "file_uploader", uploader)


def test_new_intake_form_reads_designated_excel_sheet(data_page, monkeypatch):
    """新建任务表单：仅 Excel 上传显示 sheet 输入，指定后读到对应工作表的数据。"""
    service, existing, page = data_page
    holder = {"file": None}
    patch_uploader(monkeypatch, "提供 CSV、Excel 或 JSONL", holder)
    page.run()
    assert not page.exception
    assert not any(t.label == SHEET_INPUT_LABEL for t in page.text_input)
    holder["file"] = uploaded_bytes(CSV, "工单.csv")
    page.run()
    assert not any(t.label == SHEET_INPUT_LABEL for t in page.text_input)
    holder["file"] = uploaded_bytes(
        _workbook_bytes(
            ("工单表", ("编号", "类别"), [("001", "质量")]),
            ("员工表", ("员工号", "部门"), [("E01", "质检")]),
        ),
        "花名册.xlsx",
    )
    page.run()
    assert not page.exception
    next(t for t in page.text_area if t.label == "希望模型完成什么业务工作？").input("统计员工部门")
    next(t for t in page.text_input if t.label == SHEET_INPUT_LABEL).input("员工表")
    button(page, "读取数据并开始").click().run()
    assert not page.exception
    created = next(s for s in service.list_sessions() if s.session_id != existing.session_id)
    loaded = service.load(created.session_id)
    assert loaded.source.sheet == "员工表"
    assert loaded.source.columns == ["员工号", "部门"]
    assert any(row.values.get("员工号") == "E01" for row in loaded.source.rows)
    assert loaded.source.sheet_note is not None and "工单表" in loaded.source.sheet_note


def test_add_source_form_reads_designated_excel_sheet(data_page, monkeypatch):
    """保存补充资料与新建任务对称：仅 Excel 上传显示 sheet 输入，指定后读到对应工作表。"""
    service, session, page = data_page
    holder = {"file": uploaded_bytes(CSV, "labels.csv")}
    patch_uploader(monkeypatch, "上传补充原始资料", holder)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert not any(t.label == SHEET_INPUT_LABEL for t in page.text_input)
    holder["file"] = uploaded_bytes(
        _workbook_bytes(
            ("工单表", ("编号", "类别"), [("001", "质量")]),
            ("员工表", ("员工号", "部门"), [("E01", "质检")]),
        ),
        "labels.xlsx",
    )
    page.run()
    next(t for t in page.text_input if t.label == "补充资料名称").input("labels")
    next(t for t in page.text_input if t.label == SHEET_INPUT_LABEL).input("员工表")
    button(page, "保存补充资料").click().run()
    assert not page.exception
    saved = service.load(session.session_id)
    assert saved.sources["labels"].sheet == "员工表"
    assert saved.sources["labels"].columns == ["员工号", "部门"]
    assert any(row.values.get("员工号") == "E01" for row in saved.sources["labels"].rows)
    assert saved.sources["main"].digest == session.source.digest


def test_full_upload_form_reads_designated_excel_sheet(data_page, monkeypatch):
    """全量文件读取设置与新建任务对称：全量 Excel 数据在第二个 sheet 时按指定读取。"""
    from tests.unit.test_full_data import approved

    service, _, page = data_page
    session = approved(service)
    holder = {"file": None}
    patch_uploader(monkeypatch, "提供本次任务的全量文件", holder)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert not any(t.label == SHEET_INPUT_LABEL for t in page.text_input)
    holder["file"] = uploaded_bytes(
        _workbook_bytes(
            ("说明", ("备注",), [("全量数据在下一个 sheet",)]),
            (
                "全量表",
                ("编号", "客户描述", "类别", "处理结果"),
                [
                    ("100", "收到破杯", "质量", "补发"),
                    ("101", "快递太慢", "物流", "查询"),
                    ("102", "杯把断了", "质量", "补发"),
                ],
            ),
        ),
        "全量.xlsx",
    )
    page.run()
    next(t for t in page.text_input if t.label == SHEET_INPUT_LABEL).input("全量表")
    button(page, "按已确认方案验证全量数据").click().run()
    assert not page.exception
    current = service.load(session.session_id)
    assert next_action(current) == "review_full_data"
    assert current.full_data.sources["main"].sheet == "全量表"
    assert current.full_data.source.columns == ["编号", "客户描述", "类别", "处理结果"]
    assert current.full_data.preview.counts["ready"] == 3


def _fact_note_workbook_bytes() -> bytes:
    """双 sheet + 合并区 + 隐藏行 + 全空行 + 重复表头行:一个工作簿同时触发
    sheet/merged/hidden/blank/dup_header 五条如实标注。"""
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单表"
    sheet.append(("编号", "客户描述", "类别"))
    for row in (
        ("001", "杯子破损", "质量"),
        (None, None, None),  # 第 3 行:全空行,照常读入为全空记录
        ("002", "物流未更新", None),  # C4:C5 合并,非首格读空
        ("003", "屏幕碎裂", None),
        ("004", "快递丢失", "物流"),
        ("编号", "客户描述", "类别"),  # 第 7 行:导出拼接产生的重复表头,照常读入为数据行
    ):
        sheet.append(row)
    sheet.merge_cells("C4:C5")
    sheet.row_dimensions[6].hidden = True  # 第 6 行(004,带有效标签)隐藏,照常读入
    staff = workbook.create_sheet("员工表")
    staff.append(("员工号", "部门"))
    staff.append(("E01", "质检"))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _full_fact_note_workbook_bytes() -> bytes:
    """全量侧夹具:隐藏行带着有效标签(场景 44 形态),全量验证照常通过。"""
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "全量表"
    sheet.append(("编号", "客户描述", "类别", "处理结果"))
    for row in (
        ("100", "收到破杯", "质量", "补发"),
        ("101", "快递太慢", "物流", "查询"),
        ("102", "杯把断了", "质量", "补发"),
    ):
        sheet.append(row)
    sheet.row_dimensions[3].hidden = True  # 101 行隐藏——有效数据,不拦旅程
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_fact_notes_render_after_create(data_page):
    """如实标注在任务页可见:主来源的 sheet(info)/merged/hidden/blank/dup_header(warning)
    在 scope_note 下按级渲染;CSV 任务没有 Excel 事实,一条都不渲染。"""
    service, existing, page = data_page
    # 两个会话都在首次 run 前建好:AppTest 的 selectbox 选项来自上一次渲染,
    # run 之后才 create 的会话不在旧选项里,select 会静默渲染回旧会话。
    created = service.create(
        "根据客户首次描述判断售后类别", "工单.xlsx", _fact_note_workbook_bytes()
    )
    page.run()
    page.selectbox(key="intake_select").select(existing.session_id).run()
    assert not page.exception
    # CSV 任务没有 Excel 事实:Excel 标注一条都不出现(空行/重复表头属跨格式事实,
    # 由 blank_note/dup_header_note 承担,干净 CSV 不携带)
    assert not any("合并单元格" in w.value for w in page.warning)
    assert not any("隐藏行" in w.value for w in page.warning)
    assert not any("全空行" in w.value for w in page.warning)
    assert not any("与表头完全相同" in w.value for w in page.warning)
    assert not any("sheet" in i.value for i in page.info)

    page.selectbox(key="intake_select").select(created.session_id).run()
    assert not page.exception
    assert any("1 处合并单元格" in w.value for w in page.warning), [w.value for w in page.warning]
    assert any("1 个隐藏行" in w.value for w in page.warning), [w.value for w in page.warning]
    assert any("1 个全空行" in w.value for w in page.warning), [w.value for w in page.warning]
    assert any("1 行与表头完全相同" in w.value for w in page.warning), [
        w.value for w in page.warning
    ]
    assert any("仅读取第一个" in i.value for i in page.info), [i.value for i in page.info]


def test_full_report_renders_excel_fact_notes(data_page, monkeypatch):
    """全量验证报告同样渲染 Excel 事实标注:全量文件的隐藏行在来源 caption 下可见。"""
    from tests.unit.test_full_data import approved

    service, _, page = data_page
    session = approved(service)
    holder = {"file": uploaded_bytes(_full_fact_note_workbook_bytes(), "全量.xlsx")}
    patch_uploader(monkeypatch, "提供本次任务的全量文件", holder)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    button(page, "按已确认方案验证全量数据").click().run()
    assert not page.exception
    assert next_action(service.load(session.session_id)) == "review_full_data"
    assert any("1 个隐藏行" in w.value for w in page.warning), [w.value for w in page.warning]


def test_added_source_fact_notes_listed_with_alias(data_page):
    """补充资料的 Excel 标注在「原始资料与补充文件」区按资料名前缀列出。"""
    service, session, page = data_page
    service.add_source(
        session.session_id,
        session.revision,
        "labels",
        "labels.xlsx",
        _fact_note_workbook_bytes(),
    )
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert any(w.value.startswith("labels：") and "1 个隐藏行" in w.value for w in page.warning), [
        w.value for w in page.warning
    ]
    assert any(w.value.startswith("labels：") and "1 个全空行" in w.value for w in page.warning), [
        w.value for w in page.warning
    ]
    assert any(
        w.value.startswith("labels：") and "1 行与表头完全相同" in w.value for w in page.warning
    ), [w.value for w in page.warning]


def test_composed_full_sources_form_reads_designated_excel_sheets(data_page, monkeypatch):
    """组合全量表单与单文件表单对称：每份 Excel 旁的 sheet 输入逐份透传。

    两份资料各指定不同 sheet，任一来源留空或读错 sheet 都无法蒙混成正确全量。
    """
    from tests.unit.test_multisource_cli import LABELS, MAIN, combined_plan

    service, _, page = data_page
    session = service.create("判断工单类别", "tickets.csv", MAIN)
    session = service.add_source(
        session.session_id, session.revision, "labels", "labels.csv", LABELS
    )
    session = service.apply_analysis(session, combined_plan())
    session = service.confirm(session.session_id, session.revision)
    holders = {"main": {"file": None}, "labels": {"file": None}}
    patch_uploader(monkeypatch, "全量原始资料：main", holders["main"])
    patch_uploader(monkeypatch, "全量原始资料：labels", holders["labels"])
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    assert not page.exception
    assert not any(t.label == SHEET_INPUT_LABEL for t in page.text_input)
    holders["main"]["file"] = uploaded_bytes(
        _workbook_bytes(
            ("说明", ("备注",), [("全量数据在下一个 sheet",)]),
            (
                "工单表",
                ("ticket", "text"),
                [("10", "新杯子损坏"), ("11", "新物流延迟"), ("12", "新配件缺陷")],
            ),
        ),
        "全量工单.xlsx",
    )
    holders["labels"]["file"] = uploaded_bytes(
        _workbook_bytes(
            ("类别表", ("id", "category"), [("10", "质量"), ("11", "物流"), ("12", "质量")]),
            ("说明", ("备注",), [("这份资料的数据在第一个 sheet",)]),
        ),
        "全量类别.xlsx",
    )
    page.run()
    assert not page.exception
    main_sheet = next(
        t for t in page.text_input if t.key == f"full_sheet_{session.session_id}_main"
    )
    labels_sheet = next(
        t for t in page.text_input if t.key == f"full_sheet_{session.session_id}_labels"
    )
    main_sheet.input("工单表")
    labels_sheet.input("类别表")
    button(page, "按组合方案验证全部全量资料").click().run()
    assert not page.exception
    current = service.load(session.session_id)
    assert next_action(current) == "review_full_data"
    assert current.full_data.sources["main"].sheet == "工单表"
    assert current.full_data.sources["labels"].sheet == "类别表"
    assert current.full_data.preview.counts["ready"] == 3


def _comparison_row(index: int, output: str, *, correct: bool) -> dict:
    prompt = "指" * 100
    return {
        "index": index,
        "source": {"source_row_id": f"r{index}"},
        "prompt": prompt,
        "expected": "甲",
        "output": output,
        "status": "ok",
        "correct": correct,
    }


def _comparison_report(outputs: list[str]) -> EvaluationReport:
    rows = [
        _comparison_row(index, output, correct=output.startswith("甲"))
        for index, output in enumerate(outputs)
    ]
    return EvaluationReport(
        evaluation_id="eval-echo-0001",
        created_at="2026-09-28T00:00:00+00:00",
        dataset={"purpose": "development_only"},
        protocol={"scorer": "classification_exact", "max_new_tokens": 256},
        comparison_key="",
        status="completed",
        models=[
            {
                "label": "本轮微调",
                "requested_model": {"base_model": "demo", "adapter_path": None},
                "metrics": {"total": len(rows), "scored": len(rows), "exact_match": 0.5},
                "rows": rows,
            }
        ],
    )


def _render_business_comparison(data_page, report) -> AppTest:
    """单函数渲染面：真实页面模块加载后直接调 show_business_comparison。

    页面模块名以数字开头无法常规 import；from_function 的脚本体须自包含，
    故在体内按文件加载页面模块（session_id / report / PAGE 经 args 传入），
    让函数引用到真实的模块级全局（session、agent 设置、PROJECT_ROOT）。
    """
    _, session, _ = data_page

    def render(session_id, report, page_path):
        import importlib.util

        import streamlit as st

        st.session_state["intake_id"] = session_id
        spec = importlib.util.spec_from_file_location("data_intake_page_render", page_path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.show_business_comparison(report, key="echo_triage")

    page = AppTest.from_function(
        render, default_timeout=20, args=(session.session_id, report, PAGE)
    )
    page.run()
    return page


def test_business_comparison_renders_echo_triage_warnings(data_page):
    """带回声模型的对照按 echo_triage_lines 分流渲染，替换旧的一条静态警告。"""
    prompt = "指" * 100
    page = _render_business_comparison(
        data_page,
        _comparison_report(
            [prompt[:24] + "，所以答案是甲。", "乙类问题应先转人工处理并附上说明。"]
        ),
    )
    assert not page.exception
    warnings_text = "\n".join(warning.value for warning in page.warning)
    assert "检测到指令回声——" in warnings_text
    assert "长度上限不是第一嫌疑" in warnings_text
    # 旧的静态一条式警告（全角括号拼 echo_models）不应再出现。
    assert "检测到指令回声（" not in warnings_text


def test_business_comparison_without_echo_renders_no_echo_warning(data_page):
    """无回声模型的对照不渲染任何指令回声警告。"""
    page = _render_business_comparison(
        data_page,
        _comparison_report(["甲类问题应转质量组处理。", "乙类问题应先转人工处理并附上说明。"]),
    )
    assert not page.exception
    assert not any("检测到指令回声" in warning.value for warning in page.warning)


def test_demo_task_entry_creates_real_session(data_page):
    """一键演示入口只代劳找文件+填表:创建的是真实任务,落在与真实任务相同的视图。"""
    import ui.config

    service, _session, page = data_page
    stage_demo_files(ui.config.PROJECT_ROOT)
    page.run()
    assert not page.exception
    assert any(expander.label == DEMO_SAMPLE_ENTRY_LABEL for expander in page.expander)
    button(page, DEMO_SAMPLE_BUTTON).click().run()
    assert not page.exception
    demo_sessions = [s for s in service.list_sessions() if s.goal == DEMO_GOAL]
    assert len(demo_sessions) == 1
    loaded = service.load(demo_sessions[0].session_id)
    sample_bytes = (
        ui.config.PROJECT_ROOT / "data" / "custom" / "examples" / DEMO_SAMPLE_NAME
    ).read_bytes()
    assert loaded.source.digest == hashlib.sha256(sample_bytes).hexdigest()
    assert loaded.source.scope == "sample"
    assert len(loaded.source.rows) == 2
    # 一键之后落在真实任务视图:后续每一步(基础分析、预览核对……)与真实任务完全相同
    assert any(
        expander.label == "没有 Agent 服务？用基础分析开始（产品内置判断，无需任何密钥）"
        for expander in page.expander
    )


def test_demo_full_companion_gated_by_digest_and_validates(data_page):
    """配套演示全量按钮按内容摘要门控:演示任务可见可验证,同摘要环境下的真实任务不可见。"""
    import ui.config
    from src.workbench.baseline_analysis import propose_baseline_analysis

    service, session, page = data_page
    stage_demo_files(ui.config.PROJECT_ROOT)
    sample = demo_sample(ui.config.PROJECT_ROOT)
    assert sample is not None
    assert is_demo_session(hashlib.sha256(sample[1]).hexdigest(), ui.config.PROJECT_ROOT)

    def confirmed(source_session):
        updated = service.apply_analysis(
            source_session,
            propose_baseline_analysis(
                source_session,
                target_column="类别",
                group_columns=["编号"],
                excluded_columns=["处理结果"],
            ),
            model="baseline-deterministic",
        )
        return service.confirm(updated.session_id, updated.revision)

    demo = confirmed(service.create(DEMO_GOAL, sample[0], sample[1], data_description="演示数据"))
    # 真实任务用内容不同的数据:夹具 CSV 与演示文件逐字节相同(旅程测试同一份数据),
    # 内容摘要一致的来源本来就算演示来源,不能充当「真实任务看不见按钮」的反例。
    real_bytes = (
        "编号,客户描述,类别,处理结果\n101,把手松动,质量,维修\n102,配送延迟,物流,补偿\n"
    ).encode()
    real = confirmed(service.create("真实售后分类目标", "真实工单.csv", real_bytes))

    page.run()
    page.selectbox(key="intake_select").select(real.session_id).run()
    assert not page.exception
    assert all(item.label != DEMO_FULL_BUTTON for item in page.button)

    page.selectbox(key="intake_select").select(demo.session_id).run()
    assert not page.exception
    button(page, DEMO_FULL_BUTTON).click().run()
    assert not page.exception
    loaded = service.load(demo.session_id)
    assert next_action(loaded) == "review_full_data"
    assert loaded.full_data.source.name == DEMO_FULL_NAME
    assert len(loaded.full_data.source.rows) == 10
    assert button(page, "确认全量数据含义").disabled  # 仍需逐行核对后勾选,门禁未被跳过


def test_demo_task_module_degrades_honestly_when_files_absent(tmp_path):
    """演示文件不在场时如实返回 None/False:入口隐藏,不编造演示数据。"""
    assert demo_sample(tmp_path) is None
    assert demo_full(tmp_path) is None
    assert not is_demo_session("0" * 64, tmp_path)


def test_preflight_passed_state_moves_page_exits_forward(data_page, monkeypatch):
    """预检通过后页面就绪句与方案区出口前移(R94):状态机新增 preflight_passed 态,
    页面成功行改渲染单一来源短语(与 CLI 尾行同词汇),不再回环指向已完成的预检;
    盲标核验/固定题集/业务评分/微调方案各区块在该态继续可见(可重跑、可多轮改进)。"""
    import src.workbench.training_preflight as preflight
    from tests.unit.test_full_data import FULL, approved

    service, _, page = data_page
    session = approved(service)
    session = service.validate_full_data(session.session_id, session.revision, "full.csv", FULL)
    session = service.confirm_full_data(session.session_id, session.revision)
    session = service.materialize_dataset(session.session_id, session.revision)
    assert next_action(session) == "ready_for_training_preflight"

    def report(current, tokenizer, max_length):
        return {
            "status": "passed",
            "scope_note": "UI protocol fixture",
            "issues": [],
            "splits": {"train": {"answer_lost_rows": 0}},
            "rows": [],
            "tokenizer": {"name_or_path": "local-fixture"},
        }

    monkeypatch.setattr(
        preflight, "load_local_tokenizer", lambda path, *, local_files_only: object()
    )
    monkeypatch.setattr(preflight, "preflight_dataset", report)
    page.run()
    page.selectbox(key="intake_select").select(session.session_id).run()
    next(t for t in page.text_input if t.label == "本地 tokenizer 目录或已缓存标识").input(
        "/tmp/local-fixture"
    )
    next(n for n in page.number_input if n.label == "训练最大 token 长度").set_value(32)
    button(page, "检查实际截断与答案保留").click().run()
    assert not page.exception
    assert next_action(service.load(session.session_id)) == "preflight_passed"
    # 就绪句前移:单一来源短语在场,旧的「可继续训练前检查」回环句退场
    assert any(
        "训练前检查已通过：可运行 plan-recommend" in message.value for message in page.success
    )
    assert not any("可继续训练前检查" in message.value for message in page.success)
    # 报告横幅与 CLI 同源(R95):状态句来自 summarize_preflight(「训练前检查通过：」,
    # 区别于就绪行的「已通过」),手抄短句「当前数据与 token 消费检查通过」退场。
    assert any("训练前检查通过：" in message.value for message in page.success), (
        "preflight_passed 态的报告横幅必须渲染单源状态句"
    )
    assert not any("当前数据与 token 消费检查通过" in message.value for message in page.success)
    assert any("不代表训练效果或业务达标" in cap.value for cap in page.caption)
    # 中段建议行同源铺出(r95-reviewer nit-1 采纳):无截断确认行在场
    assert any("没有内容因长度超限被截断" in entry.value for entry in page.markdown), (
        "passed 报告的无截断确认行必须在场"
    )
    # 方案区与盲标核验在该态继续渲染:通过预检不是终点,是训练侧入口
    assert any("用当前数据微调模型" in header.value for header in page.subheader)
    assert any("盲标核验" in header.value for header in page.subheader)
    # 固定题集 caption 在该态渲染(r94-reviewer nit-3):新测试数据集未冻结题集,
    # elif 分支的静态引导文案不因状态前移而消失。
    assert any("先固定当前开发与测试题目" in cap.value for cap in page.caption), (
        "preflight_passed 态的固定题集引导 caption 必须在场"
    )
    page.run()
    assert not page.exception


def test_comparison_warnings_render_shared_builders_not_hand_copies():
    """R102 源码扫描钉(r102-reviewer nit-1):对照区截断/重复输出警告必须经
    truncation_warning_sentence / dominant_warning_sentence 渲染——与语言化摘要
    同一函数返回值;页面手抄句「检测到高比例输出截断」退场后不得用任何新词
    再抄第二份(运行时钉只拦措辞漂移,拦不住重新手抄)。"""
    source = PAGE.read_text(encoding="utf-8")
    assert "truncation_warning_sentence(" in source
    assert "dominant_warning_sentence(" in source
    assert "检测到高比例输出截断" not in source


def test_acceptance_and_plan_names_source_from_report_summary_constants():
    """验收结论横幅与方案状态标签单源化(R97):页面不再手抄第二套名字——
    验收 map 换 ACCEPTANCE_DECISION_NAMES、方案 labels 换 PLAN_STATUS_NAMES;
    手抄句退场(此前已与摘要漂移:「已冻结」vs「运行前冻结」,同屏两种结论措辞)。"""
    source = PAGE.read_text(encoding="utf-8")
    assert "ACCEPTANCE_DECISION_NAMES" in source
    assert "PLAN_STATUS_NAMES" in source
    assert "达到已冻结的业务验收标准" not in source
    assert "未达到已冻结的业务验收标准" not in source
    # 对称缺口补钉(r97-reviewer nit-1):新措辞也不得手抄回页面——横幅值必须经
    # 常量运行时取,防未来用新词再抄第二份 map 而现有四钉全拦不住。
    assert "达到运行前冻结的验收标准" not in source
    # 方案三态标签的字面量定义只剩 report_summary.py 一处,页面不再留第二份
    assert '"方案可供确认"' not in source


def test_sidebar_listing_survives_corrupt_session_row(tmp_path, monkeypatch):
    """侧栏任务列表守卫钉（R131）：一行坏快照不再打崩整页——完好任务照常列出、
    新建表单可用、损坏如实 warning（结论→为什么），不落回「无任务」失实空态。"""
    import sqlite3

    import ui.config

    for name in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv(f"TUNESMITH_AGENT_{name}", raising=False)
    monkeypatch.setattr(ui.config, "PROJECT_ROOT", tmp_path)
    service = IntakeService(tmp_path / "outputs/workbench/intake")
    session = service.create("根据客户首次描述预测类别", "工单.csv", CSV)
    with sqlite3.connect(service.database) as connection:
        connection.execute(
            "INSERT INTO sessions (id, revision, snapshot) VALUES (?, 1, ?)",
            ("f" * 32, "{ 这不是 JSON"),
        )
    page = AppTest.from_file(str(PAGE), default_timeout=20)
    page.run()
    assert not page.exception
    assert any(area.label == "希望模型完成什么业务工作？" for area in page.text_area), (
        "侧栏读取失败不得打崩新建任务表单"
    )
    assert any("损坏" in w.value and "跳过" in w.value for w in page.warning)
    # options 走 format_func（goal · id[:6] 显示串），按显示串断言
    assert (
        f"{session.goal[:35]} · {session.session_id[:6]}"
        in page.selectbox(key="intake_select").options
    ), "完好任务必须保留在列表里（逐行抢救，不是整列丢弃）"


def test_corrupt_selected_task_recovers_with_in_content_exit(tmp_path, monkeypatch):
    """选中任务读取守卫钉（R131）：快照损坏的任务不再裸栈整页——如实报错 +
    页内「新建数据任务」出口（复用 new_task 清理逻辑），点击后回到新建表单。"""
    import sqlite3

    import ui.config

    for name in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv(f"TUNESMITH_AGENT_{name}", raising=False)
    monkeypatch.setattr(ui.config, "PROJECT_ROOT", tmp_path)
    service = IntakeService(tmp_path / "outputs/workbench/intake")
    session = service.create("根据客户首次描述预测类别", "工单.csv", CSV)
    with sqlite3.connect(service.database) as connection:
        connection.execute(
            "UPDATE sessions SET snapshot=? WHERE id=?", ("{ 坏的", session.session_id)
        )
    page = AppTest.from_file(str(PAGE), default_timeout=20)
    page.session_state["intake_id"] = session.session_id
    page.run()
    assert not page.exception
    assert any("读取任务失败" in e.value for e in page.error), "损坏必须如实报错，不裸栈"
    assert not any("暂无" in w.value for w in page.warning), "读不到 ≠ 没有：不得失实空态"
    next(b for b in page.button if b.label == "🆕 新建数据任务").click().run()
    assert not page.exception
    assert any(area.label == "希望模型完成什么业务工作？" for area in page.text_area), (
        "页内出口必须回到新建任务表单"
    )


def test_task_view_survives_corrupt_iterations_row(tmp_path, monkeypatch):
    """迭代记录读取守卫钉（R131）：iterations 表一行坏 JSON 不再打崩任务视图——
    页面主结构照常渲染，降级如实 warning。"""
    import sqlite3

    import ui.config
    from src.workbench.iterations import IterationService

    for name in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv(f"TUNESMITH_AGENT_{name}", raising=False)
    monkeypatch.setattr(ui.config, "PROJECT_ROOT", tmp_path)
    service = IntakeService(tmp_path / "outputs/workbench/intake")
    session = service.create("根据客户首次描述预测类别", "工单.csv", CSV)
    iteration_service = IterationService(
        tmp_path / "outputs/workbench/iterations",
        tmp_path / "outputs/workbench/training",
        tmp_path / "outputs/workbench/evaluations",
    )
    with sqlite3.connect(iteration_service.database) as connection:
        connection.execute(
            "INSERT INTO iterations (id, revision, snapshot) VALUES (?, 1, ?)",
            ("e" * 32, "{ 坏的"),
        )
    page = AppTest.from_file(str(PAGE), default_timeout=20)
    page.session_state["intake_id"] = session.session_id
    page.run()
    assert not page.exception
    assert any(h.value == session.goal for h in page.subheader), "任务视图主结构必须照常渲染"
    assert any("迭代记录读取失败" in w.value for w in page.warning)


def test_task_view_survives_corrupt_suite_file(tmp_path, monkeypatch):
    """评测套件读取守卫钉（R131）：evaluation-suites 目录一个坏 JSON 不再打崩任务视图。"""
    import ui.config

    for name in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv(f"TUNESMITH_AGENT_{name}", raising=False)
    monkeypatch.setattr(ui.config, "PROJECT_ROOT", tmp_path)
    service = IntakeService(tmp_path / "outputs/workbench/intake")
    session = service.create("根据客户首次描述预测类别", "工单.csv", CSV)
    suites_dir = tmp_path / "outputs/workbench/evaluation-suites"
    suites_dir.mkdir(parents=True)
    (suites_dir / f"{'a' * 64}.json").write_text("{ 坏的", encoding="utf-8")
    page = AppTest.from_file(str(PAGE), default_timeout=20)
    page.session_state["intake_id"] = session.session_id
    page.run()
    assert not page.exception
    assert any(h.value == session.goal for h in page.subheader)
    assert any("评测套件列表读取失败" in w.value for w in page.warning)


def test_tokenizer_input_explains_what_to_type():
    """tokenizer 输入白话解释钉（R132）：训练前检查的 tokenizer 输入必须自带
    可照抄的示例（placeholder）+ 帮助说明（help）——非专家不看文档也应知道
    填什么；空值报错同样给示例，不再只说「请填写」。"""
    source = PAGE.read_text(encoding="utf-8")
    assert 'placeholder="例如 Qwen/Qwen3-1.7B' in source, "输入框必须带可照抄示例"
    assert "分词器" in source and "不自动下载" in source, "help 必须白话解释是什么+边界"
    assert "如 Qwen/Qwen3-1.7B" in source, "空值报错必须给可照抄示例"
