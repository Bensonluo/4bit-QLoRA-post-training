"""Verification handoffs preserve inputs and the selected disclosure boundary."""

import csv
from io import BytesIO, StringIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from src.workbench.label_sheet import (
    verification_questions_to_csv,
    verification_questions_to_xlsx,
)


@pytest.mark.parametrize(
    "value", ["=1+2", "+1", "-1", "@SUM(A1)", "\t=1+2", "＝1+2", "#N/A", "😀\r\n行"]
)
def test_reproduce_sheet_preserves_literal_input_without_exposing_labels(value):
    items = [{"row_id": "r000001", "input": value, "answer": "NEVER_DISCLOSE_LABEL"}]
    payload = verification_questions_to_xlsx(items, mode="reproduce")
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook.active
    assert sheet.title == "盲标核验"
    assert sheet["A2"].value == "r000001"
    assert sheet["B2"].value == value
    assert sheet["B2"].data_type == "s"
    assert sheet["C2"].value is None
    with ZipFile(BytesIO(payload)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in xml
    assert b"NEVER_DISCLOSE_LABEL" not in xml
    raw = verification_questions_to_csv(items, mode="reproduce").decode("utf-8-sig")
    assert "NEVER_DISCLOSE_LABEL" not in raw
    assert list(csv.reader(StringIO(raw)))[1] == ["r000001", value, ""]


def test_review_sheet_includes_original_answer_as_text_and_empty_decision():
    items = [{"row_id": "r000002", "input": "=客户描述", "answer": "=可能的答案\r\n原文"}]
    payload = verification_questions_to_xlsx(items, mode="review")
    sheet = load_workbook(BytesIO(payload), data_only=False).active
    assert [cell.value for cell in sheet[1]] == [
        "行ID",
        "题目输入",
        "数据中的答复（供审阅）",
        "审阅结论（填写是或否）",
    ]
    assert [cell.value for cell in sheet[2]] == [
        "r000002",
        "=客户描述",
        "=可能的答案\r\n原文",
        None,
    ]
    assert sheet["B2"].data_type == sheet["C2"].data_type == "s"
    raw = verification_questions_to_csv(items, mode="review").decode("utf-8-sig")
    assert list(csv.reader(StringIO(raw)))[1] == ["r000002", "=客户描述", "=可能的答案\r\n原文", ""]


@pytest.mark.parametrize("export", [verification_questions_to_csv, verification_questions_to_xlsx])
def test_missing_review_answer_and_invalid_modes_are_rejected(export):
    items = [{"row_id": "r000001", "input": "题目"}]
    with pytest.raises(ValueError, match="缺少数据中的答复"):
        export(items, mode="review")
    with pytest.raises(ValueError, match="不支持的核验方式"):
        export(items, mode="unknown")


@pytest.mark.parametrize("value", ["a" * 32768, "😀" * 16384, "bad\x00value", "bad\ud800value"])
def test_excel_rejects_unrepresentable_input_instead_of_truncating(value):
    with pytest.raises(ValueError, match="Excel"):
        verification_questions_to_xlsx([{"row_id": "r1", "input": value}], mode="reproduce")


def test_excel_rejects_too_many_questions_before_iterating():
    class TooManyRows(list):
        def __len__(self):
            return 1_048_576

        def __iter__(self):
            pytest.fail("oversized export should be rejected before reading rows")

    with pytest.raises(ValueError, match="行列限制"):
        verification_questions_to_xlsx(TooManyRows())


@pytest.mark.parametrize("export", [verification_questions_to_csv, verification_questions_to_xlsx])
def test_invalid_question_identity_or_input_is_rejected(export):
    with pytest.raises(ValueError, match="缺少行ID或输入文本"):
        export([{"row_id": "r1", "input": None}], mode="reproduce")
