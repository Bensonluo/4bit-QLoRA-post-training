"""Spreadsheet handoffs preserve literal input and never create formulas."""

import csv
import io
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from src.workbench.answer_sheet import answer_sheet_csv, answer_sheet_xlsx
from src.workbench.sources import read_source
from tests.unit.test_answer_sheet_cli import MISSING, _store, invoke


@pytest.mark.parametrize(
    "value",
    ["=1+2", "+1", "-1", "@SUM(A1)", "\t=1+2", "＝1+2", "#N/A", "😀\r\n行", "a" * 32767],
)
def test_excel_preserves_untrusted_inputs_as_literal_text(value):
    row = SimpleNamespace(row_id="r000001", input=value, original={})
    payload = answer_sheet_xlsx([row], ["类别"])
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    sheet = workbook.active
    assert sheet["B2"].value == value
    assert sheet["B2"].data_type == "s"
    assert sheet["C2"].value is None
    assert sheet.freeze_panes == "B2"
    assert sheet["A1"].font.bold
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in xml
    assert (
        list(csv.reader(io.StringIO(answer_sheet_csv([row], ["类别"]).decode("utf-8-sig"))))[1][1]
        == value
    )


@pytest.mark.parametrize("value", ["a" * 32768, "😀" * 16384, "bad\x00text", "bad\ud800text"])
def test_excel_refuses_to_truncate_or_silently_change_invalid_input(value):
    row = SimpleNamespace(row_id="r000001", input=value, original={})
    with pytest.raises(ValueError, match="Excel"):
        answer_sheet_xlsx([row], ["类别"])


def test_cli_exports_excel_and_csv_without_changing_source(tmp_path):
    service, session = _store(tmp_path, MISSING)
    before = service.load(session.session_id).model_dump()
    xlsx, raw = tmp_path / "sheet.xlsx", tmp_path / "sheet.csv"
    done = invoke(
        service, "answer-sheet", session.session_id, "--export-xlsx", xlsx, "--export-csv", raw
    )
    assert done.returncode == 0, done.stderr
    sheet = load_workbook(xlsx, data_only=False).active
    assert sheet["A2"].value == "r000002"
    assert "物流未更新" in sheet["B2"].value
    assert sheet["B2"].data_type == "s"
    assert raw.exists()
    assert service.load(session.session_id).model_dump() == before


def test_excel_reader_does_not_treat_literal_formula_as_formula_or_error():
    row = SimpleNamespace(row_id="r000001", input="=SUM(1,2)", original={})
    payload = answer_sheet_xlsx([row], ["类别"])
    source = read_source("sheet.xlsx", payload)
    assert source.rows[0].values["题目输入（模型将看到的内容）"] == "=SUM(1,2)"
    assert not source.formula_note


def test_cli_rejects_overlapping_export_paths_without_writing(tmp_path):
    service, session = _store(tmp_path, MISSING)
    target = tmp_path / "same"
    done = invoke(
        service, "answer-sheet", session.session_id, "--export-xlsx", target, "--export-csv", target
    )
    assert done.returncode == 2
    assert "导出路径不能相同" in done.stderr
    assert not target.exists()


def test_cli_long_excel_input_fails_before_writing_either_export(tmp_path):
    source = MISSING.replace("物流未更新".encode(), ("长" * 32768).encode())
    service, session = _store(tmp_path, source)
    xlsx, raw = tmp_path / "sheet.xlsx", tmp_path / "sheet.csv"
    done = invoke(
        service, "answer-sheet", session.session_id, "--export-xlsx", xlsx, "--export-csv", raw
    )
    assert done.returncode == 2
    assert "超过 Excel 单元格长度限制" in done.stderr
    assert not xlsx.exists() and not raw.exists()
    assert service.load(session.session_id).revision == session.revision
