"""Excel error values are rejected with the original cell location."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from src.workbench.intake_service import IntakeService
from src.workbench.sources import read_source
from tests.unit.test_full_data import approved


def workbook_with(*rows):
    workbook = Workbook()
    workbook.active.title = "普通"
    sheet = workbook.create_sheet(" 业务表 ")
    sheet.append(["客户描述", "类别"])
    for row in rows:
        sheet.append(row)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


@pytest.mark.parametrize("error", ["#DIV/0!", "#N/A", "#REF!"])
def test_excel_error_value_reports_original_sheet_and_cell(error, tmp_path):
    data = workbook_with(["正常", "质量"], ["需要核对", error])
    service = IntakeService(tmp_path / "intake")
    with pytest.raises(ValueError) as raised:
        service.create("判断类型", "source.xlsx", data, sheet=" 业务表 ")
    message = str(raised.value)
    assert " 业务表 " in message
    assert "B3" in message
    assert "类别" in message
    assert error in message
    assert service.list_sessions() == []


def test_blank_and_uncached_formula_remain_distinct_from_excel_errors():
    data = workbook_with(["普通缺失", ""], ["未计算公式", "=1+1"])
    source = read_source("source.xlsx", data, sheet=" 业务表 ")
    assert source.rows[0].values["类别"] == ""
    assert source.rows[1].values["类别"] == ""
    assert "B3" in source.formula_note


def test_literal_error_text_in_csv_is_not_reinterpreted():
    source = read_source("source.csv", "客户描述,类别\n文字,#DIV/0!\n".encode())
    assert source.rows[0].values["类别"] == "#DIV/0!"


def test_excel_error_header_has_its_own_coordinate():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "业务"
    sheet.append(["客户描述", "#REF!"])
    sheet.append(["正常", "质量"])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match=r"业务.*B1.*#REF!"):
        read_source("source.xlsx", output.getvalue())


def test_full_upload_error_does_not_change_confirmed_session(tmp_path):
    service = IntakeService(tmp_path / "intake")
    before = approved(service)
    data = workbook_with(["出错", "#DIV/0!"])
    with pytest.raises(ValueError, match=r" 业务表 .*B2.*#DIV/0!"):
        service.validate_full_data(
            before.session_id, before.revision, "full.xlsx", data, sheet=" 业务表 "
        )
    after = service.load(before.session_id)
    assert after.model_dump() == before.model_dump()
    assert not (service.root / before.session_id / "full").exists()


def test_literal_error_text_in_excel_is_not_an_error_cell():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["客户描述", "类别"])
    sheet.append(["文字", "#DIV/0!"])
    sheet["B2"].data_type = "s"
    output = BytesIO()
    workbook.save(output)
    source = read_source("source.xlsx", output.getvalue())
    assert source.rows[0].values["类别"] == "#DIV/0!"


def test_formula_cached_as_error_is_rejected_at_the_cell():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "业务"
    sheet.append(["客户描述", "类别"])
    sheet.append(["错误结果", "=1/0"])
    original = BytesIO()
    workbook.save(original)
    output = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                old = b'<c r="B2"><f>1/0</f><v /></c>'
                assert old in content
                content = content.replace(old, b'<c r="B2" t="e"><f>1/0</f><v>#DIV/0!</v></c>')
            target.writestr(info, content)
    with pytest.raises(ValueError, match=r"业务.*B2.*#DIV/0!"):
        read_source("source.xlsx", output.getvalue())
